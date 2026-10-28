"""Admin → Warehouse → Inventory (warehouse.py) — wiring checks.

Run: python backend/tests/test_warehouse_inventory.py
Every change is a ledger movement with who/why; nothing deletes.
"""
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
failed = 0


def check(cond, label):
    global failed
    print(("ok   " if cond else "FAIL ") + label)
    if not cond:
        failed += 1


src = open(os.path.join(ROOT, "backend", "warehouse.py"), encoding="utf-8").read()
code = re.sub(r'"""[\s\S]*?"""', "", src)
inv = code[code.index("# ------------------------------------------------- warehouse inventory ---"):code.index('@wh_admin_router.get("/docs")')]
check('@wh_admin_router.get("/inventory")' in inv and '@wh_admin_router.get("/inventory/export.csv")' in inv,
      "search/read list + CSV export (under Admin → Warehouse permission)")
cnt = inv[inv.index("async def adm_set_count"):inv.index("async def adm_inventory_move")]
check('"correction_in" if diff > 0 else "correction_out"' in cnt and "actor=user" in cnt and "len(note) < 3" in cnt,
      "set count = a correction movement with the person and a required reason")
check('"product_type": {"$nin": list(NOT_COUNTED)}' in cnt, "packs/hampers can't be counted as one item")
mv = inv[inv.index("async def adm_inventory_move"):]
check('(IN_REASONS | OUT_REASONS) - {"opening"}' in mv and "0 < body.qty <= 100000" in mv, "in/out moves: known reasons, sane quantities")
inv_code = "\n".join(x for x in inv.splitlines() if not x.strip().startswith("#"))
check("delete" not in inv_code.lower() and "@wh_admin_router.delete" not in inv_code, "no delete — a wrong movement is undone in Movements")
check('"WAREHOUSE_INVENTORY_EXPORT"' in inv and '"WAREHOUSE_COUNT_SET"' in inv, "exports and counts are audited")
app = open(os.path.join(ROOT, "frontend", "src", "pages", "admin", "AdminWarehouse.jsx"), encoding="utf-8").read()
check('["inventory", "Inventory"]' in app and "<WarehouseInventory" in app, "Inventory tab in Admin → Warehouse")
print("-- shareable files --")
sys.path.insert(0, os.path.join(ROOT, "backend"))
import io  # noqa: E402
import inventory_reports as ir  # noqa: E402
rows = [{"id": "a", "title": "=HYPERLINK(evil)", "author": "A", "isbn": "9780000000011", "category": "law", "price": 476,
         "original_price": 595, "stock": 0, "wh_stock": -2, "pod": True, "owed": 1, "diff": -2, "moves": 1},
        {"id": "b", "title": "Posh Act An Overview (Tamil) போஷ் சட்டம்", "author": "B", "isbn": "9780000000028",
         "category": "law", "price": 340, "stock": 5, "wh_stock": 5, "owed": 0, "diff": None, "moves": 0},
        {"id": "c", "title": "Gone", "category": "tax", "price": 100, "stock": 0, "wh_stock": 0, "moves": 0}]
x = ir.inventory_xlsx(rows, 10, "trial", {"law": "Law", "tax": "Tax"})
import openpyxl  # noqa: E402
wb = openpyxl.load_workbook(io.BytesIO(x))
check(wb.sheetnames == ["Summary", "Inventory"] and wb["Inventory"]["A2"].value == "'=HYPERLINK(evil)", "xlsx: Summary + Inventory; formula-safe cells")
check(ir.availability(rows[0]) == "Print on demand" and ir.availability(rows[2]) == "Out of stock", "brochure availability words")
check(ir.latin(rows[1]["title"]) == "Posh Act An Overview (Tamil)", "non-Latin script trimmed for the PDF font")
pdf = ir.brochure_pdf(rows[:2], {"law": "Law"}, None)
try:
    import pypdfium2 as pdfium
    txt = "".join(pg.get_textpage().get_text_range() for pg in pdfium.PdfDocument(pdf))
    check("Rs 595" in txt and "Print on demand" in txt and "-2" not in txt.replace("5-7", ""), "brochure: MRP + availability, no internal counts")
except ImportError:
    check(pdf[:5] == b"%PDF-", "brochure builds")
check('"WAREHOUSE_BROCHURE"' in inv and "_storage_key(" in inv and "no fetching the web" in src, "brochure: audited; covers only from our own storage")
print()
if failed:
    print(f"{failed} check(s) failed")
    sys.exit(1)
print("all checks passed")
