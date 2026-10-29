"""Marketing — contacts, free email verification, campaigns (email via Amazon
SES, WhatsApp via Interakt), tracking and the analytics behind Admin → Marketing.

Pure rules live in marketing_core.py (tested without a database).

EMAIL (Amazon SES)
  * Sent from the address in Settings (info@oakbridge.in) through SES in
    SES_REGION (falls back to S3_REGION / AWS_REGION / ap-south-1), with the
    app's existing AWS keys. Optional SES_CONFIGURATION_SET routes SES's
    delivery / bounce / complaint events to an SNS topic subscribed to
    /api/m/ses-events/<token> (Settings shows the exact URL).
  * Opens and clicks are tracked by US (a 1px image and signed redirect links),
    not by SES, so no SES tracking setup is needed and links are not wrapped
    twice. Leave open/click tracking OFF on the configuration set.
  * Every email carries List-Unsubscribe + one-click (Gmail / Yahoo rules).

VERIFICATION (free, layered — see marketing_core.verdict)
  syntax / typos / role / throwaway  ->  domain MX lookup over DNS-over-HTTPS
  (cached per domain)  ->  "proven" addresses (OTP-confirmed accounts,
  delivered / opened before)  ->  suppression (bounced, complained,
  unsubscribed). Sending goes verified -> valid -> risky (risky only if asked)
  and AUTO-PAUSES above 2% bounces / 0.1% complaints.

CONSENT
  A contact is mailed only with email_consent = subscribed, messaged on
  WhatsApp only with wa_consent = subscribed. Sources: newsletter sign-up,
  checkout tick, or an import where the admin confirms the people agreed.
  Unsubscribe / STOP / complaint switch it off for good.

WHATSAPP campaigns go through interakt.send_campaign_message (Interakt's
mode switch applies: Off refuses, Test sends to the test number).
"""
from __future__ import annotations

import asyncio
import csv
import hashlib
import io
import json
import logging
import os
import re
import secrets
import time
import uuid
from datetime import datetime, timedelta, timezone
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.utils import formataddr, make_msgid
from typing import Optional
from urllib.parse import urlparse

import requests
from fastapi import APIRouter, Depends, File, Form, Header, HTTPException, Request, UploadFile

import dbutil
import jobs as _jobs
from fastapi.responses import RedirectResponse, Response
from pydantic import BaseModel

import marketing_core as core
from audit import audit_log
from extensions import db, require_admin, require_superadmin

log = logging.getLogger(__name__)

PUBLIC_API = (os.environ.get("PUBLIC_API_URL") or "https://api.oakbridge.in").rstrip("/")
SITE = (os.environ.get("SITE_URL") or "https://www.oakbridge.in").rstrip("/")
SES_REGION = (os.environ.get("SES_REGION") or os.environ.get("S3_REGION") or os.environ.get("AWS_REGION") or "ap-south-1").strip()
SES_CONFIG_SET = (os.environ.get("SES_CONFIGURATION_SET") or "").strip()

DEFAULT_SETTINGS = {
    "from_name": "Oakbridge Publishing",
    "from_email": "info@oakbridge.in",
    "reply_to": "",
    "footer": "Oakbridge Publishing Pvt. Ltd. · 934, 9th Floor, Tower B3, Spaze Itech Park, Sector 49, Gurugram 122018",
    "logo": "",
    "double_opt_in": True,
    # info@ / contact@ / admin@ at a school, college or firm is often the real
    # buyer for a publisher, so by default they count as valid. Off = risky
    # (shared inbox: whoever reads it never opted in personally).
    "include_role_addresses": True,
    "max_bounce": 0.02,
    "wa_rate_marketing": 0.96,   # ₹ per message incl. Interakt markup, before GST
    "wa_rate_utility": 0.13,
    "extra_disposable": [],
    # Amazon SES Email Validation (paid, ~$0.01 per address): checks whether
    # the mailbox exists before anyone is mailed. The monthly ₹ budget stops a
    # big import from running up a bill; anyone left unchecked is skipped
    # (with the reason shown), never mailed blind.
    "ses_validation": True,
    "ses_budget_inr": 850,       # ≈ 1,000 checks a month
    "usd_inr": 88,               # for ₹ estimates of $-priced AWS charges
    # What happens to each kind of 'risky' address when a campaign is sent:
    # send | tail (last, stops itself if it bounces) | skip.
    "risk_policy": dict(core.DEFAULT_RISK_POLICY),
    "tail_max_bounce": 0.03,
    # Confidence score lines (0-100) and self-learning from bounce outcomes.
    "score_valid": 65,
    "score_invalid": 35,
    "auto_learn": True,
    # Retire inactive addresses: counted in campaign emails since the last
    # open / click. Abandoned inboxes become spam traps. 0 = rule off.
    "inactive_after": 6,
    "sunset_after": 12,
}
DISPOSABLE_URL = "https://raw.githubusercontent.com/disposable-email-domains/disposable-email-domains/main/disposable_email_blocklist.conf"
DISPOSABLE_ALLOW_URL = "https://raw.githubusercontent.com/disposable-email-domains/disposable-email-domains/main/allowlist.conf"
RESEND_WEBHOOK_SECRET = (os.environ.get("RESEND_WEBHOOK_SECRET") or "").strip()
RECONFIRM_BLOCKS = [
    {"type": "heading", "text": "Do you still want to hear from us?"},
    {"type": "text", "text": "Hi {{first_name}}, we only email people who want our news — new books, events and offers. "
                             "Click below to stay on the list. If you don't, we'll stop emailing you in 30 days."},
    {"type": "button", "label": "Yes, keep me subscribed", "url": "https://www.oakbridge.in/"},
]
SES_CHECK_TTL_DAYS = 180  # a mailbox verdict is reused for 6 months, never paid for twice
MAX_BOUNCE_CEILING = 0.04  # our stops must trip before AWS's 5% account review (values saved above it are capped when used)
NO_VERDICT_RETRY_DAYS = 30  # SES answered without a verdict: ask again after this, not on every pass
GIF = (b"GIF89a\x01\x00\x01\x00\x80\x00\x00\x00\x00\x00\xff\xff\xff!\xf9\x04\x01\x00\x00\x00\x00,"
       b"\x00\x00\x00\x00\x01\x00\x01\x00\x00\x02\x02D\x01\x00;")


def _now() -> datetime:
    return datetime.now(timezone.utc)


def _iso() -> str:
    return _now().isoformat()


def _secret() -> str:
    # Derived from the app's JWT secret, so no extra env var is needed; the
    # salt keeps marketing links unusable as anything else.
    return hashlib.sha256((os.environ.get("JWT_SECRET", "") + "|marketing-links-v1").encode()).hexdigest()


def sns_token() -> str:
    # SES_EVENTS_TOKEN lets the SNS endpoint be rotated on its own (e.g. after
    # the URL shows up in a screenshot) without touching JWT_SECRET, which
    # would sign everyone out and break every unsubscribe link already sent.
    # Ignored when shorter than 24 characters so a typo can't make it guessable.
    override = (os.environ.get("SES_EVENTS_TOKEN") or "").strip()
    if len(override) >= 24:
        return override
    return core.sign(_secret(), "ses-events")


async def _load_settings() -> dict:
    doc = await db.integrations.find_one({"key": "marketing"}, {"_id": 0}) or {}
    return {**DEFAULT_SETTINGS, **{k: v for k, v in doc.items() if k in DEFAULT_SETTINGS}}


_SETTINGS_MEMO = None


async def get_settings(fresh: bool = False) -> dict:
    """Read on almost every verification / send step, so it is memoised for 15 s
    and dropped the moment any admin write bumps the shared cache version
    (cache.py). fresh=True after a save, which runs before that bump."""
    global _SETTINGS_MEMO
    if fresh:
        return await _load_settings()
    if _SETTINGS_MEMO is None:
        from cache import Memo, shared_version
        _SETTINGS_MEMO = Memo(shared_version(db))
    return dict(await _SETTINGS_MEMO.get("marketing_settings", 15, _load_settings))


_indexes = False


async def _index(coll, keys, **kw) -> bool:
    """Create one index; a failure is logged and never stops the others (a
    single failing index used to skip every index listed after it)."""
    try:
        await coll.create_index(keys, **kw)
        return True
    except Exception:  # noqa: BLE001
        log.warning("marketing: index %s on %s failed", keys, coll.name, exc_info=True)
        return False


async def _merge_duplicates(coll, field: str, sum_field: Optional[str] = None) -> int:
    """Collapse documents sharing `field` into one (summing `sum_field` if
    given), so a unique index can be built. Needed when rows were upserted
    concurrently before the unique index existed — e.g. two SES checks starting
    in the same instant each created this month's usage counter, which split
    the count and let the monthly budget be exceeded."""
    merged = 0
    grp = {"_id": f"${field}", "ids": {"$push": "$_id"}, "c": {"$sum": 1}}
    if sum_field:
        grp["total"] = {"$sum": {"$ifNull": [f"${sum_field}", 0]}}
    async for g in coll.aggregate([{"$group": grp}, {"$match": {"c": {"$gt": 1}}}]):
        keep, rest = g["ids"][0], g["ids"][1:]
        if sum_field:
            await coll.update_one({"_id": keep}, {"$set": {sum_field: g["total"]}})
        await coll.delete_many({"_id": {"$in": rest}})
        merged += len(rest)
    if merged:
        log.warning("marketing: merged %s duplicate %s rows on %s", merged, coll.name, field)
    return merged


async def _unique_index(coll, field: str, sum_field: Optional[str] = None, **kw) -> None:
    if not await _index(coll, field, unique=True, **kw):
        await _merge_duplicates(coll, field, sum_field)
        await _index(coll, field, unique=True, **kw)


async def _ensure_indexes() -> None:
    global _indexes
    if _indexes:
        return
    _indexes = True
    await _unique_index(db.mk_contacts, "email_norm", sparse=True)
    await _index(db.mk_contacts, "lists")
    await _index(db.mk_sends, [("campaign_id", 1), ("status", 1)])
    await _index(db.mk_sends, "message_id", sparse=True)
    await _index(db.mk_sends, [("campaign_id", 1), ("contact_id", 1)], unique=True)
    await _unique_index(db.mk_domains, "domain")
    # Lookups and updates by `id` are on every hot path (each send, each
    # contact update); without these every one of them scanned the collection.
    for coll in (db.mk_contacts, db.mk_sends, db.mk_campaigns, db.mk_lists):
        await _index(coll, "id", unique=True)
    await _index(db.mk_alerts, [("resolved", 1), ("at", -1)])
    await _index(db.mk_sends, "sent_at")
    await _index(db.mk_sends, [("campaign_id", 1), ("tail", 1), ("status", 1)])
    await _index(db.mk_contacts, [("email_consent.status", 1), ("email_status", 1)])
    await _index(db.mk_import_rows, [("import_id", 1), ("i", 1)])
    await _unique_index(db.mk_removed, "h")
    await _index(db.mk_contacts, "created_at")
    await _index(db.mk_events, "at")
    await _jobs.ensure_indexes(db)
    await _unique_index(db.mk_ses_checks, "email_norm")
    await _unique_index(db.mk_ses_usage, "month", sum_field="n")
    await _unique_index(db.mk_proven, "email_norm")
    await _unique_index(db.mk_webhook_seen, "id")
    await _index(db.mk_webhook_seen, "at", expireAfterSeconds=7 * 86400)


# ------------------------------------------------------- domain lookups ---
def _doh_mx(domain: str) -> Optional[bool]:
    """Does the domain receive mail? DNS-over-HTTPS (Cloudflare, then Google),
    so no DNS library and no UDP from the host. True / False / None
    (could not tell — never treated as invalid)."""
    for url in ("https://cloudflare-dns.com/dns-query", "https://dns.google/resolve"):
        try:
            r = requests.get(url, params={"name": domain, "type": "MX"}, headers={"accept": "application/dns-json"}, timeout=5)
            if r.status_code != 200:
                continue
            j = r.json()
            status = j.get("Status")
            if status == 3:  # NXDOMAIN: the domain does not exist
                return False
            answers = [a for a in j.get("Answer") or [] if a.get("type") == 15]
            if answers:
                # A "null MX" ("0 .", RFC 7505) means "this domain accepts no mail".
                datas = [str(a.get("data", "")).strip() for a in answers]
                return not all(d in ("0 .", "0", ".") for d in datas)
            # No MX: mail falls back to the A record (RFC 5321).
            r2 = requests.get(url, params={"name": domain, "type": "A"}, headers={"accept": "application/dns-json"}, timeout=5)
            if r2.status_code == 200:
                return bool([a for a in r2.json().get("Answer") or [] if a.get("type") == 1])
        except (requests.RequestException, ValueError):
            continue
    return None


async def _repair_domains(rows: list, mx: dict) -> None:
    """Company-domain typos ("mmtckimited.com" -> "mmtclimited.com"), for
    addresses whose domain doesn't receive email. A fix is used only when it
    is unambiguous AND the corrected domain really receives email:

    1. the closest domain that already works — elsewhere in this file, or one
       we have delivered to before — 1 edit away (2 for long names);
    2. otherwise "one key off" candidates (neighbouring key, dropped or swapped
       letter), checked for a mail server; used only if exactly one works.

    The repaired address is never trusted more than "risky" (a guess, however
    good) unless it is proven later. Bounded: 15 bad domains a batch."""
    bad = sorted({x["syntax"]["email"].split("@")[1] for x in rows
                  if x["syntax"]["ok"] and mx.get(x["syntax"]["email"].split("@")[1]) is False})[:15]
    if not bad:
        return
    known = {d for d, ok in mx.items() if ok}
    async for d in db.mk_domains.find({"delivered_n": {"$gt": 0}}, {"_id": 0, "domain": 1}):
        known.add(d["domain"])
    fixes: dict = {}
    for d in bad:
        k = core.closest_known(d, known)
        if k:
            fixes[d] = (k, "a working domain one letter away")
            continue
        cands = core.domain_typo_candidates(d)
        if not cands:
            continue
        res = await domains_mx(set(cands))
        ok = [c for c in cands if res.get(c)]
        if len(ok) == 1:
            fixes[d] = (ok[0], "one key off, and that domain receives email")
            mx[ok[0]] = True
    for x in rows:
        s = x["syntax"]
        if not s["ok"]:
            continue
        local, dom = s["email"].split("@")
        if dom in fixes:
            new, why = fixes[dom]
            x["fixed_from"] = x["fixed_from"] or s["email"]
            x["syntax"] = {**core.check_syntax(f"{local}@{new}"),
                           "reasons": [r for r in s["reasons"] if r.startswith("fixed")] + [f"fixed domain {dom} -> {new} ({why})"],
                           "domain_guess": True}


async def domains_mx(domains: set) -> dict:
    """{domain: True/False/None}, cached 30 days in mk_domains."""
    out: dict = {}
    if not domains:
        return out
    fresh_after = (_now() - timedelta(days=30)).isoformat()
    async for d in db.mk_domains.find({"domain": {"$in": list(domains)}, "checked_at": {"$gte": fresh_after}}, {"_id": 0}):
        out[d["domain"]] = d.get("mx")
    todo = [d for d in domains if d not in out]
    sem = asyncio.Semaphore(16)

    async def one(d):
        async with sem:
            v = await asyncio.to_thread(_doh_mx, d)
            out[d] = v
            if v is not None:
                await db.mk_domains.update_one({"domain": d}, {"$set": {"domain": d, "mx": v, "checked_at": _iso()}}, upsert=True)

    await asyncio.gather(*(one(d) for d in todo))
    return out


async def _proven(emails_norm: list) -> set:
    """Addresses already shown to be real mailboxes: an OTP-confirmed account,
    or a campaign email that was delivered / opened / clicked."""
    proven = set()
    async for u in db.users.find({"email_verified": True}, {"_id": 0, "email": 1}):
        proven.add(core.normalise_email(u.get("email")))
    async for c in db.mk_contacts.find({"email_norm": {"$in": emails_norm}, "proven": True}, {"_id": 0, "email_norm": 1}):
        proven.add(c["email_norm"])
    # Order / account emails Resend delivered (resend_events) — free proof.
    async for c in db.mk_proven.find({"email_norm": {"$in": emails_norm}}, {"_id": 0, "email_norm": 1}):
        proven.add(c["email_norm"])
    return proven & set(emails_norm)


def _erase_hash(norm: str) -> str:
    return hashlib.sha256((norm or "").encode()).hexdigest()


async def _erased(norms: list) -> set:
    """Which of these mailboxes were erased at a privacy request (we keep only a hash)."""
    by_hash = {_erase_hash(n): n for n in norms if n}
    out = set()
    async for e in db.mk_erased.find({"h": {"$in": list(by_hash)}}, {"_id": 0, "h": 1}):
        out.add(by_hash[e["h"]])
    return out


async def _removed(norms: list) -> dict:
    """{mailbox: reason} for addresses removed in a database clean-up (only a
    hash is kept, like an erase) — so they can never be imported or synced
    back in by accident."""
    by_hash = {_erase_hash(n): n for n in norms if n}
    out = {}
    async for e in db.mk_removed.find({"h": {"$in": list(by_hash)}}, {"_id": 0, "h": 1, "reason": 1}):
        out[by_hash[e["h"]]] = e.get("reason") or "removed"
    return out


async def learn_providers() -> dict:
    """Teach the checker from what actually happened (last 180 days):

    - bad_provider: a provider whose UNPROVEN addresses (anything not
      "verified" when sent) bounced > PROVIDER_MAX_BOUNCE over at least
      PROVIDER_MIN_SENDS sends. Its unproven addresses are then skipped.
      Cleared again when a later window is clean.
    - ses_unverifiable: a provider where SES almost never confirms a mailbox
      (Yahoo, Rediff…). Paid checks there are skipped (see ses_check).

    Pure tallies in Python — volumes are thousands of sends, not millions."""
    since = (_now() - timedelta(days=180)).isoformat()
    tally: dict = {}
    async for x in db.mk_sends.find({"sent_at": {"$gte": since}, "band": {"$ne": "verified"},
                                     "status": {"$in": ["sent", "bounced", "complained"]}},
                                    {"_id": 0, "to": 1, "status": 1}):
        dom = str(x.get("to") or "").lower().rsplit("@", 1)[-1]
        if not dom or "@" not in str(x.get("to") or ""):
            continue
        t = tally.setdefault(dom, [0, 0])
        t[0] += 1
        t[1] += 1 if x.get("status") == "bounced" else 0
    checks: dict = {}
    async for c in db.mk_ses_checks.find({"at": {"$gte": since}}, {"_id": 0, "email_norm": 1, "mailbox": 1, "overall": 1}):
        dom = str(c.get("email_norm") or "").rsplit("@", 1)[-1]
        t = checks.setdefault(dom, [0, 0])
        t[0] += 1
        t[1] += 1 if (c.get("mailbox") == "HIGH" or c.get("overall") == "HIGH") else 0
    flagged, unverif = [], []
    for dom in set(tally) | set(checks):
        sent, bounced = tally.get(dom, [0, 0])
        n_chk, n_ok = checks.get(dom, [0, 0])
        bp, uv = core.bad_provider(sent, bounced), core.ses_unverifiable(n_chk, n_ok)
        await db.mk_domains.update_one({"domain": dom}, {"$set": {
            "domain": dom, "unproven_sent": sent, "unproven_bounced": bounced, "bad_provider": bp,
            "ses_checks": n_chk, "ses_confirmed": n_ok, "ses_unverifiable": uv, "learned_at": _iso()}}, upsert=True)
        if bp:
            flagged.append({"domain": dom, "sent": sent, "bounced": bounced})
        if uv:
            unverif.append(dom)
    return {"bad_providers": sorted(flagged, key=lambda d: -d["bounced"]), "ses_unverifiable": sorted(unverif)}


async def _unverifiable(domains: set) -> set:
    return {d["domain"] async for d in db.mk_domains.find({"domain": {"$in": list(domains)}, "ses_unverifiable": True},
                                                          {"_id": 0, "domain": 1})}


async def mark_contact_proven(contact_id: str) -> None:
    """A campaign email was DELIVERED to, opened or clicked by this contact:
    the mailbox exists, so it is "verified" — sent first, never skipped as
    risky. Until Oct 2026 delivery/open/click set only `proven` and left
    email_status at "risky", so people who had already received our mail were
    skipped as "mailbox unconfirmed" in the next campaign. Never touches a
    suppressed (unsubscribed / bounced / complained) contact."""
    await db.mk_contacts.update_one({"id": contact_id, "suppressed": {"$in": [None, ""]}},
                                    {"$set": {"proven": True, "email_status": "verified", "verified_at": _iso()}})


async def heal_proven() -> int:
    """One-off repair of contacts marked proven by the old code but still
    "risky"/"valid"/"unknown". Idempotent and cheap (indexed fields, usually
    0 rows); run before every audience count and campaign preparation."""
    r = await db.mk_contacts.update_many(
        {"proven": True, "email_status": {"$nin": ["verified", "invalid"]}, "suppressed": {"$in": [None, ""]}},
        {"$set": {"email_status": "verified", "verified_at": _iso()}})
    return getattr(r, "modified_count", 0) or 0


async def note_delivered(norm: str, source: str) -> None:
    """A real email reached this mailbox: remember it (even before the person
    is a marketing contact), mark the contact verified, and count it towards
    its company domain's delivery history the first time."""
    if not norm or "@" not in norm:
        return
    r = await db.mk_proven.update_one({"email_norm": norm}, {"$setOnInsert": {"email_norm": norm, "at": _iso(), "source": source}}, upsert=True)
    await db.mk_contacts.update_one({"email_norm": norm, "suppressed": {"$in": [None, ""]}},
                                    {"$set": {"proven": True, "email_status": "verified", "verified_at": _iso()}})
    dom = norm.split("@")[-1]
    if r.upserted_id is not None and dom not in core.WEBMAIL:
        await db.mk_domains.update_one({"domain": dom}, {"$inc": {"delivered_n": 1}}, upsert=True)


_DISP: dict = {"at": 0.0, "set": frozenset()}


async def disposable_domains() -> frozenset:
    """The maintained public throwaway-domain list (refreshed weekly by
    refresh_disposable), cached in memory for an hour."""
    if time.time() - _DISP["at"] > 3600:
        d = await db.integrations.find_one({"key": "disposable_list"}, {"_id": 0, "domains": 1}) or {}
        _DISP.update(at=time.time(), set=frozenset(d.get("domains") or []))
    return _DISP["set"]


async def refresh_disposable(force: bool = False) -> dict:
    """Weekly: download the community-maintained throwaway-domain list (and
    its allowlist). Never lets it block a big provider or our own domain, and
    ignores a download that looks broken (fewer than 1,000 domains)."""
    cur = await db.integrations.find_one({"key": "disposable_list"}, {"_id": 0, "at": 1, "count": 1}) or {}
    if not force and (cur.get("at") or "") > (_now() - timedelta(days=7)).isoformat():
        return {"skipped": True, "count": cur.get("count", 0)}
    try:
        block = await asyncio.to_thread(lambda: requests.get(DISPOSABLE_URL, timeout=20))
        allow = await asyncio.to_thread(lambda: requests.get(DISPOSABLE_ALLOW_URL, timeout=20))
        block.raise_for_status()
        domains = core.parse_domain_list(block.text)
        if allow.status_code == 200:
            domains -= core.parse_domain_list(allow.text)
    except Exception as e:  # noqa: BLE001
        await raise_alert("disposable", f"Could not update the throwaway-domain list: {str(e)[:120]}", key="disposable")
        return {"ok": False}
    st = await get_settings()
    own = {(st.get("from_email") or "").split("@")[-1], "oakbridge.in"}
    domains -= set(core.WEBMAIL) | own
    if len(domains) < 1000:
        await raise_alert("disposable", f"Throwaway-domain list download looked wrong ({len(domains)} domains) — kept the old one.", key="disposable")
        return {"ok": False}
    await db.integrations.update_one({"key": "disposable_list"}, {"$set": {"key": "disposable_list", "domains": sorted(domains),
                                                                           "count": len(domains), "at": _iso()}}, upsert=True)
    _DISP["at"] = 0.0
    await db.mk_alerts.update_many({"key": "disposable", "resolved": False}, {"$set": {"resolved": True, "resolved_at": _iso()}})
    return {"ok": True, "count": len(domains)}


# ------------------------------------------------ SES email validation ---
_SES_STOP: dict = {}  # last time live checks stopped early: {"why", "at"}


def _month() -> str:
    return _now().strftime("%Y-%m")


_SES_SUPPORTED: Optional[bool] = None


def ses_validation_supported() -> bool:
    """The sesv2 GetEmailAddressInsights call only exists in boto3/botocore
    released after SES Email Validation launched (Jan 2026). An older cached
    install on Render would raise AttributeError mid-import, so check first.
    The installed SDK can't change while the process runs, so check once."""
    global _SES_SUPPORTED
    if _SES_SUPPORTED is None:
        try:
            _SES_SUPPORTED = hasattr(_ses(), "get_email_address_insights")
        except Exception:  # noqa: BLE001
            return False
    return _SES_SUPPORTED


def _cap(st: dict) -> int:
    return core.checks_for_budget(float(st.get("ses_budget_inr") or 0), float(st.get("usd_inr") or 0))


async def ses_usage() -> dict:
    st = await get_settings()
    d = await db.mk_ses_usage.find_one({"month": _month()}, {"_id": 0}) or {}
    used = int(d.get("n", 0))
    cap = _cap(st)
    rate = float(st.get("usd_inr") or 0)
    return {"month": _month(), "used": used, "cap": cap, "left": max(0, cap - used),
            "enabled": bool(st.get("ses_validation")), "supported": ses_validation_supported(),
            "budget_inr": float(st.get("ses_budget_inr") or 0), "spent_inr": round(used * 0.01 * rate, 2),
            "per_check_inr": round(0.01 * rate, 2), "near_limit": cap > 0 and used >= 0.8 * cap,
            "last_stop": _SES_STOP or None}


# --------------------------------------------------------------- alerts ---
async def raise_alert(kind: str, message: str, *, campaign_id: Optional[str] = None, key: Optional[str] = None) -> None:
    """Something an admin should look at, shown on the Marketing dashboard.
    One open alert per key, refreshed rather than repeated."""
    try:
        key = key or f"{kind}:{campaign_id or ''}"
        await db.mk_alerts.update_one(
            {"key": key, "resolved": False},
            {"$set": {"kind": kind, "message": message[:300], "campaign_id": campaign_id, "at": _iso()},
             "$setOnInsert": {"id": str(uuid.uuid4()), "created_at": _iso()}}, upsert=True)
    except Exception:  # noqa: BLE001
        log.warning("marketing: could not store alert %s", kind, exc_info=True)


_bg_tasks: set = set()


def _bg(coro) -> None:
    """Fire-and-forget, kept referenced so it can't be garbage-collected mid-run."""
    t = asyncio.create_task(coro)
    _bg_tasks.add(t)
    t.add_done_callback(_bg_tasks.discard)


async def _reserve_check(cap: int) -> bool:
    """Count one paid check against this month's cap, atomically, BEFORE the
    call — two workers can't both take the last slot."""
    m = _month()
    try:
        await db.mk_ses_usage.update_one({"month": m}, {"$setOnInsert": {"month": m, "n": 0}}, upsert=True)
    except Exception as e:  # noqa: BLE001 — a simultaneous check created it first: fine
        if "E11000" not in str(e) and "duplicate key" not in str(e).lower():
            raise
    r = await db.mk_ses_usage.update_one({"month": m, "n": {"$lt": cap}}, {"$inc": {"n": 1}})
    return r.modified_count == 1


async def _release_check() -> None:
    await db.mk_ses_usage.update_one({"month": _month(), "n": {"$gt": 0}}, {"$inc": {"n": -1}})


async def _ses_cached(norms: list) -> dict:
    fresh = (_now() - timedelta(days=SES_CHECK_TTL_DAYS)).isoformat()
    out = {}
    async for d in db.mk_ses_checks.find({"email_norm": {"$in": norms}, "at": {"$gte": fresh}}, {"_id": 0}):
        if d.get("no_verdict"):
            continue  # asked, but SES gave no verdict: not "checked" (see _recent_no_verdict)
        # Re-derive from the stored SES levels so a rule fix (e.g. "mailbox
        # does not exist" -> invalid) applies to verdicts already paid for.
        d["status"] = core.ses_status(d.get("overall", ""), d.get("mailbox", "")) or d.get("status")
        out[d["email_norm"]] = d
    return out


async def _recent_no_verdict(norms: list) -> set:
    """Addresses SES answered WITHOUT a verdict in the last
    NO_VERDICT_RETRY_DAYS: already paid for, so not asked (or billed) again yet."""
    retry = (_now() - timedelta(days=NO_VERDICT_RETRY_DAYS)).isoformat()
    return {d["email_norm"] async for d in db.mk_ses_checks.find(
        {"email_norm": {"$in": norms}, "no_verdict": True, "at": {"$gte": retry}}, {"_id": 0, "email_norm": 1})}


async def ses_check(emails: list, *, refresh: bool = False) -> dict:
    """Live SES Email Validation for addresses with no fresh verdict yet.
    Paid, so: only when switched on, only within the monthly cap, never twice
    for one mailbox within SES_CHECK_TTL_DAYS. Returns {email_norm: verdict};
    an address that could not be checked is simply absent (and stays
    'waiting' — campaigns don't mail it)."""
    st = await get_settings()
    norms = {}
    for e in emails:
        n = core.normalise_email(e)
        if n:
            norms.setdefault(n, e)
    # refresh=True: ask SES again even if we have a verdict (used when a
    # campaign recovers from a bounce spike and old verdicts are suspect).
    out = {} if refresh else await _ses_cached(list(norms))
    if not st.get("ses_validation") or not ses_validation_supported():
        return out
    cap = _cap(st)
    skip = set() if refresh else await _recent_no_verdict(list(norms))
    # Providers where SES (learned) can't confirm mailboxes: don't pay; record
    # the answer it would have given — unconfirmed — so the address is still
    # decided by the risky rules instead of being "not checked".
    nope = await _unverifiable({n.rsplit("@", 1)[-1] for n in norms if n not in out})
    for n in list(norms):
        if n not in out and n.rsplit("@", 1)[-1] in nope:
            out[n] = {"email_norm": n, "status": "risky", "overall": "MEDIUM", "mailbox": "MEDIUM", "learned": True,
                      "reasons": ["SES can't confirm mailboxes at this provider — not paid for"]}
    todo = [(n, e) for n, e in norms.items() if n not in out and n not in skip]
    sem = asyncio.Semaphore(4)   # gentle on the API; throttling is retried below
    stop = {"why": ""}
    client = _ses() if todo else None  # boto3 clients are thread-safe; one per batch

    async def one(n: str, e: str) -> None:
        async with sem:
            if stop["why"]:
                return
            if not await _reserve_check(cap):
                stop["why"] = f"this month's SES check budget (₹{st.get('ses_budget_inr')}, {cap} checks) is used up"
                return
            resp = None
            for attempt in range(3):
                try:
                    resp = await asyncio.to_thread(lambda: client.get_email_address_insights(EmailAddress=e))
                    break
                except Exception as ex:  # noqa: BLE001
                    code = ((getattr(ex, "response", None) or {}).get("Error") or {}).get("Code", "")
                    if code in ("ThrottlingException", "TooManyRequestsException", "Throttling") and attempt < 2:
                        await asyncio.sleep(1.5 * (attempt + 1))
                        continue
                    await _release_check()   # failed calls aren't billed; give the slot back
                    if code in ("AccessDeniedException", "AccessDenied", "UnrecognizedClientException",
                                "InvalidClientTokenId", "SignatureDoesNotMatch"):
                        stop["why"] = f"SES refused the check ({code}) — is ses:GetEmailAddressInsights allowed?"
                    else:
                        log.warning("marketing: SES validation failed for one address: %s", str(ex)[:200])
                    return
            if resp is None:
                await _release_check()
                return
            status, reasons = core.ses_verdict(resp)
            if not status:
                # Billed but no verdict. Remember that we asked (status
                # "unknown") so the same address is not paid for again on every
                # re-verify / campaign / sweep — until Oct 2026 nothing was
                # stored and each pass re-billed it. Retried after
                # NO_VERDICT_RETRY_DAYS; never mailed meanwhile (unknown = not sent).
                doc = {"email_norm": n, "status": "unknown", "no_verdict": True, "reasons": ["SES gave no verdict"],
                       "at": _iso(), **core.ses_details(resp)}
                await db.mk_ses_checks.update_one({"email_norm": n}, {"$set": doc}, upsert=True)
                return
            doc = {"email_norm": n, "status": status, "reasons": reasons, "at": _iso(), **core.ses_details(resp)}
            await db.mk_ses_checks.update_one({"email_norm": n}, {"$set": doc}, upsert=True)
            out[n] = doc

    await asyncio.gather(*(one(n, e) for n, e in todo))
    if stop["why"]:
        _SES_STOP.update(why=stop["why"], at=_iso())
        log.warning("marketing: SES validation stopped early: %s", stop["why"])
        await raise_alert("ses_checks", f"Mailbox checks stopped: {stop['why']}. Unchecked addresses are skipped, not mailed.",
                          key=f"ses_checks:{_month()}")
    return out


_MODEL: dict = {"at": 0.0, "doc": None}


async def scoring_model() -> dict:
    """What the self-learning job last worked out (signal adjustments and the
    'valid' threshold). Cached for 10 minutes; the job clears the cache."""
    if _MODEL["doc"] is None or time.time() - _MODEL["at"] > 600:
        _MODEL.update(at=time.time(), doc=await db.mk_model.find_one({"key": "scoring"}, {"_id": 0}) or {})
    return _MODEL["doc"]


async def _scoring_params(st: dict) -> tuple:
    model = await scoring_model()
    auto = bool(st.get("auto_learn", True))
    valid = int(st.get("score_valid") or core.DEFAULT_THRESHOLDS["valid"])
    if auto and (model.get("thresholds") or {}).get("valid"):
        valid = int(model["thresholds"]["valid"])
    th = {"valid": valid, "invalid": min(int(st.get("score_invalid") or core.DEFAULT_THRESHOLDS["invalid"]), valid - 5)}
    return th, (model.get("adjust") or {}) if auto else {}


async def _probe_catch_all(domain: str) -> Optional[bool]:
    """Does this company domain accept ANY address? Ask SES about one made-up
    mailbox there — one paid check per domain per 90 days. If SES says even
    that "exists", the domain is catch-all and "mailbox not confirmed" on a
    real address there means nothing; if SES says it doesn't exist, the
    domain checks its mailboxes and an unconfirmed real address is a doubt."""
    probe = f"zq{secrets.token_hex(5)}x@{domain}"
    norm = core.normalise_email(probe)
    v = (await ses_check([probe])).get(norm)
    await db.mk_ses_checks.delete_one({"email_norm": norm})  # not a real contact
    if not v:
        return None
    mb = v.get("mailbox", "")
    res = False if mb == "LOW" else True if mb in ("HIGH", "MEDIUM") else None
    if res is not None:
        await db.mk_domains.update_one({"domain": domain}, {"$set": {"domain": domain, "catch_all": res, "catch_all_at": _iso()}}, upsert=True)
    return res


async def verify_emails(raw: list, *, autofix: bool = True, ses: str = "cache", names: Optional[list] = None,
                        ignore_erased: bool = False) -> list:
    """Full verification for a batch of raw strings. One result per input row:
    {input, email, email_norm, status, reasons[], fixed_from, ses_checked,
     score, flags[], risk_kind}.

    Hard facts decide first (bad format, no mail server, throwaway, bounced /
    unsubscribed before, SES "mailbox does not exist", proven). Everything
    else gets a 0-100 confidence score (core.score_address) and the score
    decides valid / risky / invalid.

    ses="cache" applies SES verdicts we already paid for (free);
    ses="live" also asks SES about addresses without one (paid, capped) and
    tests company domains for catch-all once; ses="refresh" re-asks SES;
    ses="off" uses only our own free checks.
    names: optional, aligned with raw — lets "name matches address" count."""
    st = await get_settings()
    extra_disp = {d.strip().lower() for d in st.get("extra_disposable") or [] if d.strip()} | await disposable_domains()
    names = list(names or [])
    rows = []
    for i, r in enumerate(raw):
        s = core.check_syntax(r)
        fixed_from = None
        if s.get("cleaned") and s["ok"]:
            # Spreadsheet damage repaired (spaces, mailto:, Name <…>, …): shown
            # as fixed, with exactly what changed.
            fixed_from = str(r)
            s = {**s, "reasons": [f"fixed: {'; '.join(s['cleaned'])}"] + s["reasons"]}
        if autofix and s["ok"] and s["suggestion"]:
            fixed_from = fixed_from or s["email"]
            prev = s["email"]
            s = {**core.check_syntax(s["suggestion"]), "reasons": [x for x in s["reasons"] if x.startswith("fixed:")] + [f"fixed typo from {prev}"]}
        if s["ok"] and s["email"].split("@")[1] in extra_disp:
            s.update(ok=False, disposable=True, reasons=s["reasons"] + ["throwaway domain"])
        rows.append({"input": r, "syntax": s, "fixed_from": fixed_from, "name": str(names[i] if i < len(names) else "") or ""})
    domains = {x["syntax"]["email"].split("@")[1] for x in rows if x["syntax"]["ok"]}
    mx = await domains_mx(domains)
    if autofix:
        await _repair_domains(rows, mx)
    slow = {d for d in domains if mx.get(d) is None}
    if slow:
        # One immediate retry: a slow DNS answer is not a risk signal. Still
        # nothing -> 'unknown', re-checked automatically later (never 'risky').
        mx.update({d: v for d, v in (await domains_mx(slow)).items() if v is not None})
    dom_info = {d["domain"]: d async for d in db.mk_domains.find(
        {"domain": {"$in": list(domains)}}, {"_id": 0, "domain": 1, "bad": 1, "delivered_n": 1, "hard_bounced": 1,
                                              "catch_all": 1, "catch_all_at": 1, "bad_provider": 1})}
    norms = [core.normalise_email(x["syntax"]["email"]) for x in rows]
    proven = await _proven([n for n in norms if n])
    supp, idle = {}, {}
    async for c in db.mk_contacts.find({"email_norm": {"$in": norms}}, {"_id": 0, "email_norm": 1, "suppressed": 1, "since_engaged": 1}):
        if c.get("suppressed"):
            supp[c["email_norm"]] = c["suppressed"]
        idle[c["email_norm"]] = int(c.get("since_engaged") or 0)
    if not ignore_erased:
        # Erased people are blocked from re-import — show it here, so "Check
        # the file" and the import agree (they used to look valid, then
        # vanish silently on import).
        for n in await _erased([x for x in norms if x]):
            supp.setdefault(n, "erased at a privacy request")
        for n, why in (await _removed([x for x in norms if x])).items():
            supp.setdefault(n, f"removed in a clean-up ({why})")
    out = []
    for x, n in zip(rows, norms):
        s = x["syntax"]
        dom_mx = mx.get(s["email"].split("@")[1]) if s["ok"] else None
        status = core.verdict(s, dom_mx, proven=n in proven, suppressed=supp.get(n))
        reasons = list(s["reasons"])
        if s["ok"] and dom_mx is False:
            reasons.append("domain does not receive email")
        if s["ok"] and dom_mx is None and status == "risky":
            status = "unknown"
            reasons.append("domain check did not finish — retried automatically")
        if supp.get(n):
            reasons.append(f"previously {supp[n]}")
        if n in proven:
            reasons.append("proven (confirmed account or delivered before)")
        out.append({"input": x["input"], "email": s["email"], "email_norm": n, "status": status,
                    "reasons": reasons, "fixed_from": x["fixed_from"], "role": s.get("role", False), "ses_checked": False,
                    "_ses": None, "_name": x["name"],
                    # A guessed domain repair is never trusted above "risky".
                    "_suggest": (bool(s.get("suggestion")) and not x["fixed_from"]) or bool(s.get("domain_guess"))})
    if ses in ("cache", "live", "refresh"):
        # Only addresses our free checks couldn't settle are worth an SES look:
        # verified ones are proven, invalid/suppressed ones are never sent.
        need = [r for r in out if r["status"] in ("valid", "risky") and r["email_norm"]]
        if ses == "cache":
            verdicts = await _ses_cached([r["email_norm"] for r in need])
        else:
            verdicts = await ses_check([r["email"] for r in need], refresh=ses == "refresh")
        for r in need:
            v = verdicts.get(r["email_norm"])
            if v:
                r["ses_checked"] = True
                r["_ses"] = v
                if v["status"] == "invalid":  # SES LOW / mailbox does not exist: a fact, not a score
                    r["status"] = "invalid"
                    r["reasons"] = r["reasons"] + [x for x in v.get("reasons") or [] if x not in r["reasons"]]
    ca_fresh = (_now() - timedelta(days=90)).isoformat()
    if ses in ("live", "refresh"):
        doms = {r["email"].split("@")[-1] for r in out if r["status"] in ("valid", "risky") and r["_ses"]
                and r["_ses"].get("overall") == "MEDIUM" and r["_ses"].get("mailbox") in ("MEDIUM", "")}
        for d in sorted(doms):
            if d in core.WEBMAIL or ((dom_info.get(d) or {}).get("catch_all_at") or "") >= ca_fresh:
                continue
            res = await _probe_catch_all(d)
            if res is not None:
                dom_info.setdefault(d, {"domain": d}).update(catch_all=res, catch_all_at=_iso())
    th, adjust = await _scoring_params(st)
    role_ok = bool(st.get("include_role_addresses"))
    for r in out:
        v = r.pop("_ses") or {}
        nm, suggest = r.pop("_name"), r.pop("_suggest")
        r["ses_overall"] = v.get("overall", "")
        r["role"] = bool(r["role"] or v.get("role"))
        r["risk_kind"], r["flags"] = "", []
        r["score"] = 100 if r["status"] == "verified" else 0 if r["status"] in ("invalid", "suppressed") else None
        if r["status"] not in ("valid", "risky"):
            continue
        dom = r["email"].split("@")[-1]
        di = dom_info.get(dom) or {}
        catch_all = di.get("catch_all") if (di.get("catch_all_at") or "") >= ca_fresh else None
        sig = {"ses": v.get("overall", ""), "mailbox": v.get("mailbox", ""), "catch_all": catch_all,
               "random": bool(v.get("random") or "SES: looks randomly typed" in (v.get("reasons") or [])),
               "role": r["role"], "domain_bad": bool(di.get("bad")),
               "domain_good": dom not in core.WEBMAIL and int(di.get("delivered_n") or 0) >= 3 and not di.get("hard_bounced"),
               "name_match": core.name_matches(r["email"].split("@")[0], nm),
               "made_up": bool(core.looks_made_up(r["email"].split("@")[0])),
               "inactive": core.sunset_state(idle.get(r["email_norm"], 0), inactive_after=int(st.get("inactive_after") or 0),
                                             sunset_after=0) != "active"}
        sc = core.score_address(sig, adjust=adjust, thresholds=th, role_ok=role_ok)
        status = sc["status"]
        if suggest and status == "valid":
            status = "risky"  # looks like a typo we weren't allowed to fix
        r.update(status=status, score=sc["score"], flags=sc["flags"])
        r["reasons"] = [x for x in r["reasons"] if not x.startswith("SES:")] + [f"confidence {sc['score']}/100"] + sc["reasons"]
        if sig["made_up"]:
            r["reasons"].append(f"looks made-up: {core.looks_made_up(r['email'].split('@')[0])}")
        if r["role"] and role_ok:
            r["reasons"].append("role address — treated as normal (Settings)")
        bad_prov = bool(di.get("bad_provider")) and v.get("overall") != "HIGH"
        if bad_prov:
            # Learned from our own bounces (learn_providers): unconfirmed
            # addresses at this provider bounce too often to mail. An address
            # SES itself confirmed (HIGH) is not "unconfirmed" and still goes.
            status = "risky"
            r.update(status=status)
            r["reasons"].append(f"{dom}: unconfirmed addresses here bounced too often — skipped until proven")
        if status == "risky":
            if bad_prov:
                r["risk_kind"] = "bad_provider"
            elif r["role"] and not role_ok:
                r["risk_kind"] = "role"
            elif sig["domain_bad"]:
                r["risk_kind"] = "unconfirmed"
            elif catch_all is True or (catch_all is None and core.risk_kind(role=False, domain=dom, ses_overall=sig["ses"],
                                                                              mailbox=sig["mailbox"]) == "catch_all"):
                r["risk_kind"] = "catch_all"
            else:
                r["risk_kind"] = "unconfirmed"
    return out


def _status_update(v: dict) -> dict:
    """Contact fields from one verify_emails() row."""
    upd = {"email_status": v["status"], "reasons": v["reasons"][:7], "verified_at": _iso(), "risk_kind": v.get("risk_kind", ""),
           "role": bool(v.get("role")), "score": v.get("score"), "flags": v.get("flags") or []}
    if v.get("ses_checked"):
        upd.update(ses_checked=True, ses_overall=v.get("ses_overall", ""))
    return upd


_validating = False


async def validate_pending(max_n: int = 5000) -> dict:
    """Background pass over saved contacts: website contacts that arrived
    without any check ('unknown') get the free checks, and — with SES
    validation on — every not-yet-proven address gets its SES mailbox check
    (within the budget). Campaigns don't depend on it — sending checks its own
    recipients first — it just spreads the work and keeps the Contacts page current."""
    global _validating
    if _validating:
        return {"running": True}
    _validating = True
    started = _iso()
    done = 0
    try:
        st = await get_settings()
        use = await ses_usage()
        if st.get("ses_validation") and use["supported"] and use["used"] < use["cap"]:
            flt = {"email_status": {"$in": ["unknown", "valid", "risky"]}, "ses_checked": {"$ne": True}}
        else:
            flt = {"email_status": "unknown"}
        contacts = await db.mk_contacts.find(flt, {"_id": 0, "id": 1, "email": 1, "name": 1}).limit(max_n).to_list(max_n)
        for i in range(0, len(contacts), 100):
            chunk = contacts[i:i + 100]
            res = await verify_emails([c["email"] for c in chunk], autofix=False, ses="live", names=[c.get("name", "") for c in chunk])
            await dbutil.bulk_set(db.mk_contacts, [(c["id"], _status_update(v)) for c, v in zip(chunk, res)])
            done += len(chunk)
            if _SES_STOP.get("at", "") >= started:
                break  # cap reached or SES refused: the rest wait for next time
    except Exception:  # noqa: BLE001
        log.exception("marketing: validate_pending failed")
    finally:
        _validating = False
    return {"processed": done}


async def _reapply_role_rule() -> None:
    """The role-address setting changed: re-rate saved role addresses now
    (cached SES verdicts only — free), so Contacts and audience counts match
    the new rule straight away instead of at the next send."""
    try:
        flt = {"$or": [{"role": True}, {"risk_kind": "role"}, {"reasons": "role address — treated as normal (Settings)"}],
               "email_status": {"$in": ["valid", "risky"]}}
        async for chunk in dbutil.batches(db.mk_contacts, flt, {"_id": 0, "id": 1, "email": 1, "name": 1}, size=200):
            res = await verify_emails([c["email"] for c in chunk], autofix=False, ses="cache", names=[c.get("name", "") for c in chunk])
            await dbutil.bulk_set(db.mk_contacts, [(c["id"], _status_update(v)) for c, v in zip(chunk, res)])
    except Exception:  # noqa: BLE001
        log.exception("marketing: re-applying the role-address rule failed")


# ------------------------------------------------------------- dispatch ---
# Every piece of background work goes through dispatch(). WORKER_MODE=inline
# (default) runs it as a task in this process, exactly as before. With
# WORKER_MODE=external it becomes a job for worker.py, so sending, checking and
# learning never compete with shoppers for this process's CPU and memory — and
# if the worker's heartbeat has stopped, the work runs here anyway (with an
# alert), so nothing waits for a worker that isn't there.
async def _dispatch_async(kind: str, payload: dict, dedupe: Optional[str]) -> None:
    try:
        if _jobs.worker_mode() == "external":
            if await _jobs.worker_alive(db):
                await _ensure_indexes()
                await _jobs.enqueue(db, kind, payload, dedupe=dedupe)
                return
            await raise_alert("worker", "The background worker is not running, so marketing work is running in the "
                                        "website process. Check the worker service on Render.", key="worker_down")
        fn = _jobs._HANDLERS.get(kind)
        if fn:
            await fn(payload)
    except Exception:  # noqa: BLE001
        log.exception("marketing: dispatch %s failed", kind)


def dispatch(kind: str, payload: Optional[dict] = None, *, dedupe: Optional[str] = None) -> None:
    _bg(_dispatch_async(kind, payload or {}, dedupe))


async def job_dead_alert(job: dict, error: str) -> None:
    first = error.splitlines()[0][:150]
    if first.startswith("no worker could run"):
        msg = f"Background task '{job.get('kind')}' could not run: {first}"
    else:
        msg = f"Background task '{job.get('kind')}' failed 3 times and was stopped: {first}"
    await raise_alert("job", msg, key=f"job:{job.get('kind')}")


def kick_validation() -> None:
    """Check waiting addresses soon (after an import / sync / re-verify / sign-up)."""
    dispatch("mk_validate", dedupe="all")


# ------------------------------------------------------------- contacts ---
def _consent(status: str, source: str, by: str = "") -> dict:
    return {"status": status, "source": source[:120], "at": _iso(), "by": by}


async def upsert_contact(email: str, *, name: str = "", phone: str = "", source: str = "",
                         email_consent: Optional[str] = None, wa_consent: Optional[str] = None,
                         lists: Optional[list] = None, status: Optional[str] = None, reasons: Optional[list] = None,
                         proven: Optional[bool] = None, customer: Optional[bool] = None, by: str = "",
                         ses_checked: Optional[bool] = None, extra: Optional[dict] = None,
                         readd_erased: bool = False) -> Optional[dict]:
    """Create or update one contact (keyed on the real mailbox). Consent is
    only ever RAISED to subscribed by a new opt-in, never by an import of an
    already-unsubscribed contact: an unsubscribe is final."""
    await _ensure_indexes()
    s = core.check_syntax(email)
    if not s["email"]:
        return None
    norm = core.normalise_email(s["email"])
    if await db.mk_erased.find_one({"h": _erase_hash(norm)}, {"_id": 1}):
        if not readd_erased:
            return None  # erased on request: an old export must not bring them back
        # The person opted in again themselves (or a superadmin confirmed they
        # asked to be re-added): lift the block.
        await db.mk_erased.delete_one({"h": _erase_hash(norm)})
    if await db.mk_removed.find_one({"h": _erase_hash(norm)}, {"_id": 1}):
        if not readd_erased:
            # Removed in a clean-up (bounced / invalid / …): a sync of old
            # orders or an old file must not bring the address back. Their
            # own new sign-up (double opt-in) or a superadmin re-add does.
            return None
        await db.mk_removed.delete_one({"h": _erase_hash(norm)})
    cur = await db.mk_contacts.find_one({"email_norm": norm}, {"_id": 0})
    now = _iso()
    if not cur:
        cur = {"id": str(uuid.uuid4()), "email": s["email"], "email_norm": norm, "name": "", "phone": "", "lists": [],
               "source": source, "created_at": now, "email_status": "unknown", "reasons": [], "suppressed": "",
               "email_consent": _consent("none", source), "wa_consent": _consent("none", source),
               "stats": {"sent": 0, "opened": 0, "clicked": 0}}
        await db.mk_contacts.insert_one(dict(cur))
    upd: dict = {"updated_at": now}
    if name and not cur.get("name"):
        upd["name"] = name.strip()[:120]
    if phone and not cur.get("phone"):
        upd["phone"] = re.sub(r"[^\d+ ]", "", phone)[:20]
    if status:
        upd["email_status"] = status
        upd["reasons"] = (reasons or [])[:6]
        upd["verified_at"] = now
    if proven:
        upd["proven"] = True
    if ses_checked:
        upd["ses_checked"] = True
    if extra:
        upd.update({k: v for k, v in extra.items() if k in ("risk_kind", "ses_overall", "score", "flags", "role")})
    if customer:
        upd["customer"] = True
    final = {"unsubscribed", "complained"}
    if email_consent and cur.get("email_consent", {}).get("status") not in final:
        if not (email_consent == "pending" and cur.get("email_consent", {}).get("status") == "subscribed"):
            upd["email_consent"] = _consent(email_consent, source, by)
    if wa_consent and cur.get("wa_consent", {}).get("status") not in final:
        upd["wa_consent"] = _consent(wa_consent, source, by)
    ops: dict = {"$set": upd}
    if lists:
        ops["$addToSet"] = {"lists": {"$each": lists}}
    await db.mk_contacts.update_one({"id": cur["id"]}, ops)
    return {**cur, **upd}


async def suppress(contact_filter: dict, reason: str) -> None:
    """Never email again: bounced, complained (also unsubscribes)."""
    upd = {"suppressed": reason, "suppressed_at": _iso(), "email_status": "suppressed"}
    if reason in ("complained", "unsubscribed"):
        upd["email_consent"] = _consent("complained" if reason == "complained" else "unsubscribed", reason)
    await db.mk_contacts.update_one(contact_filter, {"$set": upd})


# hooks ------------------------------------------------------------------
async def on_newsletter(email: str, source: str = "newsletter") -> None:
    """Website sign-up. Queued, so the visitor's request returns at once — the
    address check (which may ask SES) and the confirmation email happen in the
    background."""
    dispatch("mk_newsletter", {"email": email, "source": source})


def is_waitlist_source(source: str) -> bool:
    """Coming-soon / product waitlists post to the same endpoint as the
    newsletter, but their purpose is narrower: "tell me when X launches".
    Under the DPDP Act consent is purpose-specific, so a waitlist sign-up is
    NOT a newsletter / marketing opt-in."""
    return "waitlist" in (source or "").lower()


async def _newsletter_signup(email: str, source: str = "newsletter") -> None:
    """Website sign-up. With double opt-in on, the contact waits as 'pending'
    until they click the confirmation link we email them. Waitlist sign-ups are
    recorded as contacts with NO marketing consent (see is_waitlist_source)."""
    try:
        if is_waitlist_source(source):
            await upsert_contact(email, source=source)
            return
        st = await get_settings()
        if st["double_opt_in"]:
            # Check the address the moment it arrives: a bot or a typo'd
            # sign-up would otherwise cost us a bounce on the confirmation email.
            v = (await verify_emails([email], autofix=False, ses="live", ignore_erased=True))[0]
            c = await upsert_contact(email, source=source or "newsletter", email_consent="pending", status=v["status"], readd_erased=True,
                                     reasons=v["reasons"], ses_checked=v.get("ses_checked") or None,
                                     extra={"risk_kind": v.get("risk_kind", ""), "ses_overall": v.get("ses_overall", ""), "score": v.get("score"),
                                    "flags": v.get("flags") or [], "role": bool(v.get("role"))})
            if v["status"] in ("invalid", "suppressed"):
                return  # kept as a record, but no email to a dead / bounced address
            if c and c.get("email_consent", {}).get("status") == "pending":
                await send_confirmation(c)
        else:
            await upsert_contact(email, source=source or "newsletter", email_consent="subscribed", readd_erased=True)
    except Exception:  # noqa: BLE001
        log.exception("marketing: newsletter hook failed for %s", email)


async def on_order_paid(order: dict) -> None:
    """Paid customer -> contact (customer + proven deliverable once their
    receipt went out); consent only if they ticked the box at checkout."""
    try:
        await upsert_contact(order.get("email", ""), name=order.get("full_name", ""), phone=order.get("phone", ""),
                             source="checkout", customer=True,
                             email_consent="subscribed" if order.get("email_marketing_optin") else None,
                             readd_erased=bool(order.get("email_marketing_optin")))  # ticking the box is a fresh opt-in
        if order.get("email_marketing_optin"):
            kick_validation()  # check the new subscriber now, not on the next campaign
    except Exception:  # noqa: BLE001
        log.exception("marketing: order hook failed for %s", order.get("order_number"))


# ------------------------------------------------------------------ SES ---
def _ses():
    import boto3
    return boto3.client("sesv2", region_name=SES_REGION)


def _send_raw(msg_bytes: bytes) -> str:
    kw = {"Content": {"Raw": {"Data": msg_bytes}}}
    if SES_CONFIG_SET:
        kw["ConfigurationSetName"] = SES_CONFIG_SET
    return _ses().send_email(**kw)["MessageId"]


def build_mime(*, st: dict, to: str, subject: str, html_doc: str, text_doc: str, unsub_post: str) -> bytes:
    msg = MIMEMultipart("alternative")
    msg["Subject"] = subject
    msg["From"] = formataddr((st["from_name"], st["from_email"]))
    msg["To"] = to
    if st.get("reply_to"):
        msg["Reply-To"] = st["reply_to"]
    msg["Message-ID"] = make_msgid(domain=st["from_email"].split("@")[-1])
    if unsub_post:
        msg["List-Unsubscribe"] = f"<{unsub_post}>, <mailto:{st['from_email']}?subject=unsubscribe>"
        msg["List-Unsubscribe-Post"] = "List-Unsubscribe=One-Click"
    msg.attach(MIMEText(text_doc, "plain", "utf-8"))
    msg.attach(MIMEText(html_doc, "html", "utf-8"))
    return msg.as_bytes()


def _links(send_id: str):
    sec = _secret()
    return {
        "open": f"{PUBLIC_API}/api/m/o/{send_id}/{core.sign(sec, 'o', send_id)}.gif",
        "unsub_page": f"{SITE}/unsubscribe?s={send_id}&t={core.sign(sec, 'u', send_id)}",
        "unsub_post": f"{PUBLIC_API}/api/m/u/{send_id}/{core.sign(sec, 'u', send_id)}",
        "click": lambda idx: f"{PUBLIC_API}/api/m/c/{send_id}/{idx}/{core.sign(sec, 'c', send_id, str(idx))}",
    }


# ------------------------------------------------- "Picked for you" ---
def _for_you_count(campaign: dict) -> int:
    return max([int(b.get("count") or 3) for b in campaign.get("blocks") or [] if b.get("type") == "for_you"] or [0])


async def release_pool(limit: int = 40) -> list:
    """Latest releases a reader can actually buy today (in stock, or a live
    pre-order), newest first, covers only — the candidates for "Picked for you"."""
    docs = await db.books.find(
        {"enabled": {"$ne": False}, "product_type": {"$nin": ["hamper"]},
         "$or": [{"stock": {"$gt": 0}}, {"coming_soon": True}],
         "cover_image": {"$nin": [None, ""]}},
        {"_id": 0, "id": 1, "title": 1, "author": 1, "subject": 1, "category": 1, "price": 1, "cover_image": 1,
         "release_rank": 1, "publication_year": 1, "new_release": 1}).sort(
        [("new_release", -1), ("release_rank", 1), ("publication_year", -1)]).limit(limit).to_list(limit)
    return docs


async def reader_profile(email_norm: str) -> dict:
    """What one person bought and showed interest in, from OUR records only:
    paid orders, their cart, and "tell me when it's back" requests. No
    browsing history is kept per person, so none is used."""
    prof = {"bought": set(), "interest": set(), "authors": {}, "subjects": {}, "categories": {}}
    if not email_norm:
        return prof
    u = await db.users.find_one({"$or": [{"email_normalised": email_norm}, {"email": email_norm}]}, {"_id": 0, "id": 1, "email": 1})
    or_ = [{"email": {"$regex": f"^{re.escape(email_norm)}$", "$options": "i"}}]
    if u and u.get("id"):
        or_.append({"user_id": u["id"]})
    ids = []
    async for o in db.orders.find({"$or": or_, "payment_status": "paid"}, {"_id": 0, "items": 1}):
        for it in o.get("items") or []:
            if it.get("book_id"):
                prof["bought"].add(it["book_id"])
                ids.append(it["book_id"])
    if u and u.get("id"):
        cart = await db.carts.find_one({"user_id": u["id"]}, {"_id": 0, "items": 1}) or {}
        prof["interest"] |= {it.get("book_id") for it in cart.get("items") or [] if it.get("book_id")}
    async for n in db.stock_notifications.find({"email": {"$regex": f"^{re.escape(email_norm)}$", "$options": "i"}}, {"_id": 0, "book_id": 1}):
        if n.get("book_id"):
            prof["interest"].add(n["book_id"])
    look = list(set(ids) | prof["interest"])
    if look:
        async for b in db.books.find({"id": {"$in": look}}, {"_id": 0, "author": 1, "subject": 1, "category": 1}):
            for key, field in (("authors", "author"), ("subjects", "subject"), ("categories", "category")):
                v = (b.get(field) or "").strip().lower()
                if v:
                    prof[key][v] = prof[key].get(v, 0) + 1
    return prof


async def attach_picks(contact: dict, campaign: dict, pool: Optional[list] = None) -> list:
    """Put this reader's picks on the contact (`_picks`) and return their
    absolute URLs (to be stored as the send's own trackable links)."""
    n = _for_you_count(campaign)
    if not n:
        return []
    pool = pool if pool is not None else await release_pool()
    prof = await reader_profile(contact.get("email_norm") or core.normalise_email(contact.get("email") or ""))
    picks = core.pick_for_you(prof, pool, n)
    A = {k for k, v in prof["authors"].items() if v}
    S = {k for k, v in prof["subjects"].items() if v}
    out = []
    for b in picks:
        why = ("You asked about this one" if b.get("id") in prof["interest"]
               else f"Because you read {b.get('author')}" if (b.get("author") or "").strip().lower() in A
               else f"New in {b.get('subject')}" if (b.get("subject") or "").strip().lower() in S else "")
        out.append({"title": b.get("title"), "author": b.get("author"), "price": f"₹{float(b.get('price') or 0):,.0f}",
                    "cover": _abs_media(b.get("cover_image") or ""), "url": f"{SITE}/books/{b['id']}", "why": why})
    contact["_picks"] = out
    return [p["url"] for p in out]


def _abs_media(u: str) -> str:
    from emailer import media_url
    return media_url(u) or ""


def render_for(campaign: dict, contact: dict, send_id: str, st: dict, links: Optional[list] = None) -> tuple[str, str]:
    """`links`: this send's own link list (campaign links + the reader's
    picks) — click tracking redirects by index into exactly this list."""
    links = links if links is not None else (campaign.get("links") or [])
    L = _links(send_id)

    def link_for(u):
        return L["click"](links.index(u)) if u in links else u

    brand = {"name": "Oakbridge Publishing", "logo": st.get("logo"), "footer": st.get("footer"), "navy": "#002B5C", "red": "#CC0033"}
    return core.render_email(campaign.get("blocks") or [], brand=brand, contact=contact, link_for=link_for,
                             open_pixel=L["open"], unsubscribe_url=L["unsub_page"],
                             preheader=core.personalise(campaign.get("preheader") or "", contact))


async def send_confirmation(contact: dict) -> None:
    """Double opt-in email (not tracked, no unsubscribe needed: one email,
    sent because they asked)."""
    st = await get_settings()
    sig = core.sign(_secret(), "confirm", contact["id"])
    url = f"{SITE}/subscribe/confirm?c={contact['id']}&t={sig}"
    blocks = [{"type": "heading", "text": "Please confirm your subscription"},
              {"type": "text", "text": "You asked to hear about new books, events and offers from Oakbridge Publishing. Click below to confirm. If this wasn't you, ignore this email and nothing happens."},
              {"type": "button", "label": "Yes, subscribe me", "url": url}]
    html_doc, text_doc = core.render_email(blocks, brand={"name": "Oakbridge Publishing", "logo": st.get("logo"), "footer": st.get("footer")},
                                           contact=contact, link_for=lambda u: u)
    raw = build_mime(st=st, to=contact["email"], subject="Confirm your subscription to Oakbridge", html_doc=html_doc,
                     text_doc=text_doc, unsub_post="")
    try:
        await asyncio.to_thread(_send_raw, raw)
        await db.mk_contacts.update_one({"id": contact["id"]}, {"$set": {"confirm_sent_at": _iso()}})
    except Exception:  # noqa: BLE001
        log.exception("marketing: confirmation email failed for %s", contact.get("email"))


# ------------------------------------------------------------- audience ---
async def _rule_emails(rule: dict) -> Optional[set]:
    """email_norm set for rules that look outside mk_contacts (orders, carts)."""
    t = rule.get("type")
    if t == "bought_category":
        since = (_now() - timedelta(days=int(rule.get("days") or 180))).isoformat()
        cat = str(rule.get("category") or "")
        ids = [b["id"] async for b in db.books.find({"category": cat}, {"_id": 0, "id": 1})]
        out = set()
        async for o in db.orders.find({"payment_status": "paid", "created_at": {"$gte": since}, "items.book_id": {"$in": ids}},
                                      {"_id": 0, "email": 1}):
            out.add(core.normalise_email(o.get("email")))
        return out
    if t == "customers":
        since = (_now() - timedelta(days=int(rule.get("days") or 3650))).isoformat()
        return {core.normalise_email(o.get("email")) async for o in db.orders.find(
            {"payment_status": "paid", "created_at": {"$gte": since}}, {"_id": 0, "email": 1})}
    if t == "abandoned_cart":
        uids = [c["user_id"] async for c in db.carts.find({"items.0": {"$exists": True}}, {"_id": 0, "user_id": 1})]
        return {core.normalise_email(u.get("email")) async for u in db.users.find({"id": {"$in": uids}}, {"_id": 0, "email": 1})}
    return None


async def audience_filter(campaign: dict) -> dict:
    """Mongo filter for who a campaign MAY go to. Consent and suppression are
    applied here, always — no option skips them. For email, which of these
    candidates are actually sent is decided when the campaign is prepared
    (mailbox checks + risk rules, core.send_decision)."""
    aud = campaign.get("audience") or {}
    flt: dict = {}
    ors = []
    if aud.get("list_ids"):
        ors.append({"lists": {"$in": aud["list_ids"]}})
    rules = aud.get("rules") or []
    and_parts = []
    for r in rules:
        if r.get("type") == "source":
            and_parts.append({"source": r.get("value")})
        elif r.get("type") == "status":
            and_parts.append({"email_status": r.get("value")})
        else:
            emails = await _rule_emails(r)
            if emails is not None:
                and_parts.append({"email_norm": {"$in": list(emails)}})
    if ors:
        flt["$or"] = ors
    if and_parts:
        flt["$and"] = and_parts
    if aud.get("exclude_list_ids"):
        flt["lists"] = {"$nin": aud["exclude_list_ids"]}
    if campaign.get("channel") == "whatsapp":
        flt["wa_consent.status"] = "subscribed"
        flt["phone"] = {"$nin": [None, ""]}
    elif campaign.get("kind") == "reconfirm":
        # A re-confirm / permission email also reaches people imported as "not
        # known to us" (permission_pending): it is the ONE email they can get,
        # and a click is their opt-in. Never anyone who said no.
        flt.setdefault("$and", []).append({"$or": [
            {"email_consent.status": "subscribed"},
            {"permission_pending": True, "email_consent.status": {"$nin": ["unsubscribed", "complained", "lapsed"]}}]})
        flt["suppressed"] = {"$in": [None, ""]}
        flt["email_status"] = {"$in": sorted(core.selected_statuses(aud)) + ["unknown"]}
    else:
        flt["email_consent.status"] = "subscribed"
        flt["suppressed"] = {"$in": [None, ""]}
        # Only the statuses ticked for this campaign (+ not-yet-checked, which
        # are checked on Send and then filtered by the same ticks).
        flt["email_status"] = {"$in": sorted(core.selected_statuses(aud)) + ["unknown"]}
    return flt


def _policy(st: dict, campaign: dict) -> dict:
    pol = core.clean_policy(st.get("risk_policy"))
    if (campaign.get("audience") or {}).get("include_risky") and pol["unconfirmed"] == "skip":
        pol["unconfirmed"] = "tail"  # older campaigns that ticked "include risky"
    return pol


# -------------------------------------------------------------- sending ---
WORKER_STALE = 120  # seconds without a heartbeat before another tick takes over
TAIL_BATCH_PAUSE = 60  # seconds between tail batches of 25, so bounces land before the next (< WORKER_STALE)
_running: set = set()


_preparing: set = set()
PREPARE_STALE = 180  # seconds without progress before the cron restarts a preparation


def start_prepare(cid: str) -> None:
    if cid not in _preparing:
        dispatch("mk_prepare", {"cid": cid}, dedupe=cid)


async def _prepare_campaign(cid: str) -> None:
    """Send = check, then send. Every not-yet-proven recipient gets the free
    checks and (if on) its SES mailbox check — within the budget — then each
    recipient is queued, put in the tail, or skipped WITH the reason, by the
    risk rules in Settings. Idempotent: a restart re-runs it safely (cached
    verdicts cost nothing, sends are unique per contact)."""
    _preparing.add(cid)
    try:
        await heal_proven()
        await learn_providers()
        c = await db.mk_campaigns.find_one({"id": cid}, {"_id": 0})
        if not c or c.get("status") != "preparing":
            return
        st = await get_settings()
        flt = await audience_filter(c)
        proj = {"_id": 0, "id": 1, "email": 1, "phone": 1, "email_status": 1, "risk_kind": 1, "email_norm": 1, "name": 1,
                "score": 1, "flags": 1, "since_engaged": 1}
        is_email = c["channel"] == "email"
        ses_on = is_email and bool(st.get("ses_validation")) and ses_validation_supported()
        pol = _policy(st, c)
        allowed = core.selected_statuses(c.get("audience"))
        reconfirm = c.get("kind") == "reconfirm"
        sunset_after = int(st.get("sunset_after") or 0)
        to_check = await db.mk_contacts.count_documents({"$and": [flt, {"email_status": {"$ne": "verified"}}]}) if is_email else 0
        checked, skipped = 0, {}
        # Pages of 500 (dbutil.batches): memory stays flat however big the
        # audience; each page is checked, decided and queued before the next.
        async for page in dbutil.batches(db.mk_contacts, flt, proj, size=500):
            if is_email:
                todo = [x for x in page if x.get("email_status") != "verified"]
                for i in range(0, len(todo), 100):
                    chunk = todo[i:i + 100]
                    res = await verify_emails([x["email"] for x in chunk], autofix=False, ses="live" if ses_on else "cache",
                                              names=[x.get("name", "") for x in chunk])
                    await dbutil.bulk_set(db.mk_contacts, [(x["id"], _status_update(v)) for x, v in zip(chunk, res)])
                    for x, v in zip(chunk, res):
                        x.update(email_status=v["status"], risk_kind=v.get("risk_kind", ""), _ses=v.get("ses_checked"),
                                 score=v.get("score"), flags=v.get("flags") or [])
                checked += len(todo)
            now = _iso()
            docs = []
            for x in page:
                if is_email:
                    act, prio, why = core.send_decision(x.get("email_status") or "unknown", x.get("risk_kind") or "", pol,
                                                        needs_ses=ses_on and x.get("email_status") != "verified",
                                                        has_ses=bool(x.get("_ses")), allowed=allowed)
                    n_idle = int(x.get("since_engaged") or 0)
                    if act != "skip" and not reconfirm and core.sunset_state(
                            n_idle, inactive_after=0, sunset_after=sunset_after) == "sunset":
                        # Retired: no opens / clicks in N campaign emails. A
                        # re-confirm campaign is the one way back in.
                        act, prio, why = "skip", 9, f"inactive — no opens or clicks in {n_idle} emails"
                else:
                    act, prio, why = "send", 1, ""
                doc = {"id": str(uuid.uuid4()), "campaign_id": cid, "contact_id": x["id"], "priority": prio, "queued_at": now,
                       "to": x.get("email") if is_email else x.get("phone")}
                if is_email:
                    # What we believed about the address when we sent: the
                    # self-learning job compares this with what actually bounced.
                    doc.update(band=x.get("email_status") or "unknown",
                               score=100 if x.get("email_status") == "verified" else x.get("score"),
                               flags=["proven"] if x.get("email_status") == "verified" else (x.get("flags") or []))
                if act == "skip":
                    doc.update(status="skipped", error=why)
                    skipped[why] = skipped.get(why, 0) + 1
                else:
                    doc.update(status="queued", tail=act == "tail")
                docs.append(doc)
            cur = await db.mk_campaigns.find_one_and_update(
                {"id": cid}, {"$set": {"prepare.checked": checked, "prepare.to_check": to_check, "prepare_heartbeat": _iso()}},
                projection={"_id": 0, "status": 1})
            if not cur or cur.get("status") != "preparing":
                # Cancelled while checking: nothing queued so far may go out.
                await db.mk_sends.update_many({"campaign_id": cid, "status": "queued"},
                                              {"$set": {"status": "skipped", "error": "cancelled"}})
                return
            try:
                await db.mk_sends.insert_many(docs, ordered=False)
            except Exception:  # noqa: BLE001 — duplicates from a restarted preparation are fine (unique index)
                pass
        queued = await db.mk_sends.count_documents({"campaign_id": cid, "status": "queued"})
        when = c.get("schedule_at") or ""
        now = _iso()
        if not queued:
            status, extra = "sent", {"finished_at": now, "note": "Nobody passed the checks — see the skipped list."}
        elif when and when > now:
            status, extra = "scheduled", {}
        else:
            status, extra = "sending", {"started_at": now, "schedule_at": now}
        await db.mk_campaigns.update_one({"id": cid, "status": "preparing"}, {"$set": {
            "status": status, "queued": queued, "prepare.finished_at": now, "prepare.skipped": skipped, **extra}})
        await refresh_campaign_stats(cid)
        if status == "sending":
            await start_campaign(cid)
    except Exception as e:  # noqa: BLE001
        log.exception("marketing: preparing campaign %s failed", cid)
        await db.mk_campaigns.update_one({"id": cid, "status": "preparing"},
                                         {"$set": {"status": "draft", "prepare_error": str(e)[:300]}})
        await raise_alert("prepare", f"A campaign could not be prepared and went back to draft: {str(e)[:160]}", campaign_id=cid)
    finally:
        _preparing.discard(cid)


async def _window_stats(cid: str, since: str) -> dict:
    """sent / bounced / complained for this campaign's sends since a moment —
    after a recovery the bounce guard judges only what was sent since."""
    base = {"campaign_id": cid, "sent_at": {"$gte": since}}
    return {"sent": await db.mk_sends.count_documents({**base, "status": {"$in": ["sent", "bounced", "complained"]}}),
            "bounced": await db.mk_sends.count_documents({**base, "status": "bounced"}),
            "complained": await db.mk_sends.count_documents({**base, "status": "complained"})}


async def _recover(cid: str, why: str) -> dict:
    """First bounce spike: instead of stopping, keep only proven addresses
    and ones SES rated HIGH (re-asking SES where that verdict is over 30
    days old); drop the rest, with the reason, and carry on."""
    fresh_cut = (_now() - timedelta(days=30)).isoformat()
    dropped, recheck = 0, []
    async for queued in dbutil.batches(db.mk_sends, {"campaign_id": cid, "status": "queued"}, {"_id": 0, "id": 1, "contact_id": 1}, size=500):
        drop: list = []
        await _recover_page(queued, fresh_cut, drop, recheck)
        dropped += await _drop_sends(drop)   # per page: memory stays flat
    return await _recover_finish(cid, why, dropped, recheck)


async def _drop_sends(ids: list) -> int:
    if not ids:
        return 0
    r = await db.mk_sends.update_many({"id": {"$in": ids}, "status": "queued"},
                                      {"$set": {"status": "skipped", "error": "dropped after a bounce spike (not proven / not high-confidence)"}})
    return r.modified_count


async def _recover_page(queued: list, fresh_cut: str, drop: list, recheck: list) -> None:
    contacts = {x["id"]: x async for x in db.mk_contacts.find({"id": {"$in": [q["contact_id"] for q in queued]}},
                                                              {"_id": 0, "id": 1, "email": 1, "email_norm": 1, "email_status": 1, "name": 1})}
    cached = await _ses_cached([x["email_norm"] for x in contacts.values() if x.get("email_status") != "verified"])
    for q in queued:
        x = contacts.get(q["contact_id"])
        if not x:
            drop.append(q["id"])
        elif x.get("email_status") == "verified":
            continue
        elif (cached.get(x["email_norm"]) or {}).get("overall") != "HIGH":
            drop.append(q["id"])
        elif cached[x["email_norm"]].get("at", "") < fresh_cut:
            recheck.append((q, x))


async def _recover_finish(cid: str, why: str, dropped: int, recheck: list) -> dict:
    for i in range(0, len(recheck), 100):
        part = recheck[i:i + 100]
        res = await verify_emails([x["email"] for _, x in part], autofix=False, ses="refresh", names=[x.get("name", "") for _, x in part])
        await dbutil.bulk_set(db.mk_contacts, [(x["id"], _status_update(v)) for (_, x), v in zip(part, res)])
        dropped += await _drop_sends([q["id"] for (q, _), v in zip(part, res)
                                      if not (v["status"] == "valid" and v.get("ses_overall") == "HIGH")])
    info = {"at": _iso(), "reason": why, "dropped": dropped, "rechecked": len(recheck)}
    await db.mk_campaigns.update_one({"id": cid}, {"$set": {"recovered_at": info["at"], "recovery": info}})
    log.warning("marketing: campaign %s recovered from a bounce spike: %s", cid, info)
    return info


async def start_campaign(cid: str) -> None:
    if cid in _running:
        return
    dispatch("mk_send", {"cid": cid}, dedupe=cid)


async def _run_campaign(cid: str) -> None:
    """Send a campaign's queued messages, rate-limited to the SES quota,
    re-checking pause / cancel and the bounce guard after every batch."""
    _running.add(cid)
    try:
        st = await get_settings()
        rate = 1.0
        camp = await db.mk_campaigns.find_one({"id": cid}, {"_id": 0})
        if not camp:
            return
        if camp["channel"] == "email":
            try:
                acct = await asyncio.to_thread(lambda: _ses().get_account())
                rate = max(1.0, float((acct.get("SendQuota") or {}).get("MaxSendRate") or 1)) * 0.8
            except Exception:  # noqa: BLE001
                log.warning("marketing: could not read SES quota; sending at 1/s", exc_info=True)
        while True:
            camp = await db.mk_campaigns.find_one({"id": cid}, {"_id": 0})
            if not camp or camp.get("status") != "sending":
                return
            await db.mk_campaigns.update_one({"id": cid}, {"$set": {"heartbeat": _iso()}})
            batch = await db.mk_sends.find({"campaign_id": cid, "status": "queued"}, {"_id": 0}).sort(
                [("priority", 1), ("queued_at", 1)]).limit(25).to_list(25)
            if not batch:
                await db.mk_campaigns.update_one({"id": cid, "status": "sending"}, {"$set": {"status": "sent", "finished_at": _iso()}})
                await refresh_campaign_stats(cid)
                return
            tail_batch = any(s.get("tail") for s in batch)
            pool = None  # latest releases, loaded once per batch when the email has "Picked for you"
            for s in batch:
                contact = await db.mk_contacts.find_one({"id": s["contact_id"]}, {"_id": 0}) or {}
                cs = contact.get("email_consent", {}).get("status")
                ok_now = (not contact.get("suppressed") and (cs == "subscribed" or (
                    camp.get("kind") == "reconfirm" and contact.get("permission_pending")
                    and cs not in ("unsubscribed", "complained", "lapsed")))) \
                    if camp["channel"] == "email" else contact.get("wa_consent", {}).get("status") == "subscribed"
                if not ok_now:  # unsubscribed / bounced since the campaign was queued
                    await db.mk_sends.update_one({"id": s["id"]}, {"$set": {"status": "skipped", "error": "no longer subscribed"}})
                    continue
                if camp["channel"] == "email":
                    try:
                        send_links = None
                        if _for_you_count(camp):
                            if pool is None:
                                pool = await release_pool()
                            extra = await attach_picks(contact, camp, pool)
                            send_links = (camp.get("links") or []) + [u for u in extra if u not in (camp.get("links") or [])]
                            await db.mk_sends.update_one({"id": s["id"]}, {"$set": {"links": send_links}})
                        html_doc, text_doc = render_for(camp, contact, s["id"], st, links=send_links)
                        raw = build_mime(st=st, to=contact["email"], subject=core.personalise(camp["subject"], contact),
                                         html_doc=html_doc, text_doc=text_doc, unsub_post=_links(s["id"])["unsub_post"])
                        mid = await asyncio.to_thread(_send_raw, raw)
                        await db.mk_sends.update_one({"id": s["id"]}, {"$set": {"status": "sent", "message_id": mid, "sent_at": _iso()}})
                        await db.mk_contacts.update_one({"id": contact["id"]}, {"$inc": {"stats.sent": 1, "since_engaged": 1}, "$set": {"last_sent_at": _iso()}})
                    except Exception as e:  # noqa: BLE001
                        await db.mk_sends.update_one({"id": s["id"]}, {"$set": {"status": "failed", "error": str(e)[:300]}})
                        if "Throttling" in str(e) or "MaxSendRate" in str(e):
                            await asyncio.sleep(2)
                    await asyncio.sleep(1.0 / rate)
                else:
                    from interakt import send_campaign_message
                    values = [core.personalise(v.get("value", ""), contact) if v.get("source") == "static"
                              else (contact.get("name") or "there").split(" ")[0] if v.get("source") == "first_name"
                              else (contact.get("name") or "there") for v in camp.get("wa_variables") or []]
                    rec = await send_campaign_message(template=camp["wa_template"], language=camp.get("wa_language") or "en",
                                                      values=values, phone=contact.get("phone"),
                                                      dedupe=f"campaign:{cid}:{contact['id']}", campaign_id=cid)
                    upd = {"status": "sent" if rec and rec.get("status") != "failed" else "failed", "sent_at": _iso(),
                           "wa_message_id": (rec or {}).get("id", ""), "error": (rec or {}).get("error", "") if rec else "not sent (WhatsApp is off?)"}
                    await db.mk_sends.update_one({"id": s["id"]}, {"$set": upd})
                    await asyncio.sleep(0.25)
            if tail_batch and camp["channel"] == "email":
                # Unconfirmed addresses go in small batches with a pause, so
                # their bounces (SES reports them seconds to minutes later) are
                # counted before the next batch — the tail can stop in time.
                await db.mk_campaigns.update_one({"id": cid}, {"$set": {"heartbeat": _iso()}})
                await asyncio.sleep(TAIL_BATCH_PAUSE)
            await refresh_campaign_stats(cid)
            camp = await db.mk_campaigns.find_one({"id": cid}, {"_id": 0, "stats": 1, "channel": 1, "recovered_at": 1, "name": 1})
            if camp["channel"] == "email":
                # The catch-all tail stops itself if it bounces; the campaign carries on.
                tail_base = {"campaign_id": cid, "tail": True}
                t_sent = await db.mk_sends.count_documents({**tail_base, "status": {"$in": ["sent", "bounced", "complained"]}})
                t_bounced = await db.mk_sends.count_documents({**tail_base, "status": "bounced"})
                if core.tail_should_stop(t_sent, t_bounced, max_rate=min(float(st.get("tail_max_bounce") or 0.03), MAX_BOUNCE_CEILING)):
                    r = await db.mk_sends.update_many({**tail_base, "status": "queued"},
                                                      {"$set": {"status": "skipped", "error": "catch-all tail stopped: too many bounces"}})
                    if r.modified_count:
                        await db.mk_campaigns.update_one({"id": cid}, {"$set": {"tail_stopped_at": _iso()}})
                since = camp.get("recovered_at")
                s_ = await _window_stats(cid, since) if since else (camp.get("stats") or {})
                why = core.should_pause(s_.get("sent", 0), s_.get("bounced", 0), s_.get("complained", 0),
                                        max_bounce=min(float(st.get("max_bounce") or 0.02), MAX_BOUNCE_CEILING))
                if why and not since:
                    await _recover(cid, why)   # first spike: clean the queue and carry on
                elif why:
                    await db.mk_campaigns.update_one({"id": cid, "status": "sending"},
                                                     {"$set": {"status": "paused", "paused_reason": f"Auto-paused again after recovering: {why}", "paused_at": _iso()}})
                    await raise_alert("paused", f"“{camp.get('name', '')}” paused: {why}, even after dropping unproven addresses. Check the bounce list before resuming.",
                                      campaign_id=cid)
                    return
    except Exception:  # noqa: BLE001
        log.exception("marketing: campaign %s worker crashed", cid)
    finally:
        _running.discard(cid)


async def refresh_campaign_stats(cid: str) -> dict:
    camp = await db.mk_campaigns.find_one({"id": cid}, {"_id": 0, "channel": 1, "id": 1})
    if not camp:
        return {}
    counts: dict = {}
    async for r in db.mk_sends.aggregate([{"$match": {"campaign_id": cid}}, {"$group": {"_id": "$status", "n": {"$sum": 1}}}]):
        counts[r["_id"]] = r["n"]
    agg = {"queued": counts.get("queued", 0), "failed": counts.get("failed", 0), "skipped": counts.get("skipped", 0),
           "blocked": counts.get("blocked", 0)}
    # "blocked" = SES Auto Validation refused it before sending: not sent, and
    # not a bounce, so it must not push the campaign towards auto-pause.
    sent_like = sum(v for k, v in counts.items() if k not in ("queued", "failed", "skipped", "blocked"))
    agg["sent"] = sent_like
    for fld in ("delivered_at", "opened_at", "clicked_at", "unsubscribed_at"):
        agg[fld.replace("_at", "")] = await db.mk_sends.count_documents({"campaign_id": cid, fld: {"$exists": True}})
    agg["bounced"] = counts.get("bounced", 0)
    agg["complained"] = counts.get("complained", 0)
    # The gap between "sent" and "delivered + bounced". A soft (Transient /
    # Undetermined) bounce keeps status "sent" — it is not a hard bounce — and
    # never gets delivered_at, so it used to vanish from every card: e.g. 186
    # sent, 181 delivered, 0.00% bounced, 5 unexplained. Counted separately so
    # the report adds up: delivered + bounced + soft_bounced + no_report = sent.
    # An open or click proves the mail arrived even when SES's Delivery
    # event never reached us, so it counts as delivered (and is never left
    # in "no report").
    unproven = {"delivered_at": {"$exists": False}, "opened_at": {"$exists": False}}
    agg["delivered"] = await db.mk_sends.count_documents(
        {"campaign_id": cid, "$or": [{"delivered_at": {"$exists": True}}, {"opened_at": {"$exists": True}}]})
    agg["soft_bounced"] = await db.mk_sends.count_documents(
        {"campaign_id": cid, "status": "sent", "bounce_type": {"$exists": True}, **unproven})
    # Sent, but SES has reported nothing back yet (normal for a few minutes;
    # after a day it means a missed Delivery event, not a failed send).
    agg["no_report"] = await db.mk_sends.count_documents(
        {"campaign_id": cid, "status": "sent", "bounce_type": {"$exists": False}, **unproven})
    if camp["channel"] == "whatsapp":
        wa = {}
        async for r in db.wa_messages.aggregate([{"$match": {"campaign_id": cid}}, {"$group": {"_id": "$status", "n": {"$sum": 1},
                                                                                             "cost": {"$sum": {"$ifNull": ["$cost", 0]}}}}]):
            wa[r["_id"]] = r
        agg["delivered"] = sum(wa.get(k, {}).get("n", 0) for k in ("delivered", "read"))
        agg["read"] = wa.get("read", {}).get("n", 0)
        agg["opened"] = agg["read"]
        agg["failed"] = agg["failed"] + wa.get("failed", {}).get("n", 0)
        agg["cost"] = round(sum(v.get("cost", 0) for v in wa.values()), 2)
    tag = core.campaign_tag(cid)
    orders = await db.orders.find({"attribution.utm_campaign": tag}, {"_id": 0, "payment_status": 1, "total": 1}).to_list(5000)
    agg["orders"] = len(orders)
    agg["paid"] = sum(1 for o in orders if o.get("payment_status") == "paid")
    agg["revenue"] = round(sum(float(o.get("total") or 0) for o in orders if o.get("payment_status") == "paid"), 2)
    await db.mk_campaigns.update_one({"id": cid}, {"$set": {"stats": agg, "stats_at": _iso()}})
    return agg


# =========================================================== public API ===
public_router = APIRouter(prefix="/api/m", tags=["marketing-public"])


@public_router.get("/o/{send_id}/{sig}.gif")
async def track_open(send_id: str, sig: str):
    if core.check_sig(_secret(), sig, "o", send_id):
        s = await db.mk_sends.find_one({"id": send_id}, {"_id": 0, "opened_at": 1, "contact_id": 1})
        if s:
            upd = {"$inc": {"open_count": 1}}
            if not s.get("opened_at"):
                upd["$set"] = {"opened_at": _iso()}
                await db.mk_contacts.update_one({"id": s["contact_id"]}, {"$inc": {"stats.opened": 1},
                                                                          "$set": {"proven": True, "since_engaged": 0, "last_engaged_at": _iso()}})
                await mark_contact_proven(s["contact_id"])
            await db.mk_sends.update_one({"id": send_id}, upd)
    return Response(content=GIF, media_type="image/gif", headers={"Cache-Control": "no-store, max-age=0"})


@public_router.get("/c/{send_id}/{idx}/{sig}")
async def track_click(send_id: str, idx: int, sig: str):
    if not core.check_sig(_secret(), sig, "c", send_id, str(idx)):
        return RedirectResponse(SITE, status_code=302)
    s = await db.mk_sends.find_one({"id": send_id}, {"_id": 0})
    camp = await db.mk_campaigns.find_one({"id": (s or {}).get("campaign_id")}, {"_id": 0, "links": 1, "id": 1, "kind": 1}) if s else None
    # Automated emails (growth.py) carry their own links per recipient.
    links = (s or {}).get("links") or (camp or {}).get("links") or []
    if not s or idx < 0 or idx >= len(links):
        return RedirectResponse(SITE, status_code=302)
    upd = {"$inc": {"click_count": 1}, "$push": {"clicks": {"$each": [{"i": idx, "at": _iso()}], "$slice": -50}}}
    sets = {}
    if not s.get("clicked_at"):
        sets["clicked_at"] = _iso()
        await db.mk_contacts.update_one({"id": s["contact_id"]}, {"$inc": {"stats.clicked": 1},
                                                                  "$set": {"proven": True, "since_engaged": 0, "last_engaged_at": _iso()}})
        await mark_contact_proven(s["contact_id"])
        if (camp or {}).get("kind") == "reconfirm":
            # A click (not an open — Apple Mail opens everything) on a
            # re-confirm email is a fresh, provable yes.
            await db.mk_contacts.update_one({"id": s["contact_id"], "suppressed": {"$in": [None, ""]},
                                             "email_consent.status": {"$nin": ["unsubscribed", "complained"]}},
                                            {"$set": {"email_consent": _consent("subscribed", "re-confirmed (campaign click)"),
                                                      "email_status": "verified", "reconfirmed_at": _iso(),
                                                      "permission_pending": False}})
    if not s.get("opened_at"):
        sets["opened_at"] = _iso()  # a click proves the open (images may be blocked)
    if sets:
        upd["$set"] = sets
    await db.mk_sends.update_one({"id": send_id}, upd)
    return RedirectResponse(core.add_utm(links[idx], camp["id"]), status_code=302)


async def _unsubscribe(send_id: str, how: str) -> Optional[dict]:
    s = await db.mk_sends.find_one({"id": send_id}, {"_id": 0})
    if not s:
        return None
    await db.mk_sends.update_one({"id": send_id, "unsubscribed_at": {"$exists": False}}, {"$set": {"unsubscribed_at": _iso()}})
    c = await db.mk_contacts.find_one({"id": s["contact_id"]}, {"_id": 0, "email": 1, "id": 1})
    if c:
        await db.mk_contacts.update_one({"id": c["id"]}, {"$set": {"email_consent": _consent("unsubscribed", how)}})
    return c


@public_router.get("/u/{send_id}/{sig}")
async def unsub_info(send_id: str, sig: str):
    if not core.check_sig(_secret(), sig, "u", send_id):
        raise HTTPException(status_code=404, detail="This link is not valid.")
    s = await db.mk_sends.find_one({"id": send_id}, {"_id": 0, "contact_id": 1})
    c = await db.mk_contacts.find_one({"id": (s or {}).get("contact_id")}, {"_id": 0, "email": 1, "email_consent": 1}) if s else None
    if not c:
        raise HTTPException(status_code=404, detail="This link is not valid.")
    local, _, dom = c["email"].partition("@")
    return {"email": f"{local[:2]}•••@{dom}", "subscribed": c.get("email_consent", {}).get("status") == "subscribed"}


@public_router.post("/u/{send_id}/{sig}")
async def unsub(send_id: str, sig: str):
    """Both the page's button and the mail app's one-click (RFC 8058) land here."""
    if not core.check_sig(_secret(), sig, "u", send_id):
        raise HTTPException(status_code=404, detail="This link is not valid.")
    c = await _unsubscribe(send_id, "unsubscribe link")
    if not c:
        raise HTTPException(status_code=404, detail="This link is not valid.")
    return {"ok": True}


class ConfirmBody(BaseModel):
    c: str
    t: str


@public_router.post("/confirm")
async def confirm_subscription(body: ConfirmBody):
    if not core.check_sig(_secret(), body.t, "confirm", body.c):
        raise HTTPException(status_code=404, detail="This link is not valid.")
    c = await db.mk_contacts.find_one({"id": body.c}, {"_id": 0})
    if not c:
        raise HTTPException(status_code=404, detail="This link is not valid.")
    if c.get("email_consent", {}).get("status") in ("unsubscribed", "complained"):
        return {"ok": True, "status": c["email_consent"]["status"]}
    await db.mk_contacts.update_one({"id": c["id"]}, {"$set": {
        "email_consent": _consent("subscribed", "confirmed (double opt-in)"), "proven": True,
        "email_status": "verified", "verified_at": _iso()}})
    return {"ok": True, "status": "subscribed"}


@public_router.post("/resend-events")
async def resend_events(request: Request):
    """Resend (order / account emails) delivery events, signed with Svix.
    delivered -> that mailbox is proven (free); bounced / suppressed ->
    never mail it again; complained -> unsubscribed + suppressed. Retries
    carry the same svix-id, so each event is applied once."""
    if not RESEND_WEBHOOK_SECRET:
        raise HTTPException(status_code=404, detail="Not found")
    body = await request.body()
    h = request.headers
    sid = h.get("svix-id", "")
    if not core.svix_verify(RESEND_WEBHOOK_SECRET, sid, h.get("svix-timestamp", ""), body, h.get("svix-signature", "")):
        raise HTTPException(status_code=400, detail="Invalid signature")
    await _ensure_indexes()
    try:
        await db.mk_webhook_seen.insert_one({"id": f"resend:{sid}", "at": datetime.now(timezone.utc)})
    except Exception:  # noqa: BLE001 — duplicate = already handled
        return {"ok": True, "duplicate": True}
    try:
        evt = json.loads(body or b"{}")
    except ValueError:
        raise HTTPException(status_code=400, detail="Bad JSON")
    typ = str(evt.get("type") or "")
    data = evt.get("data") or {}
    for addr in data.get("to") or []:
        norm = core.normalise_email(addr)
        if not norm:
            continue
        if typ == "email.delivered":
            await note_delivered(norm, "order/account email delivered (Resend)")
        elif typ in ("email.bounced", "email.suppressed"):
            await suppress({"email_norm": norm}, "bounced")
        elif typ == "email.complained":
            await suppress({"email_norm": norm}, "complained")
    await db.mk_events.insert_one({"at": _iso(), "type": f"resend:{typ}"})
    return {"ok": True}


@public_router.post("/ses-events/{token}")
async def ses_events(token: str, request: Request):
    """SNS -> SES delivery / bounce / complaint events.

    Guarded by a secret token in the URL and by the topic ARN (SNS_TOPIC_ARN,
    if set). A subscription confirmation is accepted only from an
    https://sns.<region>.amazonaws.com address — the URL is fetched, so it
    must never be somewhere an attacker chose."""
    import hmac as _h
    if not _h.compare_digest(token, sns_token()):
        raise HTTPException(status_code=404, detail="Not found")
    raw = await request.body()
    try:
        env = json.loads(raw or b"{}")
    except ValueError:
        raise HTTPException(status_code=400, detail="Bad JSON")
    want_arn = (os.environ.get("SNS_TOPIC_ARN") or "").strip()
    if want_arn and env.get("TopicArn") != want_arn:
        raise HTTPException(status_code=403, detail="Wrong topic")
    typ = env.get("Type")
    await db.mk_events.insert_one({"at": _iso(), "type": typ, "body": raw[:20000].decode("utf-8", "replace")})
    if typ == "SubscriptionConfirmation":
        url = env.get("SubscribeURL") or ""
        host = urlparse(url).hostname or ""
        if url.startswith("https://") and re.fullmatch(r"sns\.[a-z0-9-]+\.amazonaws\.com", host):
            await asyncio.to_thread(lambda: requests.get(url, timeout=10))
            return {"ok": True, "confirmed": True}
        raise HTTPException(status_code=400, detail="Unexpected SubscribeURL")
    if typ != "Notification":
        return {"ok": True}
    try:
        msg = json.loads(env.get("Message") or "{}")
    except ValueError:
        return {"ok": True}
    await handle_ses_event(msg)
    return {"ok": True}


async def handle_ses_event(msg: dict) -> None:
    kind = msg.get("eventType") or msg.get("notificationType") or ""
    mail = msg.get("mail") or {}
    mid = mail.get("messageId")
    s = await db.mk_sends.find_one({"message_id": mid}, {"_id": 0}) if mid else None
    if kind == "Delivery":
        if s:
            await db.mk_sends.update_one({"id": s["id"]}, {"$set": {"delivered_at": _iso()}})
            before = await db.mk_contacts.find_one_and_update({"id": s["contact_id"]}, {"$set": {"proven": True}},
                                                              projection={"_id": 0, "proven": 1, "email_norm": 1})
            await mark_contact_proven(s["contact_id"])
            dom = ((before or {}).get("email_norm") or "").split("@")[-1]
            if before and not before.get("proven") and dom and dom not in core.WEBMAIL:
                # One more distinct address at this company domain reached an
                # inbox: after 3 (and no hard bounces) the domain counts as good.
                await db.mk_domains.update_one({"domain": dom}, {"$inc": {"delivered_n": 1}}, upsert=True)
    elif kind == "Bounce" and (msg.get("bounce") or {}).get("bounceSubType") == "EmailValidationSuppressed":
        # SES Auto Validation stopped it before sending. Never mail it again,
        # but it never reached a mailbox provider, so it is not a bounce.
        for r in (msg.get("bounce") or {}).get("bouncedRecipients") or []:
            await suppress({"email_norm": core.normalise_email(r.get("emailAddress"))}, "invalid (SES validation)")
        if s:
            await db.mk_sends.update_one({"id": s["id"]}, {"$set": {"status": "blocked", "error": "SES validation: address unlikely to exist"}})
    elif kind == "Bounce":
        b = msg.get("bounce") or {}
        diag = " ".join(str(r.get("diagnosticCode") or "") for r in b.get("bouncedRecipients") or [])
        # Gmail "inactive / out of storage" and similar arrive as Transient but
        # mean a dead mailbox: treat as hard so it is never retried.
        permanent = b.get("bounceType") == "Permanent" or core.dead_mailbox(diag)
        for r in b.get("bouncedRecipients") or []:
            norm = core.normalise_email(r.get("emailAddress"))
            if permanent:
                await suppress({"email_norm": norm}, "bounced")
                dom = norm.split("@")[-1] if "@" in norm else ""
                if dom and dom not in core.WEBMAIL:
                    d = await db.mk_domains.find_one_and_update({"domain": dom}, {"$addToSet": {"hard_bounced": norm}},
                                                                upsert=True, return_document=True, projection={"_id": 0, "hard_bounced": 1})
                    if core.is_bad_domain(dom, len((d or {}).get("hard_bounced") or [])):
                        await db.mk_domains.update_one({"domain": dom}, {"$set": {"bad": True, "bad_at": _iso()}})
            else:
                await db.mk_contacts.update_one({"email_norm": norm}, {"$inc": {"soft_bounces": 1}})
                c = await db.mk_contacts.find_one({"email_norm": norm}, {"_id": 0, "soft_bounces": 1})
                if c and c.get("soft_bounces", 0) >= 3:
                    await suppress({"email_norm": norm}, "bounced")
        if s:
            await db.mk_sends.update_one({"id": s["id"]}, {"$set": {
                "status": "bounced" if permanent else s.get("status"), "bounce_type": b.get("bounceType"),
                "bounce_sub": b.get("bounceSubType"), "bounced_at": _iso(),
                "error": str(((b.get("bouncedRecipients") or [{}])[0]).get("diagnosticCode") or "")[:300]}})
    elif kind == "Complaint":
        for r in (msg.get("complaint") or {}).get("complainedRecipients") or []:
            await suppress({"email_norm": core.normalise_email(r.get("emailAddress"))}, "complained")
        if s:
            await db.mk_sends.update_one({"id": s["id"]}, {"$set": {"status": "complained", "complained_at": _iso()}})
    elif kind in ("Reject", "Rendering Failure"):
        if s:
            await db.mk_sends.update_one({"id": s["id"]}, {"$set": {"status": "failed", "error": kind}})


async def lapse_reconfirm() -> dict:
    """After a re-confirm campaign's window (default 30 days): everyone it
    reached who did NOT click — and hasn't re-subscribed some other way since
    — stops getting campaigns (consent 'lapsed'). The record stays; they can
    opt in again any time."""
    done = 0
    async for c in db.mk_campaigns.find({"kind": "reconfirm", "status": "sent", "lapsed_done": {"$ne": True}},
                                        {"_id": 0, "id": 1, "finished_at": 1, "started_at": 1, "reconfirm_days": 1, "name": 1}):
        due = (datetime.fromisoformat(c["finished_at"]) + timedelta(days=int(c.get("reconfirm_days") or 30))) if c.get("finished_at") else None
        if not due or due > _now():
            continue
        n = 0
        async for x in db.mk_sends.find({"campaign_id": c["id"], "status": "sent", "clicked_at": {"$exists": False}}, {"_id": 0, "contact_id": 1}):
            r = await db.mk_contacts.update_one(
                {"id": x["contact_id"], "email_consent.status": "subscribed", "email_consent.at": {"$lt": c.get("started_at") or c["finished_at"]}},
                {"$set": {"email_consent": _consent("lapsed", f"no reply to re-confirm “{c.get('name', '')}”")}})
            n += r.modified_count
            # Imported "not known to us" and did not click the permission
            # email: that was the one ask — never mailed again.
            r2 = await db.mk_contacts.update_one(
                {"id": x["contact_id"], "permission_pending": True,
                 "email_consent.status": {"$nin": ["subscribed", "unsubscribed", "complained"]}},
                {"$set": {"permission_pending": False, "permission_declined_at": _iso(),
                          "email_consent": _consent("lapsed", f"no reply to permission email “{c.get('name', '')}”")}})
            n += r2.modified_count
            # Whatever else happened (they unsubscribed instead), the one ask is used up.
            await db.mk_contacts.update_one({"id": x["contact_id"], "permission_pending": True},
                                            {"$set": {"permission_pending": False}})
        await db.mk_campaigns.update_one({"id": c["id"]}, {"$set": {"lapsed_done": True, "lapsed_count": n, "lapsed_at": _iso()}})
        await audit_log(db, "MARKETING_RECONFIRM_LAPSED", email="system", role="system", meta={"campaign": c["id"], "lapsed": n})
        done += n
    return {"lapsed": done}


async def learn_from_outcomes(force: bool = False) -> dict:
    """Self-learning, once a day. Looks at every campaign email sent in the
    last 180 days (and at least a day old, so late bounces are in) together
    with what we believed about the address when we sent it (band / score /
    signals, recorded at preparation). Then:
      * each signal's weight is nudged towards what actually bounced
        (core.learn_adjustments — bounded ±20, needs 30+ sends per signal);
      * the 'valid' line moves one step if 'valid' addresses bounced too
        much, or 'risky' ones turned out clean (core.calibrate_threshold).
    Hard rules (mailbox missing, bounced before…) are never learned away."""
    st = await get_settings()
    model = await db.mk_model.find_one({"key": "scoring"}, {"_id": 0}) or {}
    if not force and (model.get("updated_at") or "") > (_now() - timedelta(hours=24)).isoformat():
        return {"skipped": True}
    match = {"sent_at": {"$gte": (_now() - timedelta(days=180)).isoformat(), "$lte": (_now() - timedelta(days=1)).isoformat()},
             "status": {"$in": ["sent", "bounced", "complained"]}, "band": {"$exists": True}}
    total, bounced = 0, 0
    per_flag: dict = {}
    bands: dict = {}
    async for x in db.mk_sends.find(match, {"_id": 0, "status": 1, "flags": 1, "band": 1}):
        b = x.get("status") == "bounced"
        total += 1
        bounced += b
        for f in x.get("flags") or []:
            per_flag.setdefault(f, {"sent": 0, "bounced": 0})
            per_flag[f]["sent"] += 1
            per_flag[f]["bounced"] += b
        band = x.get("band") or "unknown"
        bands.setdefault(band, {"sent": 0, "bounced": 0})
        bands[band]["sent"] += 1
        bands[band]["bounced"] += b
    baseline = bounced / total if total else 0.02
    adjust = core.learn_adjustments({f: v for f, v in per_flag.items() if f in core.BASE_DELTAS}, baseline=baseline)
    cur = int((model.get("thresholds") or {}).get("valid") or st.get("score_valid") or core.DEFAULT_THRESHOLDS["valid"])
    valid_th = core.calibrate_threshold(cur, bands.get("valid", {}), bands.get("risky", {}))
    snap = {"at": _iso(), "sent": total, "baseline": round(baseline, 4), "valid": valid_th, "adjust": adjust}
    await db.mk_model.update_one({"key": "scoring"}, {
        "$set": {"key": "scoring", "adjust": adjust, "thresholds": {"valid": valid_th}, "updated_at": _iso(),
                 "stats": {"sent": total, "bounced": bounced, "baseline": round(baseline, 4), "per_flag": per_flag, "bands": bands}},
        "$push": {"history": {"$each": [snap], "$slice": -30}}}, upsert=True)
    _MODEL.update(at=0.0, doc=None)
    if valid_th != cur:
        log.info("marketing: learning moved the 'valid' line %s -> %s", cur, valid_th)
    return {"ok": True, **snap}


async def sync_aws_suppression(force: bool = False) -> dict:
    """Once a day: copy AWS's own suppression list (addresses SES itself
    refuses after a hard bounce / complaint) into Contacts, so the two never
    drift apart. Needs ses:ListSuppressedDestinations."""
    state = await db.integrations.find_one({"key": "marketing_state"}, {"_id": 0}) or {}
    last = state.get("suppression_synced_at") or ""
    if not force and last and last > (_now() - timedelta(hours=24)).isoformat():
        return {"skipped": True}
    start = datetime.fromisoformat(last) if last else datetime(2020, 1, 1, tzinfo=timezone.utc)
    n, token = 0, None
    try:
        client = _ses()
        while True:
            kw = {"PageSize": 1000, "StartDate": start}
            if token:
                kw["NextToken"] = token
            r = await asyncio.to_thread(lambda: client.list_suppressed_destinations(**kw))
            for d in r.get("SuppressedDestinationSummaries") or []:
                norm = core.normalise_email(d.get("EmailAddress"))
                c = await db.mk_contacts.find_one({"email_norm": norm, "suppressed": {"$in": [None, ""]}}, {"_id": 0, "id": 1})
                if c:
                    await suppress({"id": c["id"]}, "complained" if d.get("Reason") == "COMPLAINT" else "bounced")
                    n += 1
            token = r.get("NextToken")
            if not token:
                break
    except Exception as e:  # noqa: BLE001
        code = ((getattr(e, "response", None) or {}).get("Error") or {}).get("Code", "") or str(e)[:120]
        await raise_alert("aws_suppression", f"Could not read the AWS suppression list ({code}). "
                                             "Allow ses:ListSuppressedDestinations for the site's AWS user.", key="aws_suppression")
        return {"ok": False, "error": code}
    await db.integrations.update_one({"key": "marketing_state"}, {"$set": {"key": "marketing_state", "suppression_synced_at": _iso(),
                                                                           "suppression_last_added": n}}, upsert=True)
    await db.mk_alerts.update_many({"key": "aws_suppression", "resolved": False}, {"$set": {"resolved": True, "resolved_at": _iso()}})
    return {"ok": True, "suppressed": n}


# ============================================================ admin API ===
admin_router = APIRouter(prefix="/api/admin/marketing", tags=["marketing-admin"], dependencies=[Depends(require_admin)])


# ---- settings + health ----
@admin_router.get("/settings")
async def adm_settings(user: dict = Depends(require_admin)):
    from rbac import is_superadmin
    st = await get_settings()
    url = f"{PUBLIC_API}/api/m/ses-events/{sns_token()}"
    return {**st, "ses_region": SES_REGION, "configuration_set": SES_CONFIG_SET,
            "sns_url": url if is_superadmin(user.get("role")) else f"{PUBLIC_API}/api/m/ses-events/••••••",
            "last_ses_event": await db.mk_events.find_one({}, {"_id": 0, "at": 1, "type": 1}, sort=[("at", -1)]),
            "ses_usage": await ses_usage(),
            "waiting_checks": await db.mk_contacts.count_documents({"email_status": {"$in": ["unknown", "valid", "risky"]},
                                                                     "ses_checked": {"$ne": True}}),
            "suppression_sync": await db.integrations.find_one({"key": "marketing_state"}, {"_id": 0, "suppression_synced_at": 1,
                                                                                             "suppression_last_added": 1}),
            "resend": {"url": f"{PUBLIC_API}/api/m/resend-events", "configured": bool(RESEND_WEBHOOK_SECRET),
                       "last_event": await db.mk_events.find_one({"type": {"$regex": "^resend:"}}, {"_id": 0, "at": 1, "type": 1}, sort=[("at", -1)]),
                       "proven": await db.mk_proven.count_documents({})},
            "disposable": await db.integrations.find_one({"key": "disposable_list"}, {"_id": 0, "count": 1, "at": 1})}


class SettingsBody(BaseModel):
    from_name: Optional[str] = None
    from_email: Optional[str] = None
    reply_to: Optional[str] = None
    footer: Optional[str] = None
    logo: Optional[str] = None
    double_opt_in: Optional[bool] = None
    include_role_addresses: Optional[bool] = None
    max_bounce: Optional[float] = None
    wa_rate_marketing: Optional[float] = None
    wa_rate_utility: Optional[float] = None
    extra_disposable: Optional[list] = None
    ses_validation: Optional[bool] = None
    ses_budget_inr: Optional[float] = None
    usd_inr: Optional[float] = None
    risk_policy: Optional[dict] = None
    tail_max_bounce: Optional[float] = None
    score_valid: Optional[int] = None
    score_invalid: Optional[int] = None
    auto_learn: Optional[bool] = None
    inactive_after: Optional[int] = None
    sunset_after: Optional[int] = None


@admin_router.put("/settings")
async def adm_put_settings(body: SettingsBody, user: dict = Depends(require_superadmin)):
    upd = {k: v for k, v in body.model_dump().items() if v is not None}
    for k in ("from_email", "reply_to"):
        if upd.get(k) and not core.check_syntax(upd[k])["ok"]:
            raise HTTPException(status_code=400, detail=f"{k.replace('_', ' ')} is not a valid email address")
    # AWS places an SES account under review at a 5% bounce rate and may
    # pause it at 10%. Our own stops must trip BEFORE that, so 4% is the
    # ceiling — a setting of 5% only paused a campaign once AWS was already
    # reviewing the account (and the receipts share that account).
    if "max_bounce" in upd and not (0.005 <= upd["max_bounce"] <= MAX_BOUNCE_CEILING):
        raise HTTPException(status_code=400, detail="Auto-pause must be between 0.5% and 4% bounces (AWS reviews the account at 5%)")
    if "ses_budget_inr" in upd and not (0 <= upd["ses_budget_inr"] <= 100000):
        raise HTTPException(status_code=400, detail="Monthly mailbox-check budget must be between ₹0 and ₹1,00,000")
    if "usd_inr" in upd and not (50 <= upd["usd_inr"] <= 200):
        raise HTTPException(status_code=400, detail="USD→INR rate looks wrong (50–200)")
    if "score_valid" in upd and not (50 <= upd["score_valid"] <= 90):
        raise HTTPException(status_code=400, detail="'Valid from' score must be 50–90")
    if "score_invalid" in upd and not (10 <= upd["score_invalid"] <= 60):
        raise HTTPException(status_code=400, detail="'Invalid below' score must be 10–60")
    if upd.get("score_invalid") and upd.get("score_valid") and upd["score_invalid"] >= upd["score_valid"] - 5:
        raise HTTPException(status_code=400, detail="'Invalid below' must be at least 5 under 'Valid from'")
    for k in ("inactive_after", "sunset_after"):
        if k in upd and not (0 <= upd[k] <= 100):
            raise HTTPException(status_code=400, detail="Inactive / retire counts must be 0–100 (0 = off)")
    if upd.get("inactive_after") and upd.get("sunset_after") and upd["sunset_after"] < upd["inactive_after"]:
        raise HTTPException(status_code=400, detail="'Stop mailing after' must be at least 'Mark inactive after'")
    if "tail_max_bounce" in upd and not (0.005 <= upd["tail_max_bounce"] <= MAX_BOUNCE_CEILING):
        raise HTTPException(status_code=400, detail="'Send last' bounce limit must be between 0.5% and 4% (AWS reviews the account at 5%)")
    if "risk_policy" in upd:
        pol = upd["risk_policy"] or {}
        bad = [k for k, v in pol.items() if k in core.TAIL_ONLY_KINDS and v == "send"]
        if bad:
            raise HTTPException(status_code=400, detail="Unconfirmed and catch-all addresses can be 'Send last' or 'Skip', not mixed into the main send")
        if any(k not in core.RISK_KINDS or v not in ("send", "tail", "skip") for k, v in pol.items()):
            raise HTTPException(status_code=400, detail="Risk rules: each kind must be send, tail or skip")
        upd["risk_policy"] = {**core.DEFAULT_RISK_POLICY, **pol}
    if "extra_disposable" in upd:
        upd["extra_disposable"] = [str(d).strip().lower() for d in upd["extra_disposable"] if str(d).strip()][:500]
    before = await get_settings(fresh=True)
    await db.integrations.update_one({"key": "marketing"}, {"$set": {"key": "marketing", **upd, "updated_at": _iso()}}, upsert=True)
    if "include_role_addresses" in upd and upd["include_role_addresses"] != before.get("include_role_addresses"):
        dispatch("mk_role_rule", dedupe="all")
    await audit_log(db, "MARKETING_SETTINGS", email=user.get("email", ""), role=user.get("role", ""), meta={k: v for k, v in upd.items() if k != "footer"})
    return await get_settings(fresh=True)


_HEALTH_MEMO = None


@admin_router.get("/providers")
async def adm_providers():
    """What the checker has learned per provider (Settings): providers whose
    unconfirmed addresses bounce too often, and ones SES can't confirm."""
    learned = await learn_providers()
    rows = await db.mk_domains.find({"$or": [{"bad_provider": True}, {"ses_unverifiable": True}, {"bad": True}]},
                                    {"_id": 0, "domain": 1, "bad_provider": 1, "ses_unverifiable": 1, "bad": 1,
                                     "unproven_sent": 1, "unproven_bounced": 1, "ses_checks": 1, "ses_confirmed": 1,
                                     "hard_bounced": 1}).to_list(500)
    for r in rows:
        r["hard_bounced"] = len(r.get("hard_bounced") or [])
    return {"rows": sorted(rows, key=lambda r: -(r.get("unproven_bounced") or 0)), "summary": learned,
            "rules": {"min_sends": core.PROVIDER_MIN_SENDS, "max_bounce": core.PROVIDER_MAX_BOUNCE}}


@admin_router.get("/health")
async def adm_health():
    """Is SES ready to send from our address? Read straight from SES — two AWS
    calls (~2 s), so the answer is reused for 60 s instead of on every visit."""
    global _HEALTH_MEMO
    if _HEALTH_MEMO is None:
        from cache import Memo, shared_version
        _HEALTH_MEMO = Memo(shared_version(db))
    return await _HEALTH_MEMO.get("ses_health", 60, _ses_health)


async def _ses_health():
    st = await get_settings()
    out: dict = {"region": SES_REGION}
    try:
        a = await asyncio.to_thread(lambda: _ses().get_account())
        q = a.get("SendQuota") or {}
        out.update({"ok": True, "sending_enabled": a.get("SendingEnabled"), "production": a.get("ProductionAccessEnabled"),
                    "max_24h": q.get("Max24HourSend"), "sent_24h": q.get("SentLast24Hours"), "max_rate": q.get("MaxSendRate"),
                    "enforcement": a.get("EnforcementStatus")})
    except Exception as e:  # noqa: BLE001
        out.update({"ok": False, "error": str(e)[:300]})
    # AWS's OWN reputation numbers (what its review / pause decisions use),
    # from CloudWatch. Needs cloudwatch:GetMetricStatistics on the IAM user;
    # without it the card says so instead of guessing.
    try:
        import boto3
        cw = boto3.client("cloudwatch", region_name=SES_REGION)
        end = datetime.now(timezone.utc)
        rep_: dict = {}
        for key, metric in (("bounce_rate", "Reputation.BounceRate"), ("complaint_rate", "Reputation.ComplaintRate")):
            r = await asyncio.to_thread(lambda m=metric: cw.get_metric_statistics(
                Namespace="AWS/SES", MetricName=m, StartTime=end - timedelta(hours=24), EndTime=end,
                Period=3600, Statistics=["Maximum"]))
            pts = sorted(r.get("Datapoints") or [], key=lambda d: d["Timestamp"])
            rep_[key] = pts[-1]["Maximum"] if pts else None
        out["aws_reputation"] = {**rep_, "review_at": {"bounce": 0.05, "complaint": 0.001},
                                 "pause_at": {"bounce": 0.10, "complaint": 0.005}}
    except Exception as e:  # noqa: BLE001
        out["aws_reputation"] = {"error": str(e)[:200]}
    dom = st["from_email"].split("@")[-1]
    for ident in (st["from_email"], dom):
        try:
            i = await asyncio.to_thread(lambda: _ses().get_email_identity(EmailIdentity=ident))
            out["identity"] = {"name": ident, "verified": i.get("VerifiedForSendingStatus"),
                               "dkim": (i.get("DkimAttributes") or {}).get("Status")}
            break
        except Exception:  # noqa: BLE001
            continue
    since = (_now() - timedelta(days=30)).isoformat()
    sent = await db.mk_sends.count_documents({"sent_at": {"$gte": since}, "message_id": {"$exists": True}})
    bounced = await db.mk_sends.count_documents({"sent_at": {"$gte": since}, "status": "bounced"})
    compl = await db.mk_sends.count_documents({"sent_at": {"$gte": since}, "status": "complained"})
    out["last30"] = {"sent": sent, "bounce_rate": core.rate(bounced, sent), "complaint_rate": core.rate(compl, sent)}
    return out


# ---- verification + import ----
class VerifyBody(BaseModel):
    emails: list[str]
    autofix: bool = True
    ses: bool = True   # also ask SES whether each mailbox exists (paid, capped)


@admin_router.post("/verify")
async def adm_verify(body: VerifyBody):
    """Check up to 200 typed / pasted addresses without saving anything."""
    rows = await verify_emails([e for e in body.emails if str(e).strip()][:200], autofix=body.autofix,
                               ses="live" if body.ses else "cache")
    return {"results": rows, "counts": _count(rows)}


def _count(rows: list) -> dict:
    c = {"total": len(rows)}
    for r in rows:
        c[r["status"]] = c.get(r["status"], 0) + 1
    c["fixed"] = sum(1 for r in rows if r.get("fixed_from"))
    return c


def _read_table(data: bytes, filename: str) -> list[dict]:
    """Rows from .xlsx / .csv as dicts keyed by lower-cased header."""
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = [str(h or "").strip().lower() for h in next(it, [])]
        return [{header[i]: ("" if v is None else str(v)).strip() for i, v in enumerate(r) if i < len(header)} for r in it]
    text = data.decode("utf-8-sig", errors="replace")
    rd = csv.DictReader(io.StringIO(text))
    return [{(k or "").strip().lower(): (v or "").strip() for k, v in r.items()} for r in rd]


def _pick(cols: list, *names) -> Optional[str]:
    for n in names:
        for c in cols:
            if c == n:
                return c
    for n in names:
        for c in cols:
            if n in c:
                return c
    return None


@admin_router.post("/import")
async def adm_import(
    file: UploadFile = File(...),
    list_name: str = Form(""),
    consent_email: bool = Form(False),
    consent_whatsapp: bool = Form(False),
    consent_source: str = Form(""),
    autofix: bool = Form(True),
    dry_run: bool = Form(True),
    statuses: str = Form("verified,valid,risky"),
    ses_live: bool = Form(False),
    readd_erased: bool = Form(False),
    audience: str = Form(""),
    user: dict = Depends(require_admin),
):
    """Upload a sheet -> verification report. dry_run=True only reports;
    dry_run=False STAGES the import for approval (nothing is saved to
    Contacts until someone else with Marketing access approves it — see
    adm_import_approve).

    audience (15 Oct 2026, replaces the "these people agreed" tick box):
      known   — existing clients/customers known to the company: imported as
                consented to marketing email, source recorded at approval
                ("existing client/customer — <list>, approved by <approver>");
      unknown — not known to us: imported WITHOUT consent and marked
                permission_pending, so the only email they can get is one
                permission request (a re-confirm campaign); a click opts them
                in, no click means never mailed again."""
    from rbac import is_superadmin
    if readd_erased and not is_superadmin(user.get("role")):
        raise HTTPException(status_code=403, detail="Only a superadmin can re-add erased people")
    # Bounded read (an unbounded file.read() materialises any size on the
    # 512 MB server) and no zip bomb for openpyxl.
    import uploads_guard
    data = await uploads_guard.read_limited(file, 25 * 1024 * 1024, "File")
    uploads_guard.require_not_zip_bomb(data, "File")
    if len(data) > 8_000_000:
        raise HTTPException(status_code=413, detail="File too large (8 MB max)")
    try:
        rows = _read_table(data, file.filename)
    except Exception as e:  # noqa: BLE001
        raise HTTPException(status_code=400, detail=f"Could not read the file: {str(e)[:120]}")
    # Completely empty rows are not people: Excel often reports rows that only
    # ever had formatting (a live file had 21 after its last entry, which showed
    # up as "— empty" invalid addresses). Dropped and counted, never listed.
    blank_rows = sum(1 for r in rows if not any(str(v or "").strip() for v in r.values()))
    rows = [r for r in rows if any(str(v or "").strip() for v in r.values())]
    if not rows:
        raise HTTPException(status_code=400, detail="The file has no rows")
    if len(rows) > 50000:
        raise HTTPException(status_code=400, detail="50,000 rows at a time, please split the file")
    cols = list(rows[0].keys())
    ecol = _pick(cols, "email", "e-mail", "email address", "mail")
    if not ecol:
        raise HTTPException(status_code=400, detail=f"No email column found (columns: {', '.join(cols[:12])})")
    ncol = _pick(cols, "name", "full name", "first name")
    pcol = _pick(cols, "phone", "mobile", "whatsapp", "contact")
    if not dry_run and audience not in ("known", "unknown"):
        raise HTTPException(status_code=400, detail=(
            "Choose who these people are (known clients/customers, or not known to us). "
            "Refresh the page if you don't see the choice."))
    if consent_whatsapp and len(consent_source.strip()) < 3:
        raise HTTPException(status_code=400, detail="Say where these people agreed to WhatsApp messages (e.g. 'Law summit 2026 registration form')")
    # duplicates inside the file (same mailbox)
    seen, unique_rows, dups = set(), [], 0
    for r in rows:
        k = core.normalise_email(core.check_syntax(r.get(ecol, ""))["email"])
        if k and k in seen:
            dups += 1
            continue
        if k:
            seen.add(k)
        unique_rows.append(r)
    # Paid SES mailbox checks only when the admin ticked "Also ask Amazon SES"
    # (within the monthly budget; never twice for one address). Otherwise
    # cached verdicts only, and a real import hands the checks to the capped
    # background pass (kick_validation).
    results = await verify_emails([r.get(ecol, "") for r in unique_rows], autofix=autofix, ses="live" if ses_live else "cache", ignore_erased=readd_erased,
                                  names=[(r.get(ncol, "") if ncol else "") for r in unique_rows])
    st = await get_settings()
    counts = _count(results)
    counts["duplicates_in_file"] = dups
    counts["blank_rows"] = blank_rows
    saved = 0
    skipped: dict = {}
    lid = None
    pending_id = None
    would_save = 0
    if not dry_run:
        # Staged, not saved: what WOULD be imported is kept with its verdicts
        # (no second paid check at approval), and someone else with Marketing
        # access approves or rejects it (adm_import_approve / _reject).
        name = list_name.strip() or f"Import {_now():%d %b %Y %H:%M}"
        wanted = core.selected_statuses({"statuses": [x.strip() for x in statuses.split(",")]})

        def skip(why: str) -> None:
            skipped[why] = skipped.get(why, 0) + 1

        pending_id = str(uuid.uuid4())
        staged = []
        for r, v in zip(unique_rows, results):
            if v["status"] == "invalid":
                skip("invalid address")
                continue
            if v["status"] == "suppressed":
                skip("erased at a privacy request" if "previously erased at a privacy request" in v["reasons"]
                     else "bounced / unsubscribed before")
                continue
            if v["status"] not in wanted:
                skip(f"{v['status']} — not ticked")
                continue
            staged.append({"import_id": pending_id, "i": len(staged), "email": v["email"],
                           "name": str(r.get(ncol, "") if ncol else "")[:120], "phone": str(r.get(pcol, "") if pcol else "")[:40],
                           "v": {k: v.get(k) for k in ("status", "reasons", "email_norm", "ses_checked", "risk_kind",
                                                       "ses_overall", "score", "flags", "role")}})
        for k in range(0, len(staged), 1000):
            await db.mk_import_rows.insert_many(staged[k:k + 1000])
        would_save = len(staged)
        await db.mk_imports.insert_one({
            "id": pending_id, "status": "pending", "at": _iso(), "by": user.get("email", ""), "by_role": user.get("role", ""),
            "file": file.filename, "list_name": name[:80], "audience": audience, "consent_whatsapp": bool(consent_whatsapp),
            "consent_source": consent_source.strip()[:120], "readd_erased": bool(readd_erased),
            "counts": counts, "would_save": would_save, "skipped": skipped,
            "ses_after": sum(1 for x in staged if not x["v"].get("ses_checked") and x["v"]["status"] in ("valid", "risky"))})
        await raise_alert("import", f"Import “{name[:60]}” — {would_save} contacts "
                          f"({'known clients/customers' if audience == 'known' else 'not known to us'}) is waiting for approval. "
                          f"Uploaded by {user.get('email', '')}. Marketing → Verify & import.", key=f"import:{pending_id}")
        await audit_log(db, "MARKETING_IMPORT_STAGED", email=user.get("email", ""), role=user.get("role", ""),
                        meta={"import": pending_id, "file": file.filename, "would_save": would_save, "audience": audience,
                              "counts": counts, "skipped": skipped})
    sample: dict = {}
    for v in results:
        sample.setdefault(v["status"], [])
        if len(sample[v["status"]]) < 25:
            sample[v["status"]].append({k: v[k] for k in ("input", "email", "reasons", "fixed_from")})
    # Every repair the CRM made, so the admin can see exactly what changed.
    fixed = [{"from": v["fixed_from"], "to": v["email"], "why": next((x for x in v.get("reasons") or [] if x.startswith("fixed")), "")}
             for v in results if v.get("fixed_from")][:200]
    ses_todo = sum(1 for v in results if v["status"] in ("valid", "risky") and not v.get("ses_checked"))
    # Every row (first 2,000) with its verdict, so the admin sees each address, not just examples.
    row_report = [{"email": v["email"] or v["input"], "name": (r.get(ncol, "") if ncol else "")[:60], "status": v["status"],
                   "score": v.get("score"), "why": (v.get("reasons") or [])[:3], "risk_kind": v.get("risk_kind", "")}
                  for r, v in zip(unique_rows[:2000], results[:2000])]
    return {"columns": {"email": ecol, "name": ncol, "phone": pcol}, "counts": counts, "sample": sample, "fixed": fixed,
            "saved": saved, "list_id": lid, "dry_run": dry_run, "rows": row_report,
            "pending_id": pending_id, "would_save": would_save,
            "skipped": skipped if not dry_run else {},
            "erased": sum(1 for v in results if "previously erased at a privacy request" in v["reasons"]),
            "ses": {**(_u := await ses_usage()), "to_check": ses_todo, "est_inr": round(ses_todo * _u["per_check_inr"], 2)}}


# ------------------------------------------------- import: cost + approval ---
async def _estimate_checks(raw_emails: list, autofix: bool = True) -> dict:
    """What a paid SES check of these addresses would cost — read from our own
    records only, so the estimate itself is free.

    Counted: distinct, well-formed addresses (after the same typo fixes) that
    have no fresh SES verdict and are not already proven. That is the most a
    check can spend; it can spend less (budget cap, catch-all domains). Company
    domains may cost one extra catch-all test each (at most once per 90 days).
    """
    st = await get_settings()
    norms = set()
    for r in raw_emails:
        sx = core.check_syntax(r)
        if autofix and sx["ok"] and sx["suggestion"]:
            sx = core.check_syntax(sx["suggestion"])
        if sx["ok"]:
            n = core.normalise_email(sx["email"])
            if n:
                norms.add(n)
    norms = list(norms)
    cached = set((await _ses_cached(norms)).keys()) if norms else set()
    proven = await _proven(norms) if norms else set()
    todo = [n for n in norms if n not in cached and n not in proven]
    company = {n.split("@")[-1] for n in todo} - set(core.WEBMAIL)
    u = await ses_usage()
    per = u["per_check_inr"]
    return {"addresses": len(norms), "already_checked": len(cached | proven), "to_check": len(todo),
            "est_inr": round(len(todo) * per, 2), "company_domains": len(company),
            "domain_tests_max_inr": round(len(company) * per, 2), "per_check_inr": per,
            "budget_left_checks": u["left"], "budget_left_inr": round(u["left"] * per, 2),
            "budget_inr": u["budget_inr"], "enabled": u["enabled"] and u["supported"],
            "over_budget": len(todo) > u["left"]}


@admin_router.post("/import/estimate")
async def adm_import_estimate(file: Optional[UploadFile] = File(None), emails: str = Form(""),
                              autofix: bool = Form(True), user: dict = Depends(require_admin)):
    """Cost of checking a file (or pasted addresses) with Amazon SES, shown
    above the Check button before anything is spent."""
    raw: list = []
    if file is not None:
        import uploads_guard
        data = await uploads_guard.read_limited(file, 25 * 1024 * 1024, "File")
        uploads_guard.require_not_zip_bomb(data, "File")
        try:
            rows = _read_table(data, file.filename)
        except Exception as e:  # noqa: BLE001
            raise HTTPException(status_code=400, detail=f"Could not read the file: {str(e)[:120]}")
        if rows:
            ecol = _pick(list(rows[0].keys()), "email", "e-mail", "email address", "mail")
            if not ecol:
                raise HTTPException(status_code=400, detail="No email column found")
            raw = [r.get(ecol, "") for r in rows[:50000]]
    else:
        raw = [x.strip() for x in re.split(r"[\n,;]+", emails or "") if x.strip()][:5000]
    return await _estimate_checks(raw, autofix)


def _import_view(d: dict) -> dict:
    return {k: d.get(k) for k in ("id", "status", "at", "by", "file", "list_name", "audience", "consent_whatsapp",
                                  "consent_source", "counts", "would_save", "skipped", "ses_after", "readd_erased",
                                  "approved_by", "approved_at", "rejected_by", "rejected_at", "reject_note", "saved", "list_id")}


@admin_router.get("/imports")
async def adm_imports(status: str = "pending", user: dict = Depends(require_admin)):
    """Imports waiting for approval (or recent decided ones), with a few rows
    each so the approver can see who is in it."""
    q = {"status": status} if status in ("pending", "approved", "rejected") else {}
    out = []
    for d in await db.mk_imports.find(q, {"_id": 0}).sort("at", -1).limit(50).to_list(50):
        v = _import_view(d)
        if d.get("status") == "pending":
            v["sample"] = [{"email": r["email"], "name": r.get("name", ""), "status": r["v"]["status"]}
                           async for r in db.mk_import_rows.find({"import_id": d["id"]}, {"_id": 0}).sort("i", 1).limit(15)]
        v["can_approve"] = _approval_refusal(d, user) is None
        v["approval_refusal"] = _approval_refusal(d, user)
        out.append(v)
    return {"imports": out}


def _approval_refusal(d: dict, user: dict) -> Optional[str]:
    from rbac import is_superadmin
    if d.get("status") != "pending":
        return "Already decided."
    if (d.get("by") or "").lower() == (user.get("email") or "").lower() and not is_superadmin(user.get("role")):
        return "Uploaded by you — someone else with Marketing access must approve it."
    if d.get("readd_erased") and not is_superadmin(user.get("role")):
        return "This import re-adds erased people — only a superadmin can approve it."
    return None


class ImportDecision(BaseModel):
    note: Optional[str] = ""


@admin_router.post("/imports/{iid}/approve")
async def adm_import_approve(iid: str, user: dict = Depends(require_admin)):
    """Save a staged import into Contacts. Only someone other than the
    uploader (a superadmin may approve their own). The approver's name goes
    on the consent record of a "known clients/customers" import."""
    d = await db.mk_imports.find_one({"id": iid}, {"_id": 0})
    if not d:
        raise HTTPException(status_code=404, detail="Import not found")
    why = _approval_refusal(d, user)
    if why:
        raise HTTPException(status_code=403 if "someone else" in why or "superadmin" in why else 409, detail=why)
    claim = await db.mk_imports.update_one({"id": iid, "status": "pending"}, {"$set": {"status": "approving"}})
    if claim.modified_count != 1:
        raise HTTPException(status_code=409, detail="Someone else is deciding this import right now.")
    lid = str(uuid.uuid4())
    await db.mk_lists.insert_one({"id": lid, "name": d["list_name"], "kind": "static", "created_at": _iso(),
                                  "created_by": d.get("by", ""), "approved_by": user.get("email", ""),
                                  "audience": d.get("audience")})
    known = d.get("audience") == "known"
    src = (f"import: existing client/customer — {d['list_name']} (approved by {user.get('email', '')})" if known
           else f"import: {d['list_name']} (not known — permission email first)")[:120]
    wa = bool(d.get("consent_whatsapp"))
    readd = bool(d.get("readd_erased"))
    saved, readded, pending_perm = 0, 0, 0
    skipped = dict(d.get("skipped") or {})
    async for r in db.mk_import_rows.find({"import_id": iid}, {"_id": 0}).sort("i", 1):
        v = r["v"]
        was_erased = readd and bool(await _erased([v.get("email_norm")]))
        doc = await upsert_contact(r["email"], name=r.get("name", ""), phone=r.get("phone", ""), source=src, lists=[lid],
                                   status=v["status"], reasons=v.get("reasons") or [],
                                   proven=v["status"] == "verified" or None, ses_checked=v.get("ses_checked") or None,
                                   extra={"risk_kind": v.get("risk_kind", ""), "ses_overall": v.get("ses_overall", ""),
                                          "score": v.get("score"), "flags": v.get("flags") or [], "role": bool(v.get("role"))},
                                   email_consent="subscribed" if known else None,
                                   wa_consent="subscribed" if wa else None, by=user.get("email", ""), readd_erased=readd)
        if not doc:
            skipped["erased at a privacy request"] = skipped.get("erased at a privacy request", 0) + 1
            continue
        saved += 1
        readded += was_erased
        if not known:
            # Only people with no consent decision yet. Someone already
            # subscribed stays subscribed; someone who unsubscribed, complained
            # or lapsed is never asked again.
            pr = await db.mk_contacts.update_one(
                {"id": doc["id"], "email_consent.status": {"$nin": ["subscribed", "unsubscribed", "complained", "lapsed"]}},
                {"$set": {"permission_pending": True, "permission_list": lid, "permission_since": _iso()}})
            pending_perm += pr.modified_count
    await db.mk_import_rows.delete_many({"import_id": iid})
    await db.mk_imports.update_one({"id": iid}, {"$set": {
        "status": "approved", "approved_by": user.get("email", ""), "approved_at": _iso(), "list_id": lid,
        "saved": saved, "skipped": skipped, "permission_pending": pending_perm}})
    await db.mk_alerts.update_many({"key": f"import:{iid}", "resolved": False},
                                   {"$set": {"resolved": True, "resolved_at": _iso(), "resolved_by": user.get("email", "")}})
    await audit_log(db, "MARKETING_IMPORT", email=user.get("email", ""), role=user.get("role", ""),
                    meta={"import": iid, "file": d.get("file"), "uploaded_by": d.get("by"), "saved": saved,
                          "audience": d.get("audience"), "consent_email": known, "consent_whatsapp": wa,
                          "permission_pending": pending_perm, "skipped": skipped, "readded_erased": readded})
    kick_validation()
    return {"ok": True, "saved": saved, "list_id": lid, "skipped": skipped, "permission_pending": pending_perm}


@admin_router.post("/imports/{iid}/reject")
async def adm_import_reject(iid: str, body: ImportDecision, user: dict = Depends(require_admin)):
    """Nothing is saved; the staged rows are deleted. Anyone with Marketing
    access may reject, the uploader included (withdrawing their own)."""
    r = await db.mk_imports.update_one({"id": iid, "status": "pending"}, {"$set": {
        "status": "rejected", "rejected_by": user.get("email", ""), "rejected_at": _iso(), "reject_note": (body.note or "")[:300]}})
    if r.modified_count != 1:
        raise HTTPException(status_code=409, detail="This import is no longer waiting for approval.")
    await db.mk_import_rows.delete_many({"import_id": iid})
    await db.mk_alerts.update_many({"key": f"import:{iid}", "resolved": False},
                                   {"$set": {"resolved": True, "resolved_at": _iso(), "resolved_by": user.get("email", "")}})
    await audit_log(db, "MARKETING_IMPORT_REJECTED", email=user.get("email", ""), role=user.get("role", ""),
                    meta={"import": iid, "note": (body.note or "")[:300]})
    return {"ok": True}


@admin_router.post("/alerts/{aid}/dismiss")
async def adm_dismiss_alert(aid: str, user: dict = Depends(require_admin)):
    await db.mk_alerts.update_one({"id": aid}, {"$set": {"resolved": True, "resolved_at": _iso(), "resolved_by": user.get("email", "")}})
    return {"ok": True}


@admin_router.get("/system-health")
async def adm_system_health():
    """Admin -> Marketing -> Settings -> System health: is everything fast and
    is background work keeping up? Numbers come from this website process
    (timings, cache, memory) and from the shared job queue."""
    import perf
    from cache import PUBLIC_STORE
    q = await _jobs.stats(db)
    w = q.get("worker") or {}
    if w.get("at"):
        at = w["at"] if w["at"].tzinfo else w["at"].replace(tzinfo=timezone.utc)
        w["seconds_ago"] = round((_now() - at).total_seconds())
        w["at"] = at.isoformat()
    for d in q.get("recent_dead") or []:
        if d.get("finished_at") and hasattr(d["finished_at"], "isoformat"):
            d["finished_at"] = d["finished_at"].isoformat()
        d["error"] = (d.get("error") or "").splitlines()[0][:200] if d.get("error") else ""
    return {"queue": q, "cache": PUBLIC_STORE.stats(), "memory_mb": perf.memory_mb(),
            "uptime_min": round((time.time() - perf.STARTED) / 60), "slowest": perf.TIMINGS.snapshot(top=15),
            "storefront": perf.TIMINGS.snapshot(top=10, prefix="GET /api/b"),
            "checks": {"queue_lag_ok": (q.get("lag_seconds") or 0) < 600,
                       "worker_ok": q.get("mode") != "external" or q.get("worker_alive")}}


@admin_router.post("/disposable-refresh")
async def adm_disposable_refresh():
    return await refresh_disposable(force=True)


@admin_router.post("/learn")
async def adm_learn(user: dict = Depends(require_admin)):
    """Run the self-learning pass now (the cron runs it daily)."""
    return await learn_from_outcomes(force=True)


@admin_router.post("/suppression-sync")
async def adm_suppression_sync():
    """Copy AWS's suppression list into Contacts now (the cron does it daily)."""
    return await sync_aws_suppression(force=True)


@admin_router.post("/validate-pending")
async def adm_validate_pending():
    """'Check now': run the mailbox checks for waiting contacts immediately
    (the cron does the same every few minutes)."""
    kick_validation()
    return {"ok": True, "usage": await ses_usage()}


@admin_router.post("/reverify")
async def adm_reverify(list_id: Optional[str] = None):
    """Run verification again over saved contacts (e.g. after 6 months)."""
    flt = {"lists": list_id} if list_id else {}
    counts: dict = {"total": 0}
    async for chunk in dbutil.batches(db.mk_contacts, flt, {"_id": 0, "id": 1, "email": 1, "name": 1}, size=500):
        res = await verify_emails([c["email"] for c in chunk], autofix=False, names=[c.get("name", "") for c in chunk])
        await dbutil.bulk_set(db.mk_contacts, [(c["id"], _status_update(v)) for c, v in zip(chunk, res)])
        for k, v in _count(res).items():
            counts[k] = counts.get(k, 0) + v
    kick_validation()
    return {"counts": counts}


@admin_router.post("/sync")
async def adm_sync(user: dict = Depends(require_admin)):
    """Bring website people in as contacts: newsletter sign-ups (subscribed),
    customers with paid orders (contact only; consent only if they ticked it),
    OTP-confirmed accounts (proven)."""
    n = 0
    async for s in db.newsletter.find({}, {"_id": 0, "email": 1, "source": 1}):
        src = s.get("source") or "newsletter"
        # Waitlist sign-ups asked about one launch, not the newsletter: no consent.
        await upsert_contact(s["email"], source=src, email_consent=None if is_waitlist_source(src) else "subscribed")
        n += 1
    async for o in db.orders.find({"payment_status": "paid"}, {"_id": 0, "email": 1, "full_name": 1, "phone": 1, "email_marketing_optin": 1}):
        await upsert_contact(o.get("email", ""), name=o.get("full_name", ""), phone=o.get("phone", ""), source="checkout",
                             customer=True, email_consent="subscribed" if o.get("email_marketing_optin") else None)
        n += 1
    async for u in db.users.find({"email_verified": True}, {"_id": 0, "email": 1, "name": 1, "phone": 1, "wa_marketing_optin": 1}):
        await upsert_contact(u.get("email", ""), name=u.get("name", ""), phone=u.get("phone", ""), source="account",
                             proven=True, status="verified", reasons=["confirmed account"],
                             wa_consent="subscribed" if u.get("wa_marketing_optin") else None)
        n += 1
    await audit_log(db, "MARKETING_SYNC", email=user.get("email", ""), role=user.get("role", ""), meta={"rows": n})
    kick_validation()
    return {"ok": True, "processed": n, "contacts": await db.mk_contacts.count_documents({})}


# ---- contacts ----
# Addresses that can never be mailed: bounced / complained / blocked by SES
# (status "suppressed") and failed checks ("invalid"). Hidden from the
# Contacts list by default — NOT deleted: the record is what keeps them out
# of every future campaign and stops a later import from bringing them back.
DEAD_STATUSES = ["suppressed", "invalid"]


def _contact_filter(q: Optional[str] = None, status: Optional[str] = None, list_id: Optional[str] = None,
                    consent: Optional[str] = None, removed: str = "hide") -> dict:
    """One filter for the Contacts page AND its "select all N matching" bulk
    actions, so a bulk action can never hit more people than the admin saw.
    removed: hide (default) | show | only — bounced & invalid addresses."""
    flt: dict = {}
    if not status:
        if removed == "only":
            flt["email_status"] = {"$in": DEAD_STATUSES}
        elif removed != "show":
            flt["email_status"] = {"$nin": DEAD_STATUSES}
    if q and q.strip():
        rx = {"$regex": re.escape(q.strip()[:80]), "$options": "i"}
        flt["$or"] = [{"email": rx}, {"name": rx}, {"phone": rx}, {"source": rx}]
    if status:
        flt["email_status"] = status
    if list_id:
        flt["lists"] = list_id
    if consent:
        flt["email_consent.status"] = consent
    return flt


@admin_router.get("/contacts")
async def adm_contacts(q: Optional[str] = None, status: Optional[str] = None, list_id: Optional[str] = None,
                       consent: Optional[str] = None, page: int = 1, per: int = 50, removed: str = "hide"):
    flt = _contact_filter(q, status, list_id, consent, removed)
    per = min(max(per, 10), 200)
    total = await db.mk_contacts.count_documents(flt)
    rows = await db.mk_contacts.find(flt, {"_id": 0}).sort("created_at", -1).skip((max(page, 1) - 1) * per).limit(per).to_list(per)
    by_status = {r["_id"]: r["n"] async for r in db.mk_contacts.aggregate([{"$group": {"_id": "$email_status", "n": {"$sum": 1}}}])}
    by_consent = {r["_id"]: r["n"] async for r in db.mk_contacts.aggregate([{"$group": {"_id": "$email_consent.status", "n": {"$sum": 1}}}])}
    return {"rows": rows, "total": total, "page": page, "per": per, "by_status": by_status, "by_consent": by_consent,
            "removed_count": sum(n for k, n in by_status.items() if k in DEAD_STATUSES)}


# --------------------------------------------------------------- exports ---
EXPORT_MAX = 100_000


def _table_bytes(sheets: list, fmt: str) -> tuple[bytes, str, str]:
    """sheets: [(title, header, rows)] -> (bytes, media type, extension).
    CSV carries the first sheet only. Every cell goes through core.safe_cell
    so a name like "=HYPERLINK(...)" cannot run in Excel."""
    if fmt == "csv":
        import csv as _csv
        buf = io.StringIO()
        w = _csv.writer(buf)
        _, header, rows = sheets[0]
        w.writerow(header)
        for r in rows:
            w.writerow([core.safe_cell(v) for v in r])
        # BOM so Excel opens UTF-8 (Hindi names, ₹) correctly.
        return ("\ufeff" + buf.getvalue()).encode("utf-8"), "text/csv; charset=utf-8", "csv"
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    wb.remove(wb.active)
    for title, header, rows in sheets:
        ws = wb.create_sheet(title[:31])
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append([core.safe_cell(v) for v in r])
        ws.freeze_panes = "A2"
        for i, h in enumerate(header, start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = max(12, min(48, len(str(h)) + 6))
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"


def _download(data: bytes, media: str, name: str) -> Response:
    return Response(content=data, media_type=media, headers={
        "Content-Disposition": f'attachment; filename="{name}"', "Cache-Control": "private, no-store"})


def _slug(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", (s or "").lower()).strip("-")[:50] or "export"


@admin_router.get("/campaigns/{cid}/export")
async def adm_campaign_export(cid: str, format: str = "xlsx", user: dict = Depends(require_admin)):
    """Campaign report as a spreadsheet: Summary, every Recipient with what
    happened (delivered / opened / clicked / bounced / skipped and why), and
    Links. CSV = the recipients only. Personal data: logged."""
    fmt = "csv" if format == "csv" else "xlsx"
    c = await db.mk_campaigns.find_one({"id": cid}, {"_id": 0})
    if not c:
        raise HTTPException(status_code=404, detail="Campaign not found")
    stats = await refresh_campaign_stats(cid)
    names = {}
    rows = []
    sends = await db.mk_sends.find({"campaign_id": cid}, {"_id": 0}).limit(EXPORT_MAX).to_list(EXPORT_MAX)
    ids = list({x.get("contact_id") for x in sends if x.get("contact_id")})
    for k in range(0, len(ids), 1000):
        async for ct in db.mk_contacts.find({"id": {"$in": ids[k:k + 1000]}}, {"_id": 0, "id": 1, "name": 1}):
            names[ct["id"]] = ct.get("name", "")
    for x in sends:
        rows.append([x.get("to", ""), names.get(x.get("contact_id"), ""), core.send_outcome(x), x.get("status", ""),
                     (x.get("sent_at") or "")[:19], (x.get("delivered_at") or "")[:19], (x.get("opened_at") or "")[:19],
                     (x.get("clicked_at") or "")[:19], int(x.get("click_count") or 0), (x.get("unsubscribed_at") or "")[:19],
                     x.get("bounce_type") or "", x.get("bounce_sub") or "", x.get("error") or "", x.get("band") or "", x.get("score")])
    header = ["Email", "Name", "Outcome", "Send status", "Sent at (UTC)", "Delivered at", "Opened at", "Clicked at", "Clicks",
              "Unsubscribed at", "Bounce type", "Bounce detail", "Reason / error", "Address status when sent", "Score when sent"]
    links = c.get("links") or []
    clicks = [0] * len(links)
    for x in sends:
        for k in {y.get("i") for y in x.get("clicks") or []}:
            if isinstance(k, int) and 0 <= k < len(clicks):
                clicks[k] += 1
    rates = {"Delivery rate": core.rate(stats.get("delivered", 0), stats.get("sent", 0)),
             "Open rate": core.rate(stats.get("opened", 0), stats.get("delivered", 0) or stats.get("sent", 0)),
             "Click rate": core.rate(stats.get("clicked", 0), stats.get("delivered", 0) or stats.get("sent", 0)),
             "Bounce rate": core.rate(stats.get("bounced", 0), stats.get("sent", 0))}
    summary = [["Campaign", c.get("name", "")], ["Subject", c.get("subject", "")], ["Status", c.get("status", "")],
               ["Started", (c.get("started_at") or "")[:19]], ["Finished", (c.get("finished_at") or "")[:19]]]
    summary += [[k.replace("_", " ").capitalize(), v] for k, v in stats.items() if isinstance(v, (int, float))]
    summary += [[k, f"{round(v * 100, 1)}%" if v is not None else ""] for k, v in rates.items()]
    data, media, ext = _table_bytes([("Recipients", header, rows), ("Summary", ["Metric", "Value"], summary),
                                     ("Links", ["Link", "People who clicked"], [[u, n] for u, n in zip(links, clicks)])], fmt)
    await audit_log(db, "MARKETING_EXPORT", email=user.get("email", ""), role=user.get("role", ""),
                    meta={"what": "campaign report", "campaign": cid, "rows": len(rows), "format": ext})
    return _download(data, media, f"campaign-{_slug(c.get('name'))}-{_now():%Y-%m-%d}.{ext}")


CONTACT_EXPORT_HEADER = ["Email", "Name", "Phone", "Address status", "Score", "Why", "Email consent", "Consent source",
                         "Consent date", "WhatsApp consent", "Suppressed (why)", "Clean-up group", "Lists", "Source",
                         "Emails sent", "Opened", "Clicked", "Added"]


async def _contact_export_rows(flt: dict) -> tuple[list, list]:
    """(header, rows) for a contacts export — one shape for Contacts and for
    Lists & segments, so both files open with the same columns."""
    list_names = {x["id"]: x.get("name", "") async for x in db.mk_lists.find({}, {"_id": 0, "id": 1, "name": 1})}
    rows = []
    async for page in dbutil.batches(db.mk_contacts, flt, {"_id": 0}, size=1000):
        for ct in page:
            ec, wc = ct.get("email_consent") or {}, ct.get("wa_consent") or {}
            st = ct.get("stats") or {}
            rows.append([ct.get("email", ""), ct.get("name", ""), ct.get("phone", ""), ct.get("email_status", ""), ct.get("score"),
                         "; ".join((ct.get("reasons") or [])[:3]), ec.get("status", ""), ec.get("source", ""), (ec.get("at") or "")[:19],
                         wc.get("status", ""), ct.get("suppressed") or "", core.CLEANUP_KINDS.get(core.cleanup_kind(ct), ""),
                         ", ".join(list_names.get(i, i) for i in ct.get("lists") or []), ct.get("source", ""),
                         int(st.get("sent") or 0), int(st.get("opened") or 0), int(st.get("clicked") or 0), (ct.get("created_at") or "")[:19]])
            if len(rows) >= EXPORT_MAX:
                break
        if len(rows) >= EXPORT_MAX:
            break
    return CONTACT_EXPORT_HEADER, rows


@admin_router.get("/contacts/export")
async def adm_contacts_export(q: Optional[str] = None, status: Optional[str] = None, list_id: Optional[str] = None,
                              consent: Optional[str] = None, removed: str = "hide", format: str = "xlsx",
                              user: dict = Depends(require_admin)):
    """The Contacts list as shown (same filters), for fixing records outside
    the site or keeping a copy before a clean-up. Personal data: logged."""
    fmt = "csv" if format == "csv" else "xlsx"
    flt = _contact_filter(q, status, list_id, consent, removed)
    header, rows = await _contact_export_rows(flt)
    data, media, ext = _table_bytes([("Contacts", header, rows)], fmt)
    await audit_log(db, "MARKETING_EXPORT", email=user.get("email", ""), role=user.get("role", ""),
                    meta={"what": "contacts", "rows": len(rows), "format": ext, "filters": {"q": q, "status": status, "list": list_id,
                                                                                           "consent": consent, "removed": removed}})
    return _download(data, media, f"contacts-{removed if not status else status}-{_now():%Y-%m-%d}.{ext}")


# -------------------------------------------------------------- clean-up ---
class CleanupBody(BaseModel):
    kinds: list = ["bounced", "blocked", "invalid"]
    campaign_id: Optional[str] = None   # only people this campaign was sent to
    dry_run: bool = True
    expected: Optional[int] = None      # the count the admin confirmed


async def _cleanup_candidates(kinds: set, campaign_id: Optional[str]) -> list:
    flt: dict = {"$or": [{"suppressed": {"$nin": [None, ""]}}, {"email_status": "invalid"},
                         {"email_consent.status": {"$in": ["unsubscribed", "complained"]}}]}
    if campaign_id:
        ids = [x["contact_id"] async for x in db.mk_sends.find({"campaign_id": campaign_id}, {"_id": 0, "contact_id": 1})]
        flt = {"$and": [flt, {"id": {"$in": ids}}]}
    out = []
    async for ct in db.mk_contacts.find(flt, {"_id": 0, "id": 1, "email": 1, "email_norm": 1, "suppressed": 1,
                                              "email_status": 1, "email_consent": 1}):
        k = core.cleanup_kind(ct)
        if k in kinds:
            out.append({**ct, "kind": k})
    return out


@admin_router.post("/contacts/cleanup")
async def adm_contacts_cleanup(body: CleanupBody, user: dict = Depends(require_admin)):
    """Fix the database after sending: remove bounced / blocked / invalid
    contacts (and, if chosen, complaints and unsubscribes).

    Removed contacts leave Contacts and every list, but a one-way hash of the
    address is kept with the reason (db.mk_removed) — the same idea as an
    erase. Without it, the next order sync or an old spreadsheet would add
    the address back as new and it would be mailed (and bounce) again. Their
    own new sign-up still works. Send history (reports) is kept.

    dry_run=True counts by kind (any marketing admin); removing needs a
    superadmin and the confirmed count must still match."""
    from rbac import is_superadmin
    kinds = {k for k in body.kinds or [] if k in core.CLEANUP_KINDS}
    if not kinds:
        raise HTTPException(status_code=400, detail="Choose what to remove")
    cands = await _cleanup_candidates(kinds, body.campaign_id)
    by_kind: dict = {}
    for ct in cands:
        by_kind[ct["kind"]] = by_kind.get(ct["kind"], 0) + 1
    if body.dry_run:
        return {"total": len(cands), "by_kind": by_kind, "labels": core.CLEANUP_KINDS,
                "sample": [{"email": ct["email"], "kind": ct["kind"]} for ct in cands[:20]]}
    if not is_superadmin(user.get("role")):
        raise HTTPException(status_code=403, detail="Only a superadmin can remove contacts")
    if body.expected is not None and body.expected != len(cands):
        raise HTTPException(status_code=409, detail=f"The count changed to {len(cands)} — check again")
    done = 0
    for ct in cands:
        if not ct.get("email_norm"):
            continue
        await db.mk_removed.update_one({"h": _erase_hash(ct["email_norm"])},
                                       {"$set": {"reason": ct["kind"], "at": _iso(), "by": user.get("email", "")}}, upsert=True)
        r = await db.mk_contacts.delete_one({"id": ct["id"]})
        done += r.deleted_count
    await audit_log(db, "MARKETING_CLEANUP", email=user.get("email", ""), role=user.get("role", ""),
                    meta={"removed": done, "by_kind": by_kind, "campaign": body.campaign_id})
    return {"ok": True, "removed": done, "by_kind": by_kind}


class ContactPatch(BaseModel):
    name: Optional[str] = None
    phone: Optional[str] = None
    email_consent: Optional[str] = None   # unsubscribed only (subscribing needs the person's own action)
    wa_consent: Optional[str] = None
    add_list: Optional[str] = None
    remove_list: Optional[str] = None


@admin_router.patch("/contacts/{cid}")
async def adm_patch_contact(cid: str, body: ContactPatch, user: dict = Depends(require_admin)):
    c = await db.mk_contacts.find_one({"id": cid}, {"_id": 0})
    if not c:
        raise HTTPException(status_code=404, detail="Contact not found")
    upd: dict = {}
    ops: dict = {}
    if body.name is not None:
        upd["name"] = body.name.strip()[:120]
    if body.phone is not None:
        upd["phone"] = re.sub(r"[^\d+ ]", "", body.phone)[:20]
    for fld in ("email_consent", "wa_consent"):
        v = getattr(body, fld)
        if v is not None:
            if v != "unsubscribed":
                raise HTTPException(status_code=400, detail="Admin can only unsubscribe someone; subscribing needs their own opt-in")
            upd[fld] = _consent("unsubscribed", f"admin ({user.get('email', '')})")
    if body.add_list:
        ops["$addToSet"] = {"lists": body.add_list}
    if body.remove_list:
        ops["$pull"] = {"lists": body.remove_list}
    if upd:
        ops["$set"] = upd
    if ops:
        await db.mk_contacts.update_one({"id": cid}, ops)
    return await db.mk_contacts.find_one({"id": cid}, {"_id": 0})


@admin_router.delete("/contacts/{cid}")
async def adm_delete_contact(cid: str, erase: bool = False, user: dict = Depends(require_admin)):
    """Superadmin (DELETE). Two different things:
    delete (default) — remove the record; the address can be imported again.
    erase=true       — privacy request: also keep a hash so an old export can
                       never re-add them by accident (only their own fresh
                       opt-in, or a superadmin's deliberate re-add, can)."""
    c = await db.mk_contacts.find_one({"id": cid}, {"_id": 0, "email_norm": 1})
    if not c:
        raise HTTPException(status_code=404, detail="Contact not found")
    await db.mk_contacts.delete_one({"id": cid})
    if erase:
        await db.mk_erased.update_one({"h": _erase_hash(c["email_norm"])}, {"$set": {"at": _iso()}}, upsert=True)
    await audit_log(db, "MARKETING_CONTACT_ERASED" if erase else "MARKETING_CONTACT_DELETED", email=user.get("email", ""),
                    role=user.get("role", ""), meta={"contact": cid})
    return {"ok": True, "erased": erase}


class ContactBulk(BaseModel):
    action: str                         # unsubscribe | erase
    ids: Optional[list[str]] = None     # the ticked rows, or…
    all_matching: Optional[dict] = None  # {q, status, list_id, consent}: everyone the filter shows
    expected: int                       # how many the admin confirmed — must still match


BULK_MAX = 50000


@admin_router.post("/contacts/bulk")
async def adm_contacts_bulk(body: ContactBulk, user: dict = Depends(require_admin)):
    """Bulk unsubscribe (admin) or erase (superadmin — the same rule as the
    single erase, which is a DELETE). Unsubscribe is reversible only by the
    person opting in again; erase keeps just a hash so they can't be re-added."""
    from rbac import is_superadmin
    if body.action not in ("unsubscribe", "delete", "erase"):
        raise HTTPException(status_code=400, detail="Unknown action")
    if body.action in ("delete", "erase") and not is_superadmin(user.get("role")):
        raise HTTPException(status_code=403, detail="Only a superadmin can delete or erase contacts")
    if body.ids:
        flt = {"id": {"$in": [str(x) for x in body.ids][:BULK_MAX]}}
    elif body.all_matching is not None:
        m = body.all_matching
        flt = _contact_filter(m.get("q"), m.get("status"), m.get("list_id"), m.get("consent"), m.get("removed") or "hide")
    else:
        raise HTTPException(status_code=400, detail="Nothing selected")
    n = await db.mk_contacts.count_documents(flt)
    if n > BULK_MAX:
        raise HTTPException(status_code=400, detail=f"{n} contacts — narrow the filter (max {BULK_MAX} at a time)")
    if n != body.expected:
        raise HTTPException(status_code=409, detail=f"The selection changed to {n} — check it again")
    done = 0
    if body.action == "unsubscribe":
        r = await db.mk_contacts.update_many({**flt, "email_consent.status": {"$nin": ["unsubscribed", "complained"]}},
                                             {"$set": {"email_consent": _consent("unsubscribed", "admin (bulk)", user.get("email", ""))}})
        done = r.modified_count
    elif body.action == "delete":
        r = await db.mk_contacts.delete_many(flt)  # can be imported again later
        done = r.deleted_count
    else:
        async for c in db.mk_contacts.find(flt, {"_id": 0, "id": 1, "email_norm": 1}):
            await db.mk_erased.update_one({"h": _erase_hash(c["email_norm"])}, {"$set": {"at": _iso()}}, upsert=True)
            await db.mk_contacts.delete_one({"id": c["id"]})
            done += 1
    await audit_log(db, f"MARKETING_CONTACTS_BULK_{body.action.upper()}", email=user.get("email", ""), role=user.get("role", ""),
                    meta={"count": done, "by_filter": body.all_matching is not None})
    return {"ok": True, "done": done}


# ---- lists / segments ----
class ListBody(BaseModel):
    name: str
    kind: str = "static"          # static | segment
    rules: Optional[list] = None  # segment rules (see audience_filter)


@admin_router.get("/lists")
async def adm_lists():
    lists = await db.mk_lists.find({}, {"_id": 0}).sort("created_at", -1).to_list(500)
    pol = _policy(await get_settings(), {})
    for lst in lists:
        if lst.get("kind") == "segment":
            flt = await audience_filter({"channel": "email", "audience": {"rules": lst.get("rules") or []}})
            lst["count"] = await db.mk_contacts.count_documents(flt)
        else:
            lst["count"] = await db.mk_contacts.count_documents({"lists": lst["id"]})
            # What a campaign to this list would actually do, by the CURRENT
            # risky-address rules in Settings — not only verified/valid (which
            # read "2 can be emailed" while the rules were sending 475).
            confirmed, by_action = 0, {"send": 0, "tail": 0, "skip": 0}
            async for r in db.mk_contacts.aggregate([
                    {"$match": {"lists": lst["id"], "email_consent.status": "subscribed", "suppressed": {"$in": [None, ""]}}},
                    {"$group": {"_id": {"s": "$email_status", "k": "$risk_kind"}, "n": {"$sum": 1}}}]):
                stt, kind = (r["_id"] or {}).get("s"), (r["_id"] or {}).get("k") or "unconfirmed"
                if stt in ("verified", "valid"):
                    confirmed += r["n"]
                elif stt == "risky":
                    by_action[pol.get(kind, "skip")] += r["n"]
            lst["sendable"] = confirmed + by_action["send"] + by_action["tail"]
            lst["confirmed"] = confirmed
            lst["risky"] = by_action
    return lists


async def _list_member_filter(lid: str, q: Optional[str] = None) -> tuple[dict, dict]:
    """(list, Mongo filter for its members). A static list is whoever was
    added to it; a segment is whoever its rules match — the SAME filter the
    list's count uses, so "View" shows exactly the number on the row."""
    lst = await db.mk_lists.find_one({"id": lid}, {"_id": 0})
    if not lst:
        raise HTTPException(status_code=404, detail="List not found")
    if lst.get("kind") == "segment":
        flt = await audience_filter({"channel": "email", "audience": {"rules": lst.get("rules") or []}})
    else:
        flt = {"lists": lid}
    if q and q.strip():
        rx = {"$regex": re.escape(q.strip()[:80]), "$options": "i"}
        flt = {"$and": [flt, {"$or": [{"email": rx}, {"name": rx}, {"phone": rx}]}]}
    return lst, flt


@admin_router.get("/lists/{lid}/members")
async def adm_list_members(lid: str, q: Optional[str] = None, page: int = 1, per: int = 50):
    """Admin → Marketing → Lists & segments → View: who is in it, newest first."""
    lst, flt = await _list_member_filter(lid, q)
    per = min(max(per, 10), 200)
    total = await db.mk_contacts.count_documents(flt)
    rows = await db.mk_contacts.find(flt, {"_id": 0, "id": 1, "email": 1, "name": 1, "phone": 1, "email_status": 1,
                                           "email_consent.status": 1, "suppressed": 1, "source": 1, "created_at": 1,
                                           "customer": 1, "stats": 1}) \
        .sort("created_at", -1).skip((max(page, 1) - 1) * per).limit(per).to_list(per)
    return {"list": {k: lst.get(k) for k in ("id", "name", "kind", "rules", "created_at", "source")},
            "total": total, "page": max(page, 1), "per": per, "rows": rows}


@admin_router.get("/lists/{lid}/export")
async def adm_list_export(lid: str, q: Optional[str] = None, format: str = "xlsx", user: dict = Depends(require_admin)):
    """One list or segment as Excel / CSV — same columns as the Contacts
    export. Personal data: every export is logged."""
    fmt = "csv" if format == "csv" else "xlsx"
    lst, flt = await _list_member_filter(lid, q)
    header, rows = await _contact_export_rows(flt)
    data, media, ext = _table_bytes([((lst.get("name") or "List")[:31], header, rows)], fmt)
    await audit_log(db, "MARKETING_EXPORT", email=user.get("email", ""), role=user.get("role", ""),
                    meta={"what": "list", "list": lid, "kind": lst.get("kind"), "rows": len(rows), "format": ext})
    slug = re.sub(r"[^A-Za-z0-9-]+", "-", lst.get("name") or "list").strip("-")[:50] or "list"
    return _download(data, media, f"{slug}-{_now():%Y-%m-%d}.{ext}")


@admin_router.post("/lists")
async def adm_create_list(body: ListBody, user: dict = Depends(require_admin)):
    if body.kind not in ("static", "segment") or not body.name.strip():
        raise HTTPException(status_code=400, detail="Name and kind (static / segment) required")
    doc = {"id": str(uuid.uuid4()), "name": body.name.strip()[:80], "kind": body.kind, "rules": body.rules or [],
           "created_at": _iso(), "created_by": user.get("email", "")}
    await db.mk_lists.insert_one(dict(doc))
    return doc


class FromUsersBody(BaseModel):
    user_ids: Optional[list] = None   # specific people ticked in Admin → Users
    all_customers: bool = False       # everyone with an account (customers) + every guest buyer
    name: Optional[str] = None
    template: Optional[str] = None    # "for-you" | "blank" | None (list only)
    ask_permission: bool = False      # one permission email for those who never decided


@admin_router.post("/lists/from-users")
async def adm_list_from_users(body: FromUsersBody, user: dict = Depends(require_admin)):
    """Admin → Users → "Email selected" / "Email all customers".

    Makes a static marketing list of those people (creating contacts where
    needed) and, if asked, a draft campaign aimed at it. CONSENT IS NOT
    CHANGED: a campaign still reaches only people who agreed to marketing;
    unsubscribed / complained / bounced people are never mailed. Those who
    never decided can, if ticked, be sent ONE permission email (the re-confirm
    campaign) — the lawful way to reach them under the DPDP Act."""
    people: dict = {}
    if body.all_customers:
        async for u in db.users.find({"role": {"$in": [None, "customer"]}}, {"_id": 0, "email": 1, "name": 1, "phone": 1}):
            if u.get("email"):
                people.setdefault(core.normalise_email(u["email"]), u)
        async for o in db.orders.find({}, {"_id": 0, "email": 1, "full_name": 1, "phone": 1}):
            if o.get("email"):
                people.setdefault(core.normalise_email(o["email"]), {"email": o["email"], "name": o.get("full_name", ""), "phone": o.get("phone", "")})
    else:
        ids = [str(i) for i in (body.user_ids or [])][:5000]
        if not ids:
            raise HTTPException(status_code=400, detail="Tick at least one person, or choose all customers")
        async for u in db.users.find({"id": {"$in": ids}}, {"_id": 0, "email": 1, "name": 1, "phone": 1}):
            if u.get("email"):
                people.setdefault(core.normalise_email(u["email"]), u)
    people.pop("", None)
    name = (body.name or ("All customers" if body.all_customers else f"Selected customers · {len(people)}")).strip()[:80]
    name = f"{name} · {_now().strftime('%d %b %Y')}"
    lst = {"id": str(uuid.uuid4()), "name": name, "kind": "static", "rules": [], "created_at": _iso(),
           "created_by": user.get("email", ""), "source": "admin-users"}
    await db.mk_lists.insert_one(dict(lst))
    counts = {"people": len(people), "added": 0, "subscribed": 0, "no_decision": 0, "opted_out": 0, "blocked": 0, "asked_permission": 0}
    for norm, u in people.items():
        c = await upsert_contact(u["email"], name=u.get("name") or "", phone=u.get("phone") or "", source="website account",
                                 lists=[lst["id"]], customer=True, by=user.get("email", ""))
        if not c:
            counts["blocked"] += 1   # erased at their request / removed in a clean-up: never re-added
            continue
        counts["added"] += 1
        cs = (c.get("email_consent") or {}).get("status")
        if c.get("suppressed") or cs in ("unsubscribed", "complained", "lapsed"):
            counts["opted_out"] += 1
        elif cs == "subscribed":
            counts["subscribed"] += 1
        else:
            counts["no_decision"] += 1
            if body.ask_permission:
                r = await db.mk_contacts.update_one(
                    {"id": c["id"], "email_consent.status": {"$nin": ["subscribed", "unsubscribed", "complained", "lapsed"]}},
                    {"$set": {"permission_pending": True, "permission_list": lst["id"], "permission_since": _iso()}})
                counts["asked_permission"] += r.modified_count
    out = {"list": lst, "counts": counts, "campaigns": []}
    tmpl = (body.template or "").strip()
    if tmpl in ("for-you", "blank"):
        t = next((x for x in core.STARTER_TEMPLATES if x["key"] == "for-you"), None) if tmpl == "for-you" else None
        blocks = _clean_blocks((t or {}).get("blocks") or [])
        doc = {"id": str(uuid.uuid4()), "name": (t["name"] if t else "Email to customers") + f" — {lst['name']}", "channel": "email",
               "kind": "standard", "reconfirm_days": None, "subject": (t or {}).get("subject", ""), "preheader": (t or {}).get("preheader", ""),
               "blocks": blocks, "links": core.collect_links(blocks), "audience": {"list_ids": [lst["id"]], "rules": []},
               "wa_template": "", "wa_language": "en", "wa_variables": [], "status": "draft", "created_at": _iso(),
               "created_by": user.get("email", ""), "stats": {}}
        await db.mk_campaigns.insert_one(dict(doc))
        out["campaigns"].append({"id": doc["id"], "name": doc["name"], "kind": "standard"})
    if body.ask_permission and counts["asked_permission"]:
        blocks = [dict(b) for b in RECONFIRM_BLOCKS]
        doc = {"id": str(uuid.uuid4()), "name": f"Permission request — {lst['name']}", "channel": "email", "kind": "reconfirm",
               "reconfirm_days": 30, "subject": "Do you still want to hear from Oakbridge?", "preheader": "",
               "blocks": blocks, "links": core.collect_links(blocks), "audience": {"list_ids": [lst["id"]], "rules": []},
               "wa_template": "", "wa_language": "en", "wa_variables": [], "status": "draft", "created_at": _iso(),
               "created_by": user.get("email", ""), "stats": {}}
        await db.mk_campaigns.insert_one(dict(doc))
        out["campaigns"].append({"id": doc["id"], "name": doc["name"], "kind": "reconfirm"})
    await audit_log(db, "MARKETING_LIST_FROM_USERS", email=user.get("email", ""), role=user.get("role", ""),
                    meta={"list": lst["id"], "all": body.all_customers, **counts})
    return out


@admin_router.delete("/lists/{lid}")
async def adm_delete_list(lid: str):
    await db.mk_lists.delete_one({"id": lid})
    await db.mk_contacts.update_many({"lists": lid}, {"$pull": {"lists": lid}})
    return {"ok": True}


class IdsBody(BaseModel):
    ids: list[str]


@admin_router.post("/lists/bulk-delete")
async def adm_lists_bulk_delete(body: IdsBody, user: dict = Depends(require_superadmin)):
    """Delete several lists / segments. The contacts in them stay."""
    ids = [str(x) for x in body.ids][:500]
    r = await db.mk_lists.delete_many({"id": {"$in": ids}})
    await db.mk_contacts.update_many({"lists": {"$in": ids}}, {"$pull": {"lists": {"$in": ids}}})
    await audit_log(db, "MARKETING_LISTS_DELETED", email=user.get("email", ""), role=user.get("role", ""), meta={"count": r.deleted_count})
    return {"ok": True, "deleted": r.deleted_count}


# ---- campaigns ----
class CampaignBody(BaseModel):
    name: Optional[str] = None
    channel: Optional[str] = None          # email | whatsapp
    subject: Optional[str] = None
    preheader: Optional[str] = None
    blocks: Optional[list] = None
    audience: Optional[dict] = None        # {list_ids, exclude_list_ids, rules, include_risky}
    wa_template: Optional[str] = None
    wa_language: Optional[str] = None
    wa_variables: Optional[list] = None    # [{source: first_name|name|static, value}]
    kind: Optional[str] = None             # standard | reconfirm
    reconfirm_days: Optional[int] = None


def _clean_blocks(blocks: list) -> list:
    ok_types = {"heading", "text", "image", "button", "book", "divider", "spacer", "for_you"}
    out = []
    for b in (blocks or [])[:60]:
        if not isinstance(b, dict) or b.get("type") not in ok_types:
            continue
        c = {k: (str(v)[:5000] if isinstance(v, str) else v) for k, v in b.items()
             if k in ("type", "text", "label", "url", "src", "alt", "title", "author", "cover", "price", "book_id", "count")}
        if c.get("type") == "for_you":
            try:
                c["count"] = max(1, min(int(c.get("count") or 3), 6))
            except (TypeError, ValueError):
                c["count"] = 3
        for k in ("url", "src", "cover"):
            if c.get(k) and not re.match(r"^https?://", str(c[k]), re.I):
                c.pop(k)   # only web links / images
        out.append(c)
    return out


@admin_router.get("/campaigns")
async def adm_campaigns():
    return await db.mk_campaigns.find({"deleted": {"$ne": True}}, {"_id": 0, "blocks": 0}).sort("created_at", -1).to_list(500)


@admin_router.post("/campaigns/bulk-delete")
async def adm_campaigns_bulk_delete(body: IdsBody, user: dict = Depends(require_superadmin)):
    """Drafts (nothing sent) are deleted outright. Anything that reached a
    mailbox is ARCHIVED instead — hidden everywhere, but its send records are
    kept so the unsubscribe links in those emails keep working (a legal
    requirement) and tracking links still land on the site. A campaign that
    is preparing / sending must be cancelled first."""
    deleted, archived, refused = 0, 0, []
    for cid in [str(x) for x in body.ids][:500]:
        c = await db.mk_campaigns.find_one({"id": cid, "deleted": {"$ne": True}}, {"_id": 0, "id": 1, "status": 1, "name": 1})
        if not c:
            continue
        if c["status"] in ("preparing", "sending"):
            refused.append({"id": cid, "name": c.get("name", ""), "reason": "still sending — cancel it first"})
            continue
        # scheduled / paused: nothing more may go out
        await db.mk_sends.update_many({"campaign_id": cid, "status": "queued"}, {"$set": {"status": "skipped", "error": "campaign deleted"}})
        reached = await db.mk_sends.count_documents({"campaign_id": cid, "status": {"$in": ["sent", "bounced", "complained", "blocked"]}})
        if reached:
            await db.mk_campaigns.update_one({"id": cid}, {"$set": {"deleted": True, "deleted_at": _iso(), "deleted_by": user.get("email", ""),
                                                                    "status": c["status"] if c["status"] in ("sent", "cancelled") else "cancelled"}})
            archived += 1
        else:
            await db.mk_sends.delete_many({"campaign_id": cid})
            await db.mk_campaigns.delete_one({"id": cid})
            deleted += 1
    await audit_log(db, "MARKETING_CAMPAIGNS_DELETED", email=user.get("email", ""), role=user.get("role", ""),
                    meta={"deleted": deleted, "archived": archived, "refused": len(refused)})
    return {"ok": True, "deleted": deleted, "archived": archived, "refused": refused}


@admin_router.post("/campaigns")
async def adm_new_campaign(body: CampaignBody, user: dict = Depends(require_admin)):
    ch = body.channel or "email"
    if ch not in ("email", "whatsapp"):
        raise HTTPException(status_code=400, detail="channel must be email or whatsapp")
    kind = "reconfirm" if body.kind == "reconfirm" and ch == "email" else "standard"
    blocks = [dict(b) for b in RECONFIRM_BLOCKS] if kind == "reconfirm" else []
    doc = {"id": str(uuid.uuid4()), "name": (body.name or "Untitled campaign").strip()[:120], "channel": ch, "kind": kind,
           "reconfirm_days": max(7, min(int(body.reconfirm_days or 30), 90)) if kind == "reconfirm" else None,
           "subject": "Do you still want to hear from Oakbridge?" if kind == "reconfirm" else "", "preheader": "",
           "blocks": blocks, "links": core.collect_links(blocks), "audience": {"list_ids": [], "rules": []},
           "wa_template": "", "wa_language": "en", "wa_variables": [], "status": "draft", "created_at": _iso(),
           "created_by": user.get("email", ""), "stats": {}}
    await db.mk_campaigns.insert_one(dict(doc))
    return doc


@admin_router.get("/campaigns/{cid}")
async def adm_get_campaign(cid: str):
    c = await db.mk_campaigns.find_one({"id": cid}, {"_id": 0})
    if not c:
        raise HTTPException(status_code=404, detail="Campaign not found")
    return c


@admin_router.patch("/campaigns/{cid}")
async def adm_patch_campaign(cid: str, body: CampaignBody, user: dict = Depends(require_admin)):
    c = await db.mk_campaigns.find_one({"id": cid}, {"_id": 0, "status": 1, "edit_logged_at": 1, "edit_logged_by": 1})
    if not c:
        raise HTTPException(status_code=404, detail="Campaign not found")
    if c["status"] not in ("draft", "scheduled"):
        raise HTTPException(status_code=409, detail="Only a draft can be edited — duplicate it to make changes")
    upd = {k: v for k, v in body.model_dump().items() if v is not None and k not in ("channel", "kind")}
    if "reconfirm_days" in upd:
        upd["reconfirm_days"] = max(7, min(int(upd["reconfirm_days"]), 90))
    if "audience" in upd and "statuses" in (upd["audience"] or {}):
        upd["audience"]["statuses"] = sorted(core.selected_statuses(upd["audience"]))
    if "blocks" in upd:
        upd["blocks"] = _clean_blocks(upd["blocks"])
        upd["links"] = core.collect_links(upd["blocks"])
    for k in ("name", "subject", "preheader", "wa_template"):
        if k in upd:
            upd[k] = str(upd[k]).strip()[:300]
    changed = sorted(k for k in upd if k != "links")
    upd["updated_at"] = _iso()
    upd["updated_by"] = user.get("email", "")
    # Campaign activity log: the editor saves often, so one "edited" line per
    # person per 10 minutes, naming what was changed — not one per keystroke.
    me = user.get("email", "")
    quiet = c.get("edit_logged_by") == me and (c.get("edit_logged_at") or "") >= (_now() - timedelta(minutes=10)).isoformat()
    if not quiet:
        upd.update(edit_logged_at=upd["updated_at"], edit_logged_by=me)
    await db.mk_campaigns.update_one({"id": cid}, {"$set": upd})
    if not quiet:
        await audit_log(db, "MARKETING_CAMPAIGN_EDITED", email=me, role=user.get("role", ""),
                        meta={"campaign": cid, "fields": ", ".join(changed)[:200]})
    return await db.mk_campaigns.find_one({"id": cid}, {"_id": 0})


class RenameBody(BaseModel):
    name: str


@admin_router.patch("/campaigns/{cid}/name")
async def adm_rename_campaign(cid: str, body: RenameBody, user: dict = Depends(require_admin)):
    """Rename any campaign, whatever its status. The name is internal (never
    shown to recipients), so unlike content it is safe to change after sending.
    PATCH on its own path: POST /campaigns/{cid}/{action} would swallow it."""
    name = (body.name or "").strip()[:120]
    if not name:
        raise HTTPException(status_code=400, detail="Give the campaign a name")
    r = await db.mk_campaigns.update_one({"id": cid, "deleted": {"$ne": True}}, {"$set": {"name": name, "updated_at": _iso()}})
    if not r.matched_count:
        raise HTTPException(status_code=404, detail="Campaign not found")
    await audit_log(db, "MARKETING_CAMPAIGN_RENAMED", email=user.get("email", ""), role=user.get("role", ""), meta={"campaign": cid})
    return {"ok": True, "name": name}


@admin_router.post("/campaigns/{cid}/duplicate")
async def adm_dup_campaign(cid: str, user: dict = Depends(require_admin)):
    c = await db.mk_campaigns.find_one({"id": cid}, {"_id": 0})
    if not c:
        raise HTTPException(status_code=404, detail="Campaign not found")
    keep = {k: c.get(k) for k in ("channel", "subject", "preheader", "blocks", "links", "audience", "wa_template", "wa_language", "wa_variables")}
    doc = {**keep, "id": str(uuid.uuid4()), "name": f"{c['name']} (copy)"[:120], "status": "draft", "created_at": _iso(),
           "created_by": user.get("email", ""), "stats": {}}
    await db.mk_campaigns.insert_one(dict(doc))
    return doc


@admin_router.get("/campaigns/{cid}/audience")
async def adm_audience(cid: str):
    """Who it would go to, by verification status — shown before sending."""
    await heal_proven()
    c = await db.mk_campaigns.find_one({"id": cid}, {"_id": 0})
    if not c:
        raise HTTPException(status_code=404, detail="Campaign not found")
    flt = await audience_filter(c)
    # Counted in the database (aggregation), not by loading every contact.
    by: dict = {}
    kinds_raw: dict = {}
    async for r in db.mk_contacts.aggregate([{"$match": flt}, {"$group": {
            "_id": {"s": "$email_status", "k": "$risk_kind"}, "n": {"$sum": 1}}}]):
        stt = (r["_id"] or {}).get("s") or "unknown"
        by[stt] = by.get(stt, 0) + r["n"]
        if stt == "risky":
            k = (r["_id"] or {}).get("k") or "unconfirmed"
            kinds_raw[k] = kinds_raw.get(k, 0) + r["n"]
    total = sum(by.values())
    st = await get_settings()
    rate = float(st.get("usd_inr") or 88)
    out = {"total": total, "by_status": by}
    if c["channel"] == "whatsapp":
        out["estimated_cost_inr"] = round(total * float(st["wa_rate_marketing"]) * 1.18, 2)
        return out
    out["estimated_cost_inr"] = round(total / 1000 * 0.10 * rate, 2)
    use = await ses_usage()
    ses_on = use["enabled"] and use["supported"]
    needs = 0
    if ses_on:
        unproven_flt = {"$and": [flt, {"email_status": {"$in": ["valid", "risky", "unknown"]}}]}
        async for page in dbutil.batches(db.mk_contacts, unproven_flt, {"_id": 0, "id": 1, "email_norm": 1}, size=1000):
            norms = [r["email_norm"] for r in page if r.get("email_norm")]
            cached = await _ses_cached(norms)
            needs += sum(1 for n in norms if n not in cached)
    pol = _policy(st, c)
    kinds = {k: {"n": n, "action": pol.get(k, "skip")} for k, n in kinds_raw.items()}
    out.update(needs_check=needs, check_cost_inr=round(needs * use["per_check_inr"], 2), checks_left=use["left"],
               ses_on=ses_on, risky_by_kind=kinds, selected=sorted(core.selected_statuses(c.get("audience"))))
    return out


class TestBody(BaseModel):
    to: str


@admin_router.post("/campaigns/{cid}/test")
async def adm_test_campaign(cid: str, body: TestBody, user: dict = Depends(require_admin)):
    c = await db.mk_campaigns.find_one({"id": cid}, {"_id": 0})
    if not c:
        raise HTTPException(status_code=404, detail="Campaign not found")
    st = await get_settings()
    if c["channel"] == "email":
        if not core.check_syntax(body.to)["ok"]:
            raise HTTPException(status_code=400, detail="Not an email address")
        contact = {"name": user.get("name") or "Test", "email": body.to}
        await attach_picks(contact, c)   # the test address's own picks (or the newest, if it has no history)
        html_doc, text_doc = core.render_email(c.get("blocks") or [], brand={"name": "Oakbridge Publishing", "logo": st.get("logo"), "footer": st.get("footer")},
                                               contact=contact, link_for=lambda u: core.add_utm(u, cid), unsubscribe_url=f"{SITE}/unsubscribe",
                                               preheader=c.get("preheader") or "")
        raw = build_mime(st=st, to=body.to, subject=f"[TEST] {core.personalise(c.get('subject') or '(no subject)', contact)}",
                         html_doc=html_doc, text_doc=text_doc, unsub_post="")
        try:
            mid = await asyncio.to_thread(_send_raw, raw)
        except Exception as e:  # noqa: BLE001
            raise HTTPException(status_code=502, detail=f"SES refused it: {str(e)[:250]}")
        await audit_log(db, "MARKETING_CAMPAIGN_TEST", email=user.get("email", ""), role=user.get("role", ""),
                        meta={"campaign": cid, "channel": "email", "to": body.to[:120]})
        return {"ok": True, "message_id": mid}
    from interakt import send_campaign_message
    rec = await send_campaign_message(template=c.get("wa_template"), language=c.get("wa_language") or "en",
                                      values=[v.get("value") if v.get("source") == "static" else "Test" for v in c.get("wa_variables") or []],
                                      phone=body.to, dedupe="", campaign_id="", force_test=True)
    if not rec or rec.get("status") == "failed":
        raise HTTPException(status_code=502, detail=(rec or {}).get("error") or "Not sent")
    await audit_log(db, "MARKETING_CAMPAIGN_TEST", email=user.get("email", ""), role=user.get("role", ""),
                    meta={"campaign": cid, "channel": "whatsapp", "to": "…" + "".join(ch for ch in body.to if ch.isdigit())[-4:]})
    return {"ok": True}


@admin_router.get("/campaigns/{cid}/preview")
async def adm_preview(cid: str):
    c = await db.mk_campaigns.find_one({"id": cid}, {"_id": 0})
    if not c:
        raise HTTPException(status_code=404, detail="Campaign not found")
    st = await get_settings()
    pc = dict(PREVIEW_CONTACT)
    await attach_picks(pc, c)  # sample: the newest releases (the sample reader has no history)
    html_doc, _ = core.render_email(c.get("blocks") or [], brand={"name": "Oakbridge Publishing", "logo": st.get("logo"), "footer": st.get("footer")},
                                    contact=pc, link_for=lambda u: u, unsubscribe_url="#",
                                    preheader=core.personalise(c.get("preheader") or "", PREVIEW_CONTACT))
    return {"html": html_doc, "subject": core.personalise(c.get("subject") or "", PREVIEW_CONTACT),
            "unknown_variables": _unknown_vars(c)}


# The sample recipient for every preview (campaigns and templates).
PREVIEW_CONTACT = {"name": "Asha Verma", "email": "asha@example.com"}


def _unknown_vars(c: dict) -> list:
    texts = [c.get("subject") or "", c.get("preheader") or ""]
    for b in c.get("blocks") or []:
        texts += [b.get("text") or "", b.get("label") or ""]
    return core.unknown_variables(*texts)


# ------------------------------------------------------------- templates ---
class TemplateBody(BaseModel):
    name: Optional[str] = None
    description: Optional[str] = None
    subject: Optional[str] = None
    preheader: Optional[str] = None
    blocks: Optional[list] = None
    from_campaign_id: Optional[str] = None
    from_template_id: Optional[str] = None


def _starters() -> list:
    return [{"id": f"starter:{t['key']}", "starter": True, "name": t["name"], "description": t["description"],
             "subject": t["subject"], "preheader": t["preheader"], "blocks": [dict(b) for b in t["blocks"]]}
            for t in core.STARTER_TEMPLATES]


async def _template(tid: str) -> dict:
    if tid.startswith("starter:"):
        t = next((x for x in _starters() if x["id"] == tid), None)
    else:
        t = await db.mk_templates.find_one({"id": tid}, {"_id": 0})
    if not t:
        raise HTTPException(status_code=404, detail="Template not found")
    return t


@admin_router.get("/templates")
async def adm_templates():
    """Saved templates first, then the built-in starters; plus the merge
    variables the renderer fills (one list — marketing_core.VARIABLES)."""
    saved = await db.mk_templates.find({}, {"_id": 0}).sort("updated_at", -1).to_list(500)
    return {"templates": saved + _starters(), "variables": core.VARIABLES}


@admin_router.post("/templates")
async def adm_new_template(body: TemplateBody, user: dict = Depends(require_admin)):
    """New template: blank, a copy of a campaign ("Save as template"), or a
    copy of another template (starters are read-only, so "Duplicate")."""
    src: dict = {}
    if body.from_campaign_id:
        src = await db.mk_campaigns.find_one({"id": body.from_campaign_id}, {"_id": 0}) or {}
        if not src:
            raise HTTPException(status_code=404, detail="Campaign not found")
        if src.get("channel") != "email":
            raise HTTPException(status_code=400, detail="Only email campaigns can be saved as templates")
    elif body.from_template_id:
        src = await _template(body.from_template_id)
    name = (body.name or (f"{src.get('name')} (copy)" if src else "") or "Untitled template").strip()[:120]
    doc = {"id": str(uuid.uuid4()), "name": name,
           "description": (body.description if body.description is not None else src.get("description", ""))[:300],
           "subject": (body.subject if body.subject is not None else src.get("subject", ""))[:300],
           "preheader": (body.preheader if body.preheader is not None else src.get("preheader", ""))[:300],
           "blocks": _clean_blocks(body.blocks if body.blocks is not None else src.get("blocks") or []),
           "created_by": user.get("email", ""), "created_at": _iso(), "updated_at": _iso(), "updated_by": user.get("email", "")}
    await db.mk_templates.insert_one(dict(doc))
    await audit_log(db, "MARKETING_TEMPLATE_CREATED", email=user.get("email", ""), role=user.get("role", ""),
                    meta={"template": doc["id"], "name": name, "from_campaign": body.from_campaign_id, "from_template": body.from_template_id})
    return doc


@admin_router.patch("/templates/{tid}")
async def adm_save_template(tid: str, body: TemplateBody, user: dict = Depends(require_admin)):
    if tid.startswith("starter:"):
        raise HTTPException(status_code=400, detail="Built-in templates can't be changed — Duplicate it and edit the copy.")
    upd = {}
    for k in ("name", "description", "subject", "preheader"):
        v = getattr(body, k)
        if v is not None:
            upd[k] = v.strip()[:120 if k == "name" else 300] if isinstance(v, str) else v
    if body.blocks is not None:
        upd["blocks"] = _clean_blocks(body.blocks)
    if "name" in upd and not upd["name"]:
        raise HTTPException(status_code=400, detail="Give the template a name")
    upd.update(updated_at=_iso(), updated_by=user.get("email", ""))
    r = await db.mk_templates.update_one({"id": tid}, {"$set": upd})
    if r.matched_count == 0:
        raise HTTPException(status_code=404, detail="Template not found")
    return await db.mk_templates.find_one({"id": tid}, {"_id": 0})


@admin_router.delete("/templates/{tid}")
async def adm_delete_template(tid: str, user: dict = Depends(require_admin)):
    """Campaigns made from a template keep their own copy, so deleting one
    changes nothing already written or sent."""
    if tid.startswith("starter:"):
        raise HTTPException(status_code=400, detail="Built-in templates can't be deleted")
    r = await db.mk_templates.delete_one({"id": tid})
    if r.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Template not found")
    await audit_log(db, "MARKETING_TEMPLATE_DELETED", email=user.get("email", ""), role=user.get("role", ""), meta={"template": tid})
    return {"ok": True}


class UseTemplateBody(BaseModel):
    name: Optional[str] = None


@admin_router.post("/templates/{tid}/use")
async def adm_use_template(tid: str, body: UseTemplateBody, user: dict = Depends(require_admin)):
    """A new draft email campaign with the template's subject, preview text
    and content. The campaign is a COPY: later edits to the template don't
    change it."""
    t = await _template(tid)
    blocks = _clean_blocks(t.get("blocks") or [])
    doc = {"id": str(uuid.uuid4()), "name": (body.name or t["name"]).strip()[:120], "channel": "email", "kind": "standard",
           "reconfirm_days": None, "subject": t.get("subject", ""), "preheader": t.get("preheader", ""),
           "blocks": blocks, "links": core.collect_links(blocks), "audience": {"list_ids": [], "rules": []},
           "wa_template": "", "wa_language": "en", "wa_variables": [], "status": "draft", "created_at": _iso(),
           "created_by": user.get("email", ""), "stats": {}, "template_id": tid, "template_name": t["name"]}
    await db.mk_campaigns.insert_one(dict(doc))
    return doc


EMAIL_IMAGE_MAX = 8 * 1024 * 1024
EMAIL_IMAGE_WIDTH = 1200   # emails show images at 600 px; 2x for sharp phone screens


def _email_image(data: bytes, ext: str) -> tuple[bytes, str]:
    """Shrink to EMAIL_IMAGE_WIDTH and re-encode, so a 6 MB camera photo does
    not land in thousands of inboxes (Gmail clips heavy emails) and anything
    smuggled after the picture data is dropped. Animated GIFs are kept as
    they are (re-encoding would freeze them)."""
    if ext == "gif":
        return data, "gif"
    try:
        from io import BytesIO
        from PIL import Image
        im = Image.open(BytesIO(data))
        im.load()
        if im.width > EMAIL_IMAGE_WIDTH:
            im = im.resize((EMAIL_IMAGE_WIDTH, round(im.height * EMAIL_IMAGE_WIDTH / im.width)))
        out = BytesIO()
        if im.mode in ("RGBA", "LA", "P") and ext == "png":
            im.save(out, "PNG", optimize=True)
            return out.getvalue(), "png"
        im.convert("RGB").save(out, "JPEG", quality=82, optimize=True, progressive=True)
        return out.getvalue(), "jpg"
    except Exception:  # noqa: BLE001 — Pillow can't read it (e.g. AVIF): keep the checked original
        return data, ext


@admin_router.post("/images")
async def adm_email_image(file: UploadFile = File(...), user: dict = Depends(require_admin)):
    """Image for an email block (drag & drop in the editor). Stored as a
    public file and returned as an absolute https URL — the only kind a mail
    client can load (and the only kind _clean_blocks keeps)."""
    import uploads_guard
    from features import put_object, APP_NAME
    data = await uploads_guard.read_limited(file, EMAIL_IMAGE_MAX, "Image")
    ext = uploads_guard.require_image(data, "Image")
    # BMP/AVIF are converted to JPEG below when Pillow can read them.
    data, ext = await asyncio.to_thread(_email_image, data, ext)
    if ext in ("bmp", "avif"):
        raise HTTPException(status_code=400, detail="Use a JPEG, PNG or GIF — many mail apps can't show this format.")
    path = f"{APP_NAME}/email/{uuid.uuid4()}.{ext}"
    await asyncio.to_thread(put_object, path, data, f"image/{'jpeg' if ext == 'jpg' else ext}")
    await audit_log(db, "MARKETING_IMAGE_UPLOADED", email=user.get("email", ""), role=user.get("role", ""),
                    meta={"path": path, "bytes": len(data)})
    return {"url": f"{PUBLIC_API}/api/files/{path}", "bytes": len(data)}


class TemplatePreviewBody(BaseModel):
    subject: Optional[str] = ""
    preheader: Optional[str] = ""
    blocks: Optional[list] = None


@admin_router.post("/templates/preview")
async def adm_template_preview(body: TemplatePreviewBody):
    """Render unsaved template content for the sample recipient."""
    st = await get_settings()
    blocks = _clean_blocks(body.blocks or [])
    html_doc, _ = core.render_email(blocks, brand={"name": "Oakbridge Publishing", "logo": st.get("logo"), "footer": st.get("footer")},
                                    contact=PREVIEW_CONTACT, link_for=lambda u: u, unsubscribe_url="#",
                                    preheader=core.personalise(body.preheader or "", PREVIEW_CONTACT))
    return {"html": html_doc, "subject": core.personalise(body.subject or "", PREVIEW_CONTACT),
            "unknown_variables": _unknown_vars({"subject": body.subject, "preheader": body.preheader, "blocks": blocks})}


class SendBody(BaseModel):
    schedule_at: Optional[str] = None   # ISO; empty = now
    confirm_count: int                  # the number the admin saw — must still match


@admin_router.post("/campaigns/{cid}/send")
async def adm_send(cid: str, body: SendBody, user: dict = Depends(require_admin)):
    c = await db.mk_campaigns.find_one({"id": cid}, {"_id": 0})
    if not c:
        raise HTTPException(status_code=404, detail="Campaign not found")
    if c["status"] not in ("draft", "scheduled"):
        raise HTTPException(status_code=409, detail="This campaign has already been sent")
    if c["channel"] == "email" and (not c.get("subject") or not c.get("blocks")):
        raise HTTPException(status_code=400, detail="Add a subject and some content first")
    bad = _unknown_vars(c) if c["channel"] == "email" else []
    if bad:
        # Unknown {{variables}} are blanked when sent ("Dear ,"), so stop here.
        raise HTTPException(status_code=400, detail=(
            f"Unknown variable(s): {', '.join('{{' + b + '}}' for b in bad)}. "
            f"Use one of: {', '.join('{{' + v['key'] + '}}' for v in core.VARIABLES)}."))
    if c["channel"] == "whatsapp" and not c.get("wa_template"):
        raise HTTPException(status_code=400, detail="Choose the approved WhatsApp template first")
    return await queue_send(c, body.schedule_at, by_email=user.get("email", ""), by_role=user.get("role", ""),
                            confirm_count=body.confirm_count)


async def queue_send(c: dict, schedule_at: Optional[str], *, by_email: str, by_role: str,
                     confirm_count: Optional[int] = None) -> dict:
    """Start sending a draft (or scheduled) campaign: the one path for both the
    admin's Send button and approved newsletter issues, so every send gets the
    same checks (audience, mailbox checks, risk rules, bounce guard)."""
    cid = c["id"]
    n = await db.mk_contacts.count_documents(await audience_filter(c))
    if not n:
        raise HTTPException(status_code=400, detail="Nobody in this audience can be sent to (no one has opted in, or all are unsubscribed / bounced)")
    if confirm_count is not None and abs(n - confirm_count) > max(5, n // 50):
        raise HTTPException(status_code=409, detail=f"The audience changed to {n} — check it again before sending")
    await _ensure_indexes()
    now = _iso()
    # Status guard: a double click can't start two preparations.
    r = await db.mk_campaigns.update_one({"id": cid, "status": {"$in": ["draft", "scheduled"]}}, {"$set": {
        "status": "preparing", "schedule_at": schedule_at or "", "sent_by": by_email, "prepare_error": "",
        "prepare": {"started_at": now, "candidates": n, "checked": 0}, "prepare_heartbeat": now}})
    if not r.modified_count:
        raise HTTPException(status_code=409, detail="This campaign is already being sent")
    await audit_log(db, "MARKETING_CAMPAIGN_SEND", email=by_email, role=by_role,
                    meta={"campaign": cid, "channel": c["channel"], "candidates": n, "schedule_at": schedule_at or ""})
    start_prepare(cid)
    return {"ok": True, "status": "preparing", "recipients": n}


@admin_router.post("/campaigns/{cid}/{action}")
async def adm_campaign_action(cid: str, action: str, user: dict = Depends(require_admin)):
    if action not in ("pause", "resume", "cancel"):
        raise HTTPException(status_code=404, detail="Unknown action")
    c = await db.mk_campaigns.find_one({"id": cid}, {"_id": 0, "status": 1})
    if not c:
        raise HTTPException(status_code=404, detail="Campaign not found")
    allowed = {"pause": ("sending", "scheduled"), "resume": ("paused",), "cancel": ("preparing", "sending", "scheduled", "paused")}[action]
    if c["status"] not in allowed:
        raise HTTPException(status_code=409, detail=f"Cannot {action} a campaign that is {c['status']}")
    new = {"pause": "paused", "resume": "sending", "cancel": "cancelled"}[action]
    await db.mk_campaigns.update_one({"id": cid}, {"$set": {"status": new, f"{new}_at": _iso(), "paused_reason": "" if action == "resume" else None}})
    if action == "cancel":
        await db.mk_sends.update_many({"campaign_id": cid, "status": "queued"}, {"$set": {"status": "skipped", "error": "cancelled"}})
    if action == "resume":
        await start_campaign(cid)
    await audit_log(db, f"MARKETING_CAMPAIGN_{action.upper()}", email=user.get("email", ""), role=user.get("role", ""), meta={"campaign": cid})
    return {"ok": True, "status": new}


_ACT_WORDS = {
    "MARKETING_CAMPAIGN_SEND": "Pressed Send", "MARKETING_CAMPAIGN_PAUSE": "Paused", "MARKETING_CAMPAIGN_RESUME": "Resumed",
    "MARKETING_CAMPAIGN_CANCEL": "Cancelled", "MARKETING_CAMPAIGN_RENAMED": "Renamed", "MARKETING_CAMPAIGN_EDITED": "Edited",
    "MARKETING_CAMPAIGN_TEST": "Sent a test", "MARKETING_AI_ACTION_DRAFT": "Created from 'Do this next'",
}


def _mins(a: Optional[str], b: Optional[str]) -> Optional[float]:
    try:
        return round((datetime.fromisoformat(b) - datetime.fromisoformat(a)).total_seconds() / 60, 1) if a and b else None
    except ValueError:
        return None


@admin_router.get("/campaigns/{cid}/activity")
async def adm_campaign_activity(cid: str):
    """Who did what, when, to this campaign — and where its audience came from.

    One timeline from three sources: the campaign's own fields (created,
    checks started/finished, sending started/finished, automatic pauses),
    the audit log (Send, Pause, Resume, Cancel, edits, test sends, renames —
    each with the person's email), and the sends themselves (first delivery,
    open, click, order). The audience section names every list with who
    created it, who APPROVED its import and when, and the file it came from.
    """
    c = await db.mk_campaigns.find_one({"id": cid}, {"_id": 0, "blocks": 0})
    if not c:
        raise HTTPException(status_code=404, detail="Campaign not found")
    ev = [{"at": c.get("created_at"), "who": c.get("created_by") or "", "what": "Created",
           "detail": f"from template “{c['template_name']}”" if c.get("template_name") else (c.get("note") or "")[:160]}]
    try:
        await db.audit_events.create_index("meta.campaign")
    except Exception:  # noqa: BLE001
        pass
    async for a in db.audit_events.find({"meta.campaign": cid}, {"_id": 0}).sort("at", 1).limit(500):
        m = a.get("meta") or {}
        detail = ""
        if a["action"] == "MARKETING_CAMPAIGN_SEND":
            detail = (f"to {m.get('candidates')} possible recipients"
                      + (f", scheduled for {m['schedule_at']}" if m.get("schedule_at") else ""))
        elif a["action"] == "MARKETING_CAMPAIGN_EDITED":
            detail = f"changed: {m.get('fields') or '—'}"
        elif a["action"] == "MARKETING_CAMPAIGN_TEST":
            detail = f"to {m.get('to')}"
        ev.append({"at": a.get("at"), "who": a.get("email") or "", "what": _ACT_WORDS.get(a["action"], a["action"].title()),
                   "detail": detail})
    p = c.get("prepare") or {}
    if p.get("started_at"):
        ev.append({"at": p["started_at"], "who": "system", "what": "Checking addresses",
                   "detail": f"{p.get('candidates', 0)} candidates"})
    if p.get("finished_at"):
        sk = p.get("skipped") or {}
        ev.append({"at": p["finished_at"], "who": "system", "what": "Checks finished",
                   "detail": f"{c.get('queued', 0)} queued to send"
                             + (f", {sum(int(v or 0) for v in sk.values()) if isinstance(sk, dict) else sk} skipped" if sk else "")})
    if c.get("started_at"):
        ev.append({"at": c["started_at"], "who": "system", "what": "Sending started", "detail": ""})
    if c.get("paused_reason") and c.get("status") == "paused":
        ev.append({"at": c.get("paused_at"), "who": "system", "what": "Paused automatically", "detail": c["paused_reason"][:200]})
    if c.get("tail_stopped_at"):
        ev.append({"at": c["tail_stopped_at"], "who": "system", "what": "Stopped the unconfirmed tail",
                   "detail": "bounces from unconfirmed addresses passed the limit"})
    if c.get("finished_at"):
        ev.append({"at": c["finished_at"], "who": "system", "what": "Finished sending", "detail": ""})
    firsts = {}
    for f, label in (("delivered_at", "First delivered"), ("opened_at", "First opened"), ("clicked_at", "First click")):
        x = await db.mk_sends.find_one({"campaign_id": cid, f: {"$exists": True, "$ne": None}}, {"_id": 0, f: 1}, sort=[(f, 1)])
        if x and x.get(f):
            firsts[f] = x[f]
            ev.append({"at": x[f], "who": "recipient", "what": label, "detail": ""})
    tag = core.campaign_tag(cid)
    o = await db.orders.find_one({"attribution.utm_campaign": tag, "payment_status": "paid"}, {"_id": 0, "created_at": 1, "order_number": 1},
                                 sort=[("created_at", 1)])
    if o:
        ev.append({"at": o.get("created_at"), "who": "customer", "what": "First paid order", "detail": o.get("order_number", "")})
    ev = sorted([e for e in ev if e.get("at")], key=lambda e: e["at"])

    # ---- where the audience came from ----
    aud = c.get("audience") or {}
    lists = []
    for lid in aud.get("list_ids") or []:
        lst = await db.mk_lists.find_one({"id": lid}, {"_id": 0, "id": 1, "name": 1, "kind": 1, "created_by": 1, "created_at": 1,
                                                       "approved_by": 1, "source": 1, "rules": 1, "audience": 1})
        if not lst:
            lists.append({"id": lid, "name": "(deleted list)", "missing": True})
            continue
        imp = await db.mk_imports.find_one({"list_id": lid}, {"_id": 0, "file": 1, "by": 1, "at": 1, "approved_by": 1,
                                                             "approved_at": 1, "saved": 1, "audience": 1})
        lists.append({**lst, "members": await db.mk_contacts.count_documents({"lists": lid}),
                      "import": imp or None})
    excluded = [x.get("name") async for x in db.mk_lists.find({"id": {"$in": aud.get("exclude_list_ids") or []}}, {"_id": 0, "name": 1})]

    # ---- metrics ----
    s = c.get("stats") or {}
    sent = int(s.get("sent") or 0)
    dur = _mins(c.get("started_at"), c.get("finished_at"))
    metrics = {
        "possible_recipients": p.get("candidates"), "queued": c.get("queued"), "sent": sent,
        "delivered": int(s.get("delivered") or 0), "opened": int(s.get("opened") or 0), "clicked": int(s.get("clicked") or 0),
        "bounced": int(s.get("bounced") or 0), "complained": int(s.get("complained") or 0), "unsubscribed": int(s.get("unsubscribed") or 0),
        "orders": int(s.get("paid") or 0), "revenue": round(float(s.get("revenue") or 0), 2),
        "revenue_per_100_sent": round(float(s.get("revenue") or 0) * 100 / sent, 2) if sent else None,
        "checks_minutes": _mins(p.get("started_at"), p.get("finished_at")),
        "sending_minutes": dur, "per_minute": round(sent / dur, 1) if dur and dur > 0 else None,
        "minutes_to_first_open": _mins(c.get("started_at"), firsts.get("opened_at")),
        "minutes_to_first_click": _mins(c.get("started_at"), firsts.get("clicked_at")),
        "sent_by": c.get("sent_by") or "", "created_by": c.get("created_by") or "",
        "last_edited_by": c.get("updated_by") or "", "last_edited_at": c.get("updated_at"),
    }
    return {"events": ev, "lists": lists, "excluded_lists": excluded, "rules": aud.get("rules") or [], "metrics": metrics}


@admin_router.get("/campaigns/{cid}/report")
async def adm_report(cid: str):
    c = await db.mk_campaigns.find_one({"id": cid}, {"_id": 0})
    if not c:
        raise HTTPException(status_code=404, detail="Campaign not found")
    stats = await refresh_campaign_stats(cid)
    links = c.get("links") or []
    clicks = [0] * len(links)
    async for s in db.mk_sends.find({"campaign_id": cid, "clicks.0": {"$exists": True}}, {"_id": 0, "clicks": 1}):
        for k in {x["i"] for x in s.get("clicks") or []}:
            if 0 <= k < len(clicks):
                clicks[k] += 1
    timeline: dict = {}
    async for s in db.mk_sends.find({"campaign_id": cid}, {"_id": 0, "opened_at": 1, "clicked_at": 1}):
        for f in ("opened_at", "clicked_at"):
            if s.get(f):
                h = s[f][:13]
                timeline.setdefault(h, {"opened": 0, "clicked": 0})
                timeline[h][f[:-3]] += 1
    # Soft bounces too (status stays "sent"), so every non-delivered address
    # can be seen, not just the hard ones.
    problems = await db.mk_sends.find({"campaign_id": cid, "$or": [
                                          {"status": {"$in": ["bounced", "complained", "failed", "blocked"]}},
                                          {"status": "sent", "bounce_type": {"$exists": True}, "delivered_at": {"$exists": False}}]},
                                      {"_id": 0, "to": 1, "status": 1, "bounce_type": 1, "error": 1}).limit(200).to_list(200)
    by_provider: dict = {}
    async for x in db.mk_sends.find({"campaign_id": cid, "status": {"$in": ["sent", "bounced", "complained"]}},
                                    {"_id": 0, "to": 1, "status": 1, "delivered_at": 1}):
        dom = str(x.get("to") or "").lower().rsplit("@", 1)[-1] or "?"
        p = by_provider.setdefault(dom, {"provider": dom, "sent": 0, "bounced": 0, "delivered": 0})
        p["sent"] += 1
        p["bounced"] += 1 if x.get("status") == "bounced" else 0
        p["delivered"] += 1 if x.get("delivered_at") else 0
    providers = sorted(by_provider.values(), key=lambda p: (-p["bounced"], -p["sent"]))[:20]
    for p in providers:
        p["bounce_rate"] = core.rate(p["bounced"], p["sent"])
    skipped = [{"reason": r["_id"] or "skipped", "n": r["n"]} async for r in db.mk_sends.aggregate([
        {"$match": {"campaign_id": cid, "status": "skipped"}}, {"$group": {"_id": "$error", "n": {"$sum": 1}}}, {"$sort": {"n": -1}}])]
    return {"campaign": {k: c.get(k) for k in ("id", "name", "channel", "subject", "status", "started_at", "finished_at", "paused_reason",
                                               "prepare", "prepare_error", "recovery", "tail_stopped_at", "note",
                                               "kind", "reconfirm_days", "lapsed_count", "lapsed_at")},
            "skipped": skipped, "providers": providers,
            "stats": stats, "funnel": core.funnel(stats), "links": [{"url": u, "clicks": n} for u, n in zip(links, clicks)],
            "timeline": dict(sorted(timeline.items())), "problems": problems,
            "rates": {"delivery": core.rate(stats.get("delivered", 0), stats.get("sent", 0)),
                      "open": core.rate(stats.get("opened", 0), stats.get("delivered", 0) or stats.get("sent", 0)),
                      "click": core.rate(stats.get("clicked", 0), stats.get("delivered", 0) or stats.get("sent", 0)),
                      "ctor": core.rate(stats.get("clicked", 0), stats.get("opened", 0)),
                      "bounce": core.rate(stats.get("bounced", 0), stats.get("sent", 0)),
                      "soft_bounce": core.rate(stats.get("soft_bounced", 0), stats.get("sent", 0)),
                      "unsubscribe": core.rate(stats.get("unsubscribed", 0), stats.get("delivered", 0) or stats.get("sent", 0))}}


async def _accuracy() -> dict:
    """How right were we? Hard-bounce rate by what we believed when we sent
    (last 90 days, settled a day), plus what the learning job has changed."""
    match = {"sent_at": {"$gte": (_now() - timedelta(days=90)).isoformat(), "$lte": (_now() - timedelta(days=1)).isoformat()},
             "status": {"$in": ["sent", "bounced", "complained"]}, "band": {"$exists": True}}
    bands: dict = {}
    async for r in db.mk_sends.aggregate([{"$match": match}, {"$group": {
            "_id": "$band", "sent": {"$sum": 1}, "bounced": {"$sum": {"$cond": [{"$eq": ["$status", "bounced"]}, 1, 0]}}}}]):
        bands[r["_id"] or "unknown"] = {"sent": r["sent"], "bounced": r["bounced"], "rate": core.rate(r["bounced"], r["sent"])}
    st = await get_settings()
    th, adjust = await _scoring_params(st)
    model = await scoring_model()
    return {"bands": bands, "thresholds": th, "auto": bool(st.get("auto_learn", True)), "adjust": adjust,
            "labels": core.SIGNAL_LABELS, "base": core.BASE_DELTAS, "learned_at": model.get("updated_at"),
            "samples": ((model.get("stats") or {}).get("per_flag") or {})}


# ---- dashboard ----
_DASH_MEMO = None


@admin_router.get("/dashboard")
async def adm_dashboard(days: int = 30):
    """Served from a 60-second memo (dropped by any admin write); the numbers
    themselves are counted inside MongoDB, so opening the dashboard no longer
    streams every send and contact through this process."""
    global _DASH_MEMO
    days = max(1, min(days, 365))
    if _DASH_MEMO is None:
        from cache import Memo, shared_version
        _DASH_MEMO = Memo(shared_version(db))
    return await _DASH_MEMO.get(f"dash:{days}", 60, lambda: _dashboard(days))


async def _by_day(coll, field: str, match: dict, value=None) -> dict:
    """{YYYY-MM-DD: count (or sum of `value`)} grouped in the database."""
    grp = {"_id": {"$substrBytes": [f"${field}", 0, 10]}, "n": {"$sum": value if value else 1}}
    return {r["_id"]: r["n"] async for r in coll.aggregate([{"$match": match}, {"$group": grp}]) if r["_id"]}


async def refresh_recent_stats() -> int:
    """Background (tick): keep stored stats of campaigns from the last 30 days
    fresh, so the dashboard never recomputes them while someone waits."""
    stale = (_now() - timedelta(minutes=10)).isoformat()
    since = (_now() - timedelta(days=30)).isoformat()
    n = 0
    async for c in db.mk_campaigns.find({"started_at": {"$gte": since}, "status": {"$in": ["sending", "sent", "paused"]},
                                         "deleted": {"$ne": True}, "$or": [{"stats_at": {"$lt": stale}}, {"stats_at": {"$exists": False}}]},
                                        {"_id": 0, "id": 1}):
        await refresh_campaign_stats(c["id"])
        n += 1
    return n


async def _dashboard(days: int) -> dict:
    since = (_now() - timedelta(days=days)).isoformat()
    camps = await db.mk_campaigns.find({"started_at": {"$gte": since}, "deleted": {"$ne": True}}, {"_id": 0, "blocks": 0}).sort("started_at", -1).to_list(200)
    gone = [c["id"] async for c in db.mk_campaigns.find({"deleted": True}, {"_id": 0, "id": 1})]
    gone_tags = [core.campaign_tag(x) for x in gone]
    tot = {"email": {}, "whatsapp": {}}
    for c in camps:
        t = tot[c["channel"]]
        for k, v in (c.get("stats") or {}).items():
            if isinstance(v, (int, float)):
                t[k] = round(t.get(k, 0) + v, 2)
    series: dict = {}
    live = {"campaign_id": {"$nin": gone}}  # archived (deleted) campaigns drop out of the charts too
    parts = {
        "sent": await _by_day(db.mk_sends, "sent_at", {**live, "sent_at": {"$gte": since}}),
        "opened": await _by_day(db.mk_sends, "opened_at", {**live, "opened_at": {"$gte": since}}),
        "clicked": await _by_day(db.mk_sends, "clicked_at", {**live, "clicked_at": {"$gte": since}}),
        "bounced": await _by_day(db.mk_sends, "sent_at", {**live, "sent_at": {"$gte": since}, "status": "bounced"}),
        "revenue": await _by_day(db.orders, "created_at", {"created_at": {"$gte": since}, "payment_status": "paid",
                                                            "attribution.utm_campaign": {"$regex": "^mk-", "$nin": gone_tags}},
                                 value={"$ifNull": ["$total", 0]}),
    }
    for key, days_map in parts.items():
        for day, n in days_map.items():
            series.setdefault(day, {})[key] = round(n, 2) if key == "revenue" else n
    growth = await _by_day(db.mk_contacts, "created_at", {"created_at": {"$gte": since}})
    unsubs = await _by_day(db.mk_contacts, "email_consent.at", {"email_consent.status": {"$in": ["unsubscribed", "complained"]},
                                                                 "email_consent.at": {"$gte": since}})
    e = tot["email"]
    kpi = {
        "contacts": await db.mk_contacts.count_documents({}),
        "subscribed": await db.mk_contacts.count_documents({"email_consent.status": "subscribed", "suppressed": {"$in": [None, ""]}}),
        "wa_subscribed": await db.mk_contacts.count_documents({"wa_consent.status": "subscribed"}),
        "campaigns": len(camps),
        "email_sent": e.get("sent", 0), "delivery_rate": core.rate(e.get("delivered", 0), e.get("sent", 0)),
        "open_rate": core.rate(e.get("opened", 0), e.get("delivered", 0) or e.get("sent", 0)),
        "click_rate": core.rate(e.get("clicked", 0), e.get("delivered", 0) or e.get("sent", 0)),
        "bounce_rate": core.rate(e.get("bounced", 0), e.get("sent", 0)),
        "complaint_rate": core.rate(e.get("complained", 0), e.get("sent", 0)),
        "unsubscribe_rate": core.rate(e.get("unsubscribed", 0), e.get("delivered", 0) or e.get("sent", 0)),
        "wa_sent": tot["whatsapp"].get("sent", 0), "wa_read_rate": core.rate(tot["whatsapp"].get("read", 0), tot["whatsapp"].get("sent", 0)),
        "wa_cost": tot["whatsapp"].get("cost", 0),
        "revenue": round(sum((c.get("stats") or {}).get("revenue", 0) for c in camps), 2),
        "orders": sum((c.get("stats") or {}).get("paid", 0) for c in camps),
    }
    by_status = {r["_id"]: r["n"] async for r in db.mk_contacts.aggregate([{"$group": {"_id": "$email_status", "n": {"$sum": 1}}}])}
    return {"days": days, "kpi": kpi, "series": dict(sorted(series.items())), "growth": dict(sorted(growth.items())),
            "unsubs": dict(sorted(unsubs.items())), "funnel_email": core.funnel(e), "funnel_whatsapp": core.funnel(tot["whatsapp"]),
            "campaigns": [{k: c.get(k) for k in ("id", "name", "channel", "status", "started_at", "stats")} for c in camps],
            "verification": by_status,
            "alerts": await db.mk_alerts.find({"resolved": False}, {"_id": 0}).sort("at", -1).limit(20).to_list(20),
            "accuracy": await _accuracy(),
            "ses_usage": await ses_usage()}


# ================================================================ tasks ===
tasks_router = APIRouter(prefix="/api/tasks", tags=["tasks"])


@tasks_router.post("/marketing-tick")
async def marketing_tick(x_task_token: Optional[str] = Header(None)):
    """Cron (every few minutes). With a worker running, the worker does this
    itself every minute too; dedupe keys make the overlap harmless."""
    expected = os.environ.get("TASK_TOKEN")
    if not expected or x_task_token != expected:
        raise HTTPException(status_code=401, detail="Unauthorized")
    return await periodic()


async def fix_waitlist_consent() -> int:
    """One-time correction (idempotent): earlier syncs marked waitlist sign-ups
    as subscribed. Withdraw that consent unless they confirmed the newsletter
    themselves (double opt-in) or opted in elsewhere since."""
    r = await db.mk_contacts.update_many(
        {"email_consent.status": "subscribed", "email_consent.source": {"$regex": "waitlist", "$options": "i"}},
        {"$set": {"email_consent": _consent("none", "waitlist sign-up is not a newsletter opt-in (DPDP purpose limitation)")}})
    if r.modified_count:
        log.warning("marketing: withdrew %s waitlist-only marketing consents", r.modified_count)
        await audit_log(db, "MARKETING_WAITLIST_CONSENT_FIXED", email="system", role="system", meta={"count": r.modified_count})
    return r.modified_count


async def periodic() -> dict:
    """Start scheduled campaigns that are due, resume ones whose sender stopped
    (e.g. a deploy), and queue the routine background work."""
    now = _iso()
    started = []
    async for c in db.mk_campaigns.find({"status": "scheduled", "schedule_at": {"$lte": now}}, {"_id": 0, "id": 1}):
        r = await db.mk_campaigns.update_one({"id": c["id"], "status": "scheduled"}, {"$set": {"status": "sending", "started_at": now}})
        if r.modified_count:
            await start_campaign(c["id"])
            started.append(c["id"])
    stale = (_now() - timedelta(seconds=WORKER_STALE)).isoformat()
    async for c in db.mk_campaigns.find({"status": "sending", "$or": [{"heartbeat": {"$lt": stale}}, {"heartbeat": {"$exists": False}}]}, {"_id": 0, "id": 1}):
        await start_campaign(c["id"])
        started.append(c["id"])
    stale_prep = (_now() - timedelta(seconds=PREPARE_STALE)).isoformat()
    async for c in db.mk_campaigns.find({"status": "preparing", "prepare_heartbeat": {"$lt": stale_prep}}, {"_id": 0, "id": 1}):
        start_prepare(c["id"])   # e.g. a deploy restarted the server mid-check
        started.append(c["id"])
    kick_validation()
    if _jobs.worker_mode() == "external":
        try:
            lag = (await _jobs.stats(db)).get("lag_seconds") or 0
            if lag > 600:
                await raise_alert("queue_lag", f"Background work is {lag // 60} minutes behind. Check the worker on Render.",
                                  key="queue_lag")
        except Exception:  # noqa: BLE001
            log.warning("marketing: queue lag check failed", exc_info=True)
    hour = _now().strftime("%Y%m%d%H")
    ten_min = _now().strftime("%Y%m%d%H") + str(_now().minute // 10)
    dispatch("mk_stats", dedupe=ten_min)
    dispatch("nl_tick", dedupe=ten_min)          # newsletter: draft / approve / send on schedule (newsletter.py)
    dispatch("mk_waitlist_fix", dedupe=_now().strftime("%Y%m%d"))
    dispatch("mk_suppression", dedupe=hour)   # no-op unless 24h have passed
    dispatch("mk_learn", dedupe=hour)         # likewise daily
    dispatch("mk_disposable", dedupe=hour)    # weekly
    dispatch("mk_lapse", dedupe=hour)
    dispatch("mk_auto", dedupe=hour)          # abandoned-checkout + review emails (growth.py), when switched on
    try:
        import amazon_spapi
        amazon_spapi.schedule_tick()          # Amazon FBA stock, every 6 h (only when the SPAPI_* keys are set)
    except Exception:  # noqa: BLE001
        log.warning("amazon: schedule tick failed", exc_info=True)
    return {"ok": True, "started": started}


async def resume_on_startup() -> None:
    """After a deploy, carry on with anything that was mid-send. With a live
    worker this is the worker's job (it runs periodic() every minute)."""
    try:
        await asyncio.sleep(10)
        if _jobs.worker_mode() == "external" and await _jobs.worker_alive(db):
            return
        async for c in db.mk_campaigns.find({"status": "sending"}, {"_id": 0, "id": 1}):
            await start_campaign(c["id"])
        async for c in db.mk_campaigns.find({"status": "preparing"}, {"_id": 0, "id": 1}):
            start_prepare(c["id"])
    except Exception:  # noqa: BLE001
        log.exception("marketing: resume on startup failed")


# ------------------------------------------------------------ job kinds ---
# What worker.py (or, inline, this process) runs for each dispatch() kind.
# Every handler is safe to run twice: a job can be retried after a crash.
@_jobs.handler("mk_validate")
async def _job_validate(p: dict) -> None:
    await validate_pending()


@_jobs.handler("mk_prepare")
async def _job_prepare(p: dict) -> None:
    await _prepare_campaign(p["cid"])


@_jobs.handler("mk_send")
async def _job_send(p: dict) -> None:
    await _run_campaign(p["cid"])


@_jobs.handler("mk_role_rule")
async def _job_role_rule(p: dict) -> None:
    await _reapply_role_rule()


@_jobs.handler("mk_suppression")
async def _job_suppression(p: dict) -> None:
    await sync_aws_suppression()


@_jobs.handler("mk_learn")
async def _job_learn(p: dict) -> None:
    await learn_from_outcomes()


@_jobs.handler("mk_disposable")
async def _job_disposable(p: dict) -> None:
    await refresh_disposable()


@_jobs.handler("mk_lapse")
async def _job_lapse(p: dict) -> None:
    await lapse_reconfirm()


@_jobs.handler("mk_newsletter")
async def _job_newsletter(p: dict) -> None:
    await _newsletter_signup(p.get("email", ""), p.get("source") or "newsletter")


@_jobs.handler("mk_stats")
async def _job_stats(p: dict) -> None:
    await refresh_recent_stats()


@_jobs.handler("mk_waitlist_fix")
async def _job_waitlist_fix(p: dict) -> None:
    await fix_waitlist_consent()
