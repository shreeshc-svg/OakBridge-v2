"""Warehouse inventory as files to share (Oct 2026). Pure: rows in, bytes out.

  inventory_xlsx   for founders / internal reports — a Summary sheet (totals,
                   out of stock, low, differences) and the full Inventory
                   sheet, formatted, filterable, header frozen.
  brochure_pdf     for book vendors — a catalogue of titles with cover, author,
                   ISBN, edition, binding, pages, MRP and availability, grouped
                   by category. It deliberately shows NO internal numbers (no
                   shelf counts, no owed copies, no differences): vendors see
                   what they can order, not how much we hold.

Called by warehouse.py (Admin → Warehouse → Inventory → Export / Brochure).
"""
from __future__ import annotations

import io
import re
from datetime import datetime
from typing import Optional
from xml.sax.saxutils import escape

INVENTORY_HEADER = ["Title", "Author", "ISBN", "Category", "Price (Rs)", "MRP (Rs)", "Warehouse count", "Website stock",
                    "Owed to orders", "Difference", "Status", "Last movement", "Last reason", "Movements"]


def status_of(r: dict, threshold: int) -> str:
    if r.get("coming_soon"):
        return "Pre-order"
    wh = r.get("wh_stock")
    if wh is None:
        return "Never counted"
    if wh <= 0:
        return "Print on demand" if r.get("pod") else "Out of stock"
    return "Low" if wh <= threshold else "In stock"


def _safe(v):
    """No formula injection when the file is opened in Excel."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


def inventory_xlsx(rows: list, threshold: int, mode: str, cat_names: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    navy = PatternFill("solid", fgColor="002B5C")
    white_b = Font(bold=True, color="FFFFFF")
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    units = sum(max(0, r.get("wh_stock") or 0) for r in rows)
    value = sum(max(0, r.get("wh_stock") or 0) * float(r.get("price") or 0) for r in rows)
    st = [status_of(r, threshold) for r in rows]
    summary = [
        ("Oakbridge Publishing — warehouse inventory", ""),
        ("Generated", datetime.now().strftime("%d %b %Y, %H:%M")),
        ("Warehouse mode", mode),
        ("", ""),
        ("Titles", len(rows)),
        ("Units on the shelf", units),
        ("Shelf value at selling price (Rs)", round(value, 2)),
        ("In stock", st.count("In stock")),
        (f"Low (≤ {threshold})", st.count("Low")),
        ("Out of stock", st.count("Out of stock")),
        ("Print on demand (no copies, sold to order)", st.count("Print on demand")),
        ("Pre-order", st.count("Pre-order")),
        ("Never counted", st.count("Never counted")),
        ("Different from website stock", sum(1 for r in rows if r.get("diff"))),
        ("Copies owed to waiting orders", sum(int(r.get("owed") or 0) for r in rows)),
    ]
    for k, v in summary:
        ws.append([k, v])
    ws["A1"].font = Font(bold=True, size=14, color="002B5C")
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 22
    ws["B7"].number_format = "#,##0.00"

    inv = wb.create_sheet("Inventory")
    inv.append(INVENTORY_HEADER)
    for c in inv[1]:
        c.fill, c.font = navy, white_b
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        inv.append([_safe(r.get("title", "")), _safe(r.get("author") or ""), r.get("isbn") or "",
                    cat_names.get(r.get("category"), r.get("category") or ""), float(r.get("price") or 0),
                    float(r.get("original_price") or r.get("price") or 0),
                    r.get("wh_stock"), int(r.get("stock") or 0), int(r.get("owed") or 0), r.get("diff"),
                    status_of(r, threshold), (r.get("last_at") or "")[:16].replace("T", " "), r.get("last_reason") or "",
                    int(r.get("moves") or 0)])
    widths = [48, 28, 16, 18, 11, 11, 14, 13, 13, 11, 16, 17, 16, 11]
    for i, w in enumerate(widths, start=1):
        inv.column_dimensions[get_column_letter(i)].width = w
    for row in inv.iter_rows(min_row=2, min_col=5, max_col=6):
        for c in row:
            c.number_format = "#,##0"
    red = Font(color="CC0033", bold=True)
    for row in inv.iter_rows(min_row=2):
        wh = row[6].value
        if isinstance(wh, (int, float)) and wh < 0:
            row[6].font = red
        if row[9].value:
            row[9].font = red
    inv.freeze_panes = "B2"
    inv.auto_filter.ref = f"A1:{get_column_letter(len(INVENTORY_HEADER))}{max(1, inv.max_row)}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------- brochure ---
def latin(s) -> str:
    """Helvetica has no Indic or ₹ glyphs: keep the Latin part of a title
    ("Posh Act An Overview (Tamil) போஷ் …" -> "Posh Act An Overview (Tamil)")."""
    s = str(s or "").replace("₹", "Rs ").replace("–", "-").replace("—", "-")
    s = re.sub(r"[^\x00-ɏ‘’“”…]+", " ", s)
    return re.sub(r"\s+", " ", s).strip(" -·")


def availability(r: dict) -> str:
    if r.get("coming_soon"):
        return "Pre-order"
    if int(r.get("stock") or 0) > 0:
        return "In stock"
    return "Print on demand" if r.get("pod") else "Out of stock"


def brochure_pdf(rows: list, cat_names: dict, covers: Optional[dict] = None, *, title: str = "Catalogue",
                 contact: str = "info@oakbridge.in · www.oakbridge.in") -> bytes:
    """A4 catalogue: grouped by category, one row per title, cover on the left
    (when given), MRP and availability on the right. `covers`: {book_id: jpeg bytes}."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.lib.utils import ImageReader
    from reportlab.platypus import (Image, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table,
                                    TableStyle)

    navy, amber, grey = colors.HexColor("#002B5C"), colors.HexColor("#B45309"), colors.HexColor("#4B5563")
    h_title = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=26, leading=30, textColor=navy)
    h_sub = ParagraphStyle("s", fontName="Helvetica", fontSize=11, leading=15, textColor=grey)
    h_cat = ParagraphStyle("c", fontName="Helvetica-Bold", fontSize=15, leading=19, textColor=navy, spaceBefore=6, spaceAfter=6)
    b_title = ParagraphStyle("bt", fontName="Helvetica-Bold", fontSize=10.5, leading=13, textColor=navy)
    b_meta = ParagraphStyle("bm", fontName="Helvetica", fontSize=8.5, leading=11, textColor=grey)
    b_price = ParagraphStyle("bp", fontName="Helvetica-Bold", fontSize=11, leading=14, textColor=navy, alignment=2)
    b_av = ParagraphStyle("ba", fontName="Helvetica", fontSize=8, leading=10, textColor=amber, alignment=2)

    buf = io.BytesIO()
    stamp = datetime.now().strftime("%B %Y")

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(grey)
        canvas.drawString(15 * mm, 9 * mm, f"Oakbridge Publishing · {title} · {stamp}")
        canvas.drawRightString(A4[0] - 15 * mm, 9 * mm, f"{contact} · page {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=14 * mm,
                            bottomMargin=16 * mm, title=f"Oakbridge Publishing — {title}", author="Oakbridge Publishing")
    flow = [Spacer(1, 40 * mm), Paragraph("Oakbridge Publishing", h_title), Spacer(1, 4 * mm),
            Paragraph(escape(f"{title} · {stamp}"), h_sub), Spacer(1, 2 * mm),
            Paragraph(escape(f"{len(rows)} titles · Law, Tax, Academic and General · {contact}"), h_sub),
            Spacer(1, 8 * mm),
            Paragraph("Prices are MRP in Indian rupees. Availability as on the date above; print-on-demand titles are "
                      "printed to order and dispatched in 5-7 working days. For trade terms and bulk orders write to "
                      "info@oakbridge.in.", h_sub),
            PageBreak()]
    groups: dict = {}
    for r in rows:
        groups.setdefault(r.get("category") or "", []).append(r)
    order = sorted(groups, key=lambda c: (list(cat_names).index(c) if c in cat_names else 99, c))
    for cat in order:
        flow.append(Paragraph(escape(latin(cat_names.get(cat) or cat.title() or "Other")), h_cat))
        for r in sorted(groups[cat], key=lambda x: latin(x.get("title")).lower()):
            meta = " · ".join(x for x in [
                f"ISBN {r['isbn']}" if r.get("isbn") else "",
                f"{r['edition']}/e" if str(r.get("edition") or "").strip() not in ("", "1", "1.0", "None") else "",
                r.get("binding") or "", f"{r['pages']} pages" if r.get("pages") else "",
                str(r.get("publication_year") or "")] if x)
            mrp = float(r.get("original_price") or r.get("price") or 0)
            text = [Paragraph(escape(latin(r.get("title"))), b_title),
                    Paragraph(escape(latin(r.get("author"))), b_meta), Paragraph(escape(latin(meta)), b_meta)]
            right = [Paragraph(f"Rs {mrp:,.0f}" if mrp else "", b_price), Paragraph(availability(r), b_av)]
            img = None
            data = (covers or {}).get(r.get("id"))
            if data:
                try:
                    ir = ImageReader(io.BytesIO(data))
                    w, h = ir.getSize()
                    iw = 16 * mm
                    img = Image(io.BytesIO(data), width=iw, height=min(24 * mm, iw * h / max(w, 1)))
                except Exception:  # noqa: BLE001 — a bad cover never stops the brochure
                    img = None
            cells = [[img or "", text, right]] if covers is not None else [[text, right]]
            widths = [20 * mm, 120 * mm, 40 * mm] if covers is not None else [140 * mm, 40 * mm]
            t = Table(cells, colWidths=widths)
            t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#E5E7EB")),
                                   ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                                   ("LEFTPADDING", (0, 0), (-1, -1), 2)]))
            flow.append(KeepTogether(t))
        flow.append(Spacer(1, 4 * mm))
    doc.build(flow, onFirstPage=lambda c, d: None, onLaterPages=footer)
    return buf.getvalue()


def thumbnail(data: bytes, px: int = 220) -> Optional[bytes]:
    """Covers shrunk for print (keeps the PDF a few MB, not a few hundred)."""
    try:
        from PIL import Image as PImage
        im = PImage.open(io.BytesIO(data)).convert("RGB")
        im.thumbnail((px, px * 2))
        out = io.BytesIO()
        im.save(out, "JPEG", quality=78)
        return out.getvalue()
    except Exception:  # noqa: BLE001
        return None
